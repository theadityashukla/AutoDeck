"""The GATE 3 review surface (Phase 3b, task 3b.10): the final deck, judged as rendered.

GATE 3 is the owner approving the finished file. This module assembles what they need to
do that honestly — every check code can make on the **rendered** PPTX, and the five checks
code cannot make, listed as a checklist for a person — and never decides. Same stance as
`gate1.py` and `gate2`: a clean run here is not the gate; the owner is.

## What is checked, and on what

Everything is computed from the files on disk at the moment `autodeck gate3` runs: the
latest IR and `runs/<run>/deck.pptx`. Nothing is carried over from the render command's
memory, so the report always describes the deck whose fingerprint `approve ... final_render`
will record (A7) — not a deck that existed earlier in the session.

- **Post-render audit** (`audit/post_render.py`): every claim survived onto its slide, every
  number on a rendered slide is traced (A2, re-linted on extracted text — task 3b.8's
  detection half), the slide set matches the IR, no stray image parts.
- **Deterministic QA** (`render/qa/deterministic.py`): overlap, safe area, minimum size,
  contrast — arithmetic, no model.
- **Grammar** (`design/grammar.py`): the IR lints for every slide's assigned mode, plus icon
  adjacency on the rendered presentation.
- **Header flow** (`design/headers/flow.py`) on the final deck, with the brief, so the owner
  reads the headers as the client will see them after art direction and critique.
- **Milestone shape** (PHASE-3B.md): at least one native diagram and at least one
  theme-recolourable icon in the rendered file — the two things checks 1 and 2 need to
  exist before a person can try them.

## What is not checked here

PHASE-3B.md's checks 1-5 are hands-on-keyboard in PowerPoint: selecting and recolouring a
diagram node and an icon, editing a chart's data, seeing the theme in the Design UI, and
adding a slide that inherits the master. They are printed as `HUMAN_CHECKS`, unticked,
every time. The phase doc says why: "D1, D10 and D11 all fail silently if verified any
other way." Nothing in this module may mark them passed.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Final

from pptx import Presentation
from pptx.presentation import Presentation as PresentationType

from autodeck.audit.manifest import canonical_pptx_digest
from autodeck.audit.post_render import PostRenderReport, post_render_audit
from autodeck.design import grammar
from autodeck.design.grammar import GrammarFinding
from autodeck.design.headers.flow import HeaderFlowReport, flow_report
from autodeck.design.headers.profile import HeaderStyleProfile
from autodeck.design.icons.consistency import ICON_NAME_PREFIX
from autodeck.design.theme.tokens import DesignTokens
from autodeck.ir.models import Deck, DeckBrief
from autodeck.render.qa.deterministic import QAFinding, run_deterministic_qa

HUMAN_CHECKS: Final[tuple[str, ...]] = (
    "A diagram node can be selected and recoloured.",
    "An icon can be selected and recoloured from the theme palette.",
    "A chart's underlying data can be edited.",
    "Theme colours and fonts appear in PowerPoint's own Design UI.",
    "Adding a new slide by hand inherits the master's look.",
)
"""PHASE-3B.md GATE 3 checks 1-5, verbatim in substance. Printed, never evaluated."""


@dataclass
class FinalAssessment:
    """Everything checkable about one rendered deck, computed from disk."""

    pptx: Path
    deck_digest: str
    """`canonical_pptx_digest(pptx)` — the value `approve ... final_render` will record, shown
    so the owner can see the report and the approval name the same file."""
    post_render: PostRenderReport
    qa: list[QAFinding] = field(default_factory=list)
    grammar: list[GrammarFinding] = field(default_factory=list)
    flow: HeaderFlowReport = field(default_factory=HeaderFlowReport)
    diagram_count: int = 0
    """Native diagrams in the rendered file (shapes the diagram engine named)."""
    icon_count: int = 0
    """Theme-coloured icon shapes in the rendered file (`ICON_NAME_PREFIX`)."""

    def checkable(self) -> list[tuple[str, bool, str]]:
        """`(label, passed, detail)` for each checkable criterion, in this order:

        1. "Post-render audit passes (every claim on its slide, every number traced)" —
           `post_render.passes`; detail: first findings.
        2. "No deterministic QA findings" — `not qa`; detail: count and first finding.
        3. "No blocking grammar findings" — no `severity == "blocking"`; detail likewise.
        4. "At least one native diagram and one theme-recolourable icon" — both counts > 0;
           detail: the two counts. (Without them, human checks 1 and 2 cannot be tried.)
        Header-flow findings are advisory (flow.py) and are reported, never a criterion.
        """
        post_render_lines = _post_render_lines(self.post_render)
        blocking_grammar = [f for f in self.grammar if f.severity == "blocking"]
        return [
            (
                "Post-render audit passes (every claim on its slide, every number traced)",
                self.post_render.passes,
                _first(post_render_lines, "no findings"),
            ),
            (
                "No deterministic QA findings",
                not self.qa,
                f"{len(self.qa)} finding(s); first: {_qa_line(self.qa[0])}"
                if self.qa
                else "no findings",
            ),
            (
                "No blocking grammar findings",
                not blocking_grammar,
                f"{len(blocking_grammar)} blocking; first: {blocking_grammar[0]}"
                if blocking_grammar
                else "no blocking findings",
            ),
            (
                "At least one native diagram and one theme-recolourable icon",
                self.diagram_count > 0 and self.icon_count > 0,
                f"{self.diagram_count} diagram(s), {self.icon_count} icon(s)",
            ),
        ]

    @property
    def passes(self) -> bool:
        return all(passed for _label, passed, _detail in self.checkable())


def assess_final(
    deck: Deck,
    pptx: Path,
    *,
    tokens: DesignTokens,
    brief: DeckBrief | None,
    profile: HeaderStyleProfile,
) -> FinalAssessment:
    """Compute every checkable GATE 3 criterion for `pptx`, rendered from `deck`.

    Contract: pure reads, writes nothing. `post_render_audit(deck, pptx)`;
    `run_deterministic_qa(pptx, tokens)`; `grammar.lint_deck(deck).findings` plus
    `grammar.check_icon_adjacency(Presentation(pptx), tokens)`; `flow_report(deck, tokens,
    profile, brief=brief)`; counts read from the PPTX's shape names (reuse the prefixes the
    diagram engine and icon placer already write — find them, do not invent new ones);
    `deck_digest = canonical_pptx_digest(pptx)`. A slide-id mismatch between `deck` and
    `pptx` is not raised here — `post_render_audit` reports it as a finding.
    """
    post_render = post_render_audit(deck, pptx)
    qa = run_deterministic_qa(pptx, tokens)
    prs = Presentation(str(pptx))
    grammar_findings = [
        *grammar.lint_deck(deck).findings,
        *grammar.check_icon_adjacency(prs, tokens),
    ]
    return FinalAssessment(
        pptx=pptx,
        deck_digest=canonical_pptx_digest(pptx),
        post_render=post_render,
        qa=qa,
        grammar=grammar_findings,
        flow=flow_report(deck, tokens, profile, brief=brief),
        diagram_count=_count_diagrams(prs),
        icon_count=_count_icons(prs),
    )


def _count_icons(prs: PresentationType) -> int:
    """Icon instances in the rendered file, from the `icon:<family>:<name>` shape names.

    `place_icon` makes one shape per subpath, all sharing a name and bounding box, so shapes
    with the same slide, name and box count once. Two placements of the same glyph at
    different positions count twice — the key `grammar._icon_shape_key` uses.
    """
    instances: set[tuple[int, str, float, float, float, float]] = set()
    for slide_index, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if not shape.name.startswith(ICON_NAME_PREFIX):
                continue
            instances.add(
                (
                    slide_index,
                    shape.name,
                    round(shape.left.pt if shape.left is not None else 0.0, 1),
                    round(shape.top.pt if shape.top is not None else 0.0, 1),
                    round(shape.width.pt if shape.width is not None else 0.0, 1),
                    round(shape.height.pt if shape.height is not None else 0.0, 1),
                )
            )
    return len(instances)


def _count_diagrams(prs: PresentationType) -> int:
    """Native diagrams in the rendered file.

    SCAFFOLD CONTRADICTION: the contract says to count diagrams from shape names "the
    diagram engine already writes". It writes none — `autodeck/design/diagrams.py` draws
    plain autoshapes and connectors through `draw.add_autoshape` / `add_connector` with
    python-pptx's default names ("Chevron 3", "Rectangle 7"). Only icons carry a name hook
    (`icon:`). Counting diagrams from the file needs a prefix the engine does not yet write,
    which this module may not invent. Left for the caller to decide; see the report.
    """
    raise NotImplementedError("the diagram engine writes no shape-name prefix to count")


def render_final_report(assessment: FinalAssessment, *, audit_report: str) -> str:
    """The markdown written to `runs/<run>/final_audit_report.md` and printed by `gate3`.

    Contract, sections in order: the deck path and `deck_digest`; the checkable criteria as
    `[PASS]`/`[FAIL]` lines with details; every post-render, QA and grammar finding in full;
    the header flow (`assessment.flow.render()`); `audit_report` (the claim-level audit,
    `audit.report.render(build_audit_report(...))`, unchanged); and last, `HUMAN_CHECKS` as
    unticked `- [ ]` items under a heading that says they must be done in PowerPoint by a
    person and that nothing in this report has checked them. Deterministic.
    """
    checks = assessment.checkable()
    failed = [label for label, passed, _detail in checks if not passed]
    lines: list[str] = [
        "# GATE 3 - final deck review",
        "",
        f"Deck: `{assessment.pptx}`",
        f"Deck digest: `{assessment.deck_digest}`",
        "",
        "`autodeck approve <run> final_render` records this digest. If the deck changes "
        "after this report was written, the approval no longer counts.",
        "",
        "## Checkable criteria",
        "",
    ]
    for label, passed, detail in checks:
        lines.append(f"[{'PASS' if passed else 'FAIL'}] {label} - {detail}")
    lines.extend(
        [
            "",
            "All four checkable criteria pass."
            if not failed
            else f"{len(failed)} of {len(checks)} checkable criteria fail.",
            "",
            "## Findings",
            "",
            "### Post-render audit",
            "",
        ]
    )
    lines.extend(_bullets(_post_render_lines(assessment.post_render, advisory=True)))
    lines.extend(["", "### Deterministic QA", ""])
    lines.extend(_bullets([_qa_line(finding) for finding in assessment.qa]))
    lines.extend(["", "### Grammar", ""])
    lines.extend(_bullets([str(finding) for finding in assessment.grammar]))
    lines.extend(["", "## Header flow", "", assessment.flow.render(), ""])
    lines.extend(["## Claim-level audit", "", audit_report.rstrip(), ""])
    lines.extend(
        [
            "## Human checks - do these in PowerPoint, by a person",
            "",
            "These five checks must be done in PowerPoint by a person. Nothing in this "
            "report has checked them, and a clean run of every criterion above does not "
            "stand in for them.",
            "",
        ]
    )
    lines.extend(f"- [ ] {check}" for check in HUMAN_CHECKS)
    return "\n".join(lines).rstrip() + "\n"


def _first(lines: list[str], when_empty: str) -> str:
    if not lines:
        return when_empty
    more = f" (+{len(lines) - 1} more)" if len(lines) > 1 else ""
    return f"{lines[0]}{more}"


def _post_render_lines(report: PostRenderReport, *, advisory: bool = False) -> list[str]:
    """Post-render findings as text: file-level findings, then the numeric re-lint's.

    Criterion 1 fails on `report.passes`, which counts blocking numeric findings only, so the
    one-line detail lists those; the full section also lists the advisory ones.
    """
    lines = [
        f"[{finding.check}] slide {finding.slide_id}: {finding.detail}"
        for finding in report.findings
    ]
    numeric = report.numeric.findings if advisory else report.numeric.blocking
    lines.extend(f"numeric re-lint {finding}" for finding in numeric)
    return lines


def _qa_line(finding: QAFinding) -> str:
    return (
        f"[{finding.check}] slide {finding.slide_id} - {', '.join(finding.shapes)}: measured "
        f"{finding.measured:g} against {finding.threshold:g}; fix by {finding.remedy} "
        f"({finding.remedy_detail})"
    )


def _bullets(lines: list[str]) -> list[str]:
    return [f"- {line}" for line in lines] if lines else ["None."]
