"""Grammar lints (D13): text/icon/diagram balance, enforced deterministically.

D13's own wording is *"text/icon/diagram balance enforced by deterministic grammar lints,
not model taste."* `framing_linter.py` is the register this module matches: findings, not
decisions (an orchestrator blocks); severities named and justified per rule; and — the two
lessons that phase carries forward —

* **Closed lists fail on ordinary phrasing outside the list.** That risk is A5's, not this
  module's: nothing here classifies natural language. Every rule below counts something
  mechanical — words via `str.split()` (the same operationalisation
  `DiagramSpec._labels_fit_the_budget` already uses), block kinds, or rendered geometry.
  The analogous risk here is different and is named per rule: a mechanical count standing in
  for a judgement ("5-7 concepts", "adjacent") a human would sometimes make differently.
* **A check that compares a prediction to itself proves nothing.** Every threshold below is
  an independent number — one of D13's own figures, or a constant derived from the deck's
  own theme tokens (`Spacing.gutter`) rather than from anything this module itself computed
  — compared against an actual count taken from the IR or the rendered slide. None of these
  rules recomputes its own input as its own check.

## The one thing harder than it looks: icon adjacency

D13 says *"every icon adjacent to a text label"*. 3a.7 did not add a structural
icon-plus-label pairing to `layout_kit` (`Stack.icon_row(...)` would be the honest fix — see
its own carried finding), so an icon's `IconRef` block carries no reference to a label
block, and `place_icon` writes only a name (`icon:<family>:<name>`) onto the shapes it
creates. **Adjacency is therefore a geometric judgement made after render, not a structural
fact read off the IR** — see `check_icon_adjacency` for the threshold and what it does not
catch. This module does not build `Stack.icon_row`; that stays 3a.7's deliberate gap, to be
closed by whoever next needs icon+label pairing to be a constructible thing rather than an
inferred one.

## Two seams, not one

Word budgets, diagram count and the icon/chart/diagram pileup rule read only the IR
(`Deck`/`Slide`/`Block`) and need no render — `lint_deck` runs them before a slide is ever
drawn, the same way A2/A5 run on IR text. Icon adjacency needs real shape geometry and can
only run on a rendered `Presentation` — `check_icon_adjacency` is the second, later seam,
mirroring `numeric_linter.py`'s own `lint_deck`/`lint_rendered_slides` split for the same
reason: the two inputs are not interchangeable, so the checks that need each one are kept
apart rather than one function pretending to accept both.

Owning phase: 3a (task 3a.9), D13. Routed to Sonnet per B28: `autodeck/design/` is not a
guardrail path in `docs/MODEL_ROUTING.md`, so the path — not the phase brief's Opus tag —
sets the tier (B28's own rule, "guardrail paths keep their tier; the rest de-escalate").
"""

from __future__ import annotations

import math
from dataclasses import dataclass, field
from typing import Literal

from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide as PptxSlide

from autodeck.design.icons.consistency import ICON_NAME_PREFIX
from autodeck.design.theme.tokens import DesignTokens
from autodeck.ir.models import Block, BlockKind, CommunicationMode, Deck, Slide

Severity = Literal["blocking", "advisory"]


@dataclass(frozen=True)
class GrammarFinding:
    """One thing D13 has to say about one slide. Same shape as `gate1.Finding` and
    `FramingFinding`: this module reports, an orchestrator blocks."""

    check: str
    severity: Severity
    detail: str
    location: str = ""

    def __str__(self) -> str:
        marker = "BLOCKING" if self.severity == "blocking" else "advisory"
        where = f" · {self.location}" if self.location else ""
        return f"[{marker}] {self.check}{where}: {self.detail}"


@dataclass
class GrammarReport:
    """What D13's grammar lints found, across the IR pass, the render pass, or both."""

    findings: list[GrammarFinding] = field(default_factory=list)
    slides_checked: int = 0

    @property
    def blocking(self) -> list[GrammarFinding]:
        return [f for f in self.findings if f.severity == "blocking"]

    @property
    def advisory(self) -> list[GrammarFinding]:
        return [f for f in self.findings if f.severity == "advisory"]

    @property
    def blocks_build(self) -> bool:
        return bool(self.blocking)

    def extend(self, other: GrammarReport) -> None:
        self.findings.extend(other.findings)
        self.slides_checked += other.slides_checked

    def render(self) -> str:
        lines = [
            "D13 — grammar lint",
            f"{self.slides_checked} slide(s) checked · {len(self.blocking)} blocking · "
            f"{len(self.advisory)} advisory",
            "",
        ]
        if not self.findings:
            lines.append("Every slide holds D13's text/icon/diagram balance.")
        else:
            lines.extend(str(f) for f in self.blocking + self.advisory)
        return "\n".join(lines)


# ---------------------------------------------------------------------------
# Word budgets per communication mode
# ---------------------------------------------------------------------------

#: D13's own figures (PHASE-3A.md task 3a.9; plan §6.11.2). A slide with no
#: `communication_mode` assigned yet has no budget to check against — see
#: `check_word_budget`.
WORD_BUDGET_PER_MODE: dict[CommunicationMode, int] = {
    "icon_anchored": 50,
    "diagram_led": 60,
    "text_led": 90,
}


def _block_word_count(block: Block) -> int:
    """Words this block puts in front of a reader, across every field that carries text.

    `str.split()` is the definition of "a word" already used one level down, in
    `DiagramNode.word_count()` and `DiagramSpec._labels_fit_the_budget` — reusing it here
    rather than inventing a second tokeniser is the same "one definition, several readers"
    discipline `numeric_linter.py`'s `derivation_input_is_traceable` argues for.

    Counted: `block.text` (framing/section_header prose); `block.claim.text` (the
    assertion, not its citations); every diagram node's label
    (`DiagramNode.word_count()` — the diagram geometry's own count, not re-derived here);
    a diagram's title and structural text (axis poles, transition labels, via
    `DiagramGeometry.structural_texts()`); a chart's headline (`ChartSpec.title`); a
    figure's caption.

    **Not counted, on purpose:** a chart's category and series names, and a citation's
    quoted text. A six-category bar chart legitimately carries six or more short labels
    that read, to a viewer, as one glance at a chart rather than six more things to read —
    counting them would make every chart-bearing slide fail a "reading load" budget for
    data density, which is a different problem D13 does not name. This is the rule's own
    stated gap: a slide that pads a chart's `title` with a paragraph, or pairs a chart with
    enough OTHER text to look like a wall of words, still trips this check on that text.
    """
    words = 0
    if block.text:
        words += len(block.text.split())
    if block.claim is not None:
        words += len(block.claim.text.split())
    if block.figure is not None and block.figure.caption:
        words += len(block.figure.caption.split())
    if block.chart is not None and block.chart.title:
        words += len(block.chart.title.split())
    if block.diagram is not None:
        diagram = block.diagram
        words += sum(node.word_count() for node in diagram.nodes)
        if diagram.title:
            words += len(diagram.title.split())
        words += sum(len(site.text.split()) for site in diagram.payload.structural_texts())
    return words


def check_word_budget(slide: Slide, *, location: str = "") -> list[GrammarFinding]:
    """Blocking: catches a slide whose face carries more reading than D13's mode ceiling.

    Face blocks only (`slide.blocks`), not speaker notes — the budget is about what a
    reader scans on the slide, and a spoken note is not read off it, unlike A5's fence,
    which follows a fact wherever it is written.

    Does not catch: a slide with `communication_mode` unset. Mode assignment is Phase 3b's
    decision (task 3b.7, the art-direction pass), so a slide built before that pass has run
    has nothing to check its word count against yet — silently skipping it, rather than
    blocking on an absent field this phase does not set, is the same posture
    `check_overflow` takes toward a component the catalog has not registered.

    Does not catch: a slide within budget that is still a wall of words for its mode
    because the 50/60/90 ceilings are themselves the brief's own compromise, not a per-slide
    measurement of legibility — D13 sets the ceiling, this function only checks it.
    """
    if slide.communication_mode is None:
        return []
    budget = WORD_BUDGET_PER_MODE[slide.communication_mode]
    total = sum(_block_word_count(block) for block in slide.blocks)
    if total <= budget:
        return []
    return [
        GrammarFinding(
            check="word budget",
            severity="blocking",
            detail=(
                f"{total} word(s) on a {slide.communication_mode!r} slide, over D13's "
                f"{budget}-word ceiling for that mode (icon_anchored ≤ 50, "
                "diagram_led ≤ 60, text_led ≤ 90). Cut prose or split the slide; "
                "shrinking the type to fit more words defeats the budget rather than "
                "meeting it."
            ),
            location=location,
        )
    ]


# ---------------------------------------------------------------------------
# Diagram count and the icon/chart/diagram pileup
# ---------------------------------------------------------------------------

#: D13: "≤1 diagram/slide". A `DiagramSpec` is already a composed geometry of several
#: nodes and edges; two of them sharing a slide is two competing geometries, not one
#: richer one, and the slide-geometry skill's own load-bearing test (would a plain list
#: lose information?) is never asked twice on the same slide by construction here.
MAX_DIAGRAMS_PER_SLIDE = 1

#: D13: "no icon+chart+diagram pileup" — three distinct kinds of non-prose visual
#: furniture. Deliberately the three named in the brief and nothing wider: an icon next to
#: a chart alone, or two charts, is a layout choice this rule has no opinion about. See
#: `check_no_icon_chart_diagram_pileup` for what that narrowness gives up.
PILEUP_KINDS: frozenset[BlockKind] = frozenset({"icon", "chart", "diagram"})


def check_diagram_count(slide: Slide, *, location: str = "") -> list[GrammarFinding]:
    """Blocking: catches more than one `kind="diagram"` block on one slide's face.

    Exact and closed: a diagram block either exists on the slide or it does not, so there
    is no judgement call here to get wrong on an ordinary slide the way a word-count
    ceiling or a concept count can. Face blocks only — a diagram cannot render twice from
    one slide's speaker notes, which carry no geometry at all.

    Does not catch: two diagrams that are visually small and genuinely uncluttered, or one
    diagram that is itself overloaded (`DiagramSpec._labels_fit_the_budget` already blocks
    an over-budget label; there is no "too many nodes" check here or anywhere else in this
    module, and it would need its own calibration if the phase after this one wants it).
    """
    diagrams = [b for b in slide.blocks if b.kind == "diagram"]
    if len(diagrams) <= MAX_DIAGRAMS_PER_SLIDE:
        return []
    return [
        GrammarFinding(
            check="diagram count",
            severity="blocking",
            detail=(
                f"{len(diagrams)} diagram blocks on one slide "
                f"({', '.join(b.id for b in diagrams)}); D13 allows at most "
                f"{MAX_DIAGRAMS_PER_SLIDE}. Two geometries competing for one slide is two "
                "arguments, not one — split the slide."
            ),
            location=location,
        )
    ]


def check_no_icon_chart_diagram_pileup(
    slide: Slide, *, location: str = ""
) -> list[GrammarFinding]:
    """Blocking: catches a slide whose face carries an icon block AND a chart block AND a
    diagram block all at once — D13's named pileup, read straight off block kinds.

    Exact and closed, for the same reason `check_diagram_count` is: kind membership is a
    fact, not an estimate. It fires only on the three-way combination the brief names.

    Does not catch: two of the three together without the third (an icon-heavy slide next
    to a chart, with no diagram, can still be busy); more than one of any single kind
    contributing to the same pileup feeling; or a genuinely cluttered slide built from
    kinds outside this set entirely (several `claim` blocks, say). Widening the set to "any
    two distinct visual kinds" was considered and rejected — a headline `claim` next to a
    `chart` is most of this catalog's components working as designed, and a rule that fired
    on that would be the closed-list failure `framing_linter.py` warns about, just on
    kinds instead of words.
    """
    present = {block.kind for block in slide.blocks} & PILEUP_KINDS
    if present != PILEUP_KINDS:
        return []
    return [
        GrammarFinding(
            check="icon+chart+diagram pileup",
            severity="blocking",
            detail=(
                "This slide carries an icon block, a chart block and a diagram block "
                "together — D13's named pileup. Three different kinds of non-prose visual "
                "furniture compete for the same attention; drop one or split the slide."
            ),
            location=location,
        )
    ]


# ---------------------------------------------------------------------------
# Concept count — advisory
# ---------------------------------------------------------------------------

#: D13's own figure is a range ("5-7"), not a single number. 7, the top of that range, is
#: used as the ceiling: anything at or under 7 is unambiguously within what the brief calls
#: fine, and firing at the range's low end would flag slides D13 itself does not object to.
MAX_CONCEPTS_ADVISORY = 7

#: Slot names that are metadata rather than a "concept" a reader is being asked to hold —
#: the citation strip every component with a `source` slot carries. Kept as an explicit,
#: named exclusion (rather than inferring it from the slot's role) because a metadata slot
#: is the one place counting every block would double-count the same information already
#: sitting behind the claims it labels.
EXCLUDED_CONCEPT_SLOTS = frozenset({"source"})


def check_concept_count(slide: Slide, *, location: str = "") -> list[GrammarFinding]:
    """Advisory: flags a slide whose face carries an unusually large number of IR blocks.

    **Why advisory, not blocking.** "Concept" is a judgement about distinct ideas; this
    function counts `Block`s, which is the only unit the IR actually offers and a real but
    imperfect proxy for it — a one-word data-card label and a fully argued bullet point
    both count as one "concept" here, and a component's own structural chrome (two column
    titles, say) counts exactly like a load-bearing point does. A correctly filled
    `two_column_compare` (two titles plus three points a side) or `data_card_grid` (six
    cards) can legitimately clear 7 blocks while reading as clean, intentional layouts — the
    false-positive risk the spec asks this module to weigh against blocking. Per D13's own
    "prefer advisory where you are unsure" instruction, this stays a name-and-report rather
    than a build-stopper.

    Does not catch: a slide with few, dense blocks that is conceptually overloaded anyway
    (one `claim` block whose text argues five unrelated points), or a slide that pads its
    block count with metadata this function already excludes in some other way it did not
    anticipate.
    """
    concepts = [block for block in slide.blocks if block.slot not in EXCLUDED_CONCEPT_SLOTS]
    if len(concepts) <= MAX_CONCEPTS_ADVISORY:
        return []
    return [
        GrammarFinding(
            check="concept count",
            severity="advisory",
            detail=(
                f"{len(concepts)} content blocks on one slide, over the "
                f"{MAX_CONCEPTS_ADVISORY}-block proxy for D13's 5-7 'concepts' guidance. "
                "Block count is a rough stand-in for distinct ideas, not a measurement of "
                "them (see this function's docstring) — a human should look at whether the "
                "slide actually reads as crowded before splitting it."
            ),
            location=location,
        )
    ]


# ---------------------------------------------------------------------------
# The IR-only pass
# ---------------------------------------------------------------------------


def lint_slide_ir(slide: Slide, *, location: str = "") -> list[GrammarFinding]:
    """Every IR-only grammar check for one slide: word budget, diagram count, the
    icon/chart/diagram pileup, and the advisory concept count. No render required."""
    findings: list[GrammarFinding] = []
    findings.extend(check_word_budget(slide, location=location))
    findings.extend(check_diagram_count(slide, location=location))
    findings.extend(check_no_icon_chart_diagram_pileup(slide, location=location))
    findings.extend(check_concept_count(slide, location=location))
    return findings


def lint_deck(deck: Deck) -> GrammarReport:
    """Run every IR-only grammar check over a whole deck.

    Icon adjacency is not in this report — it needs rendered geometry `Deck` does not
    carry. See `check_icon_adjacency` and combine the two reports at the call site once a
    render exists (`GrammarReport.extend`).
    """
    report = GrammarReport()
    for slide in deck.slides:
        report.slides_checked += 1
        report.findings.extend(lint_slide_ir(slide, location=f"slide {slide.id}"))
    return report


# ---------------------------------------------------------------------------
# Icon adjacency — the rendered pass
# ---------------------------------------------------------------------------

#: How far a text shape may sit from an icon shape and still read as its label, as a
#: multiple of the deck's own `Spacing.gutter` token — never a bare point value invented
#: here (`catalog.py`'s own rule: geometry is read off something real, not guessed).
#: `gutter`, not `baseline`, is the right theme quantity to measure against: this codebase's
#: own slot builders reach for `canvas.baseline` (2-3x, ~12-18pt at the default scale) for
#: the gap *inside* one visual unit — `_callout_takeaway_slots`' label-to-takeaway gap,
#: `_quote_slots`' mark-to-quote gap — and reach for `canvas.gutter` only when splitting a
#: slide into independent regions (`_evidence_with_figure_slots` and `_before_after_slots`
#: both call `split_columns(2, canvas.gutter)`; `_two_column_compare_slots` uses
#: `canvas.gutter * 1.5`). So a text shape further from an icon than one whole gutter is, by
#: this codebase's own convention, sitting in a different region of the slide, not beside
#: it. 1.0x (a full gutter, not a fraction of one) is chosen because no registered component
#: places an icon yet (3a.7's carried finding — see this module's docstring) — there is no
#: rendered example to calibrate against the way `budget_check.py`'s
#: `RENDER_HEADROOM_FRACTION` was measured from real LibreOffice output. A full gutter is
#: generous enough that it should not reject a real pairing once one is built; the cost is a
#: genuinely mis-placed icon several gutters from any text still passing. Recalibrate this
#: against a real render the first time a component actually calls `Frame.icon`.
ICON_ADJACENCY_GUTTER_MULTIPLE = 1.0


@dataclass(frozen=True)
class _ShapeRect:
    """A shape's bounding box, in points — the only geometry adjacency needs."""

    left: float
    top: float
    width: float
    height: float


def _gap_pt(a: _ShapeRect, b: _ShapeRect) -> float:
    """The shortest distance between two axis-aligned rectangles, in points.

    0.0 when they touch or overlap. The standard AABB-to-AABB gap: the horizontal and
    vertical clearances are each `max(0, ...)` (zero when the boxes overlap on that axis),
    and the two are combined as a Euclidean distance so a box sitting diagonally off a
    corner is not treated as adjacent merely because one axis's clearance is small.
    """
    dx = max(a.left - (b.left + b.width), b.left - (a.left + a.width), 0.0)
    dy = max(a.top - (b.top + b.height), b.top - (a.top + a.height), 0.0)
    return math.hypot(dx, dy)


def _icon_shape_key(name: str, rect: _ShapeRect) -> tuple[str, float, float, float, float]:
    """De-duplication key for one icon *instance*.

    `place_icon` creates one shape per subpath, all sharing one name and — by construction
    of `place_icon`'s own `square` — the exact same bounding box (see `custgeom.py`). Without
    this, one icon with three strokes would produce three identical findings or three
    identical passes for what a reader sees as one glyph. Two icons of the same glyph placed
    twice on a slide are NOT merged — the key includes position, not just the name — because
    `icons/consistency.py`'s own `placed_icons` treats repeated placements as distinct too.
    """
    return (
        name,
        round(rect.left, 1),
        round(rect.top, 1),
        round(rect.width, 1),
        round(rect.height, 1),
    )


def _text_shapes(slide: PptxSlide) -> list[_ShapeRect]:
    """Every non-icon shape on `slide` carrying visible text, as its bounding box.

    `has_text_frame` alone is not enough: `add_autoshape` zeroes an autoshape's margins but
    leaves its (empty) text frame in place for chrome shapes like rules and panels, so an
    empty-after-strip text frame is excluded — it is not a label, it is a rectangle that
    happens to have a `TextFrame` object because python-pptx gives every autoshape one.
    """
    rects: list[_ShapeRect] = []
    for shape in slide.shapes:
        if shape.name.startswith(ICON_NAME_PREFIX):
            continue
        if not shape.has_text_frame:
            continue
        # `.text_frame` is only on the concrete shape classes that actually have one, not
        # on the `BaseShape` that iterating `slide.shapes` is typed as — the same narrowing
        # gap `icons/consistency.py` works around for `.line` on a `Shape`. `has_text_frame`
        # just confirmed this shape has one.
        text_frame = getattr(shape, "text_frame")  # noqa: B009
        if not text_frame.text.strip():
            continue
        rects.append(_ShapeRect(shape.left.pt, shape.top.pt, shape.width.pt, shape.height.pt))
    return rects


def check_icon_adjacency(prs: PresentationType, tokens: DesignTokens) -> list[GrammarFinding]:
    """Blocking: catches an icon shape with no text shape within
    `ICON_ADJACENCY_GUTTER_MULTIPLE` gutters of it, on the same slide.

    Reads every shape `place_icon` created (`ICON_NAME_PREFIX`, the same hook
    `icons/consistency.py` uses) and, for each one, the shortest gap to every other
    text-bearing shape on that slide. **The icon naming convention is a strong hook for
    finding icons — every icon shape is named and none is missed — but it says nothing
    about which text is that icon's OWN label**, because no structural pairing exists (see
    this module's docstring). So this checks the weaker, geometric question: is *some* text
    close enough to read as this icon's label. See `ICON_ADJACENCY_GUTTER_MULTIPLE` for the
    threshold and why it is generous rather than tight.

    Does not catch: an icon sitting near text that is not actually its label — a caption or
    source line that happens to be the nearest text shape on a sparse slide would satisfy
    this check even though it is not what the icon illustrates. That gap is exactly what a
    structural `Stack.icon_row(...)` pairing would close and a geometric heuristic cannot;
    see this module's docstring for why that fix is named rather than built here. Also does
    not catch a text label that is adjacent but wrong (an icon meant for "risk" sitting next
    to a "cost" label) — adjacency is a layout fact, not a semantic one.

    Location is positional (`slide {n}`, 1-based), not the IR's `Slide.id`: a rendered
    `Presentation` carries no reference back to the IR slide it came from.
    """
    threshold = tokens.spacing.gutter * ICON_ADJACENCY_GUTTER_MULTIPLE
    findings: list[GrammarFinding] = []
    for slide_index, slide in enumerate(prs.slides):
        texts = _text_shapes(slide)
        seen: set[tuple[str, float, float, float, float]] = set()
        for shape in slide.shapes:
            if not shape.name.startswith(ICON_NAME_PREFIX):
                continue
            rect = _ShapeRect(shape.left.pt, shape.top.pt, shape.width.pt, shape.height.pt)
            key = _icon_shape_key(shape.name, rect)
            if key in seen:
                continue
            seen.add(key)

            if texts and min(_gap_pt(rect, text) for text in texts) <= threshold:
                continue
            findings.append(
                GrammarFinding(
                    check="icon adjacent to text label",
                    severity="blocking",
                    detail=(
                        f"{shape.name!r} has no text shape within {threshold:.1f}pt "
                        f"({ICON_ADJACENCY_GUTTER_MULTIPLE:g}x the deck's gutter) of it. "
                        "D13 requires every icon to sit next to a text label; an icon this "
                        "far from any text reads as decoration with nothing explaining it."
                    ),
                    location=f"slide {slide_index + 1} · {shape.name}",
                )
            )
    return findings
