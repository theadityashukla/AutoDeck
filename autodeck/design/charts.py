"""Native editable charts from `ChartSpec` (D10, task 3a.5).

## The invariant

A `ChartSpec` is what `prompts/content.md` calls it: *"a dense set of factual
assertions"* — a ten-point series is ten claims, not one illustration. Two things follow,
and this module is built around both.

**No image fallback exists.** `place_chart` either builds the chart type it is asked for as
a native OOXML chart part, or it raises `ChartConstructionError` — never a `.png`. D10's
whole point is that a pasted picture is the #1 client edit request, and a chart that
degrades to a picture the moment its data is awkward would reintroduce exactly that on a
schedule nobody controls.

**The source table cell is the citation.** `ChartSpec.source_citations` already carries
`Citation` objects resolved against the document store (`DocumentStore.resolve_cell`,
task 1.4) — resolving one is not this module's job, and it does not attempt it. What this
module owns is turning citations that already exist into something a reader in the room can
act on: `_source_line` formats them into the caption band under the chart, in the same
`doc_id p.N` shape `autodeck.audit.numeric_linter._citation_label` and
`autodeck.audit.report._ref` already use for the same purpose elsewhere. Three modules
agreeing on that shape by convention, not by importing one function across a design/audit
boundary that does not otherwise exist, is a deliberate, cheap kind of consistency — not the
sort of "resolution" the task warns against duplicating.

## Why this module reads `ChartSpec` directly

Every other registered component (`catalog.py`) takes its own small content dataclass
(`QuoteContent`, `BigNumberContent`, ...), decoupled from `autodeck.ir`; a future assembler
maps `Block` to those. Charts are different for the same reason diagrams are (task 3a.6):
`ChartSpec` already *is* the render-ready content — categories, series, axis labels and
citations, nothing left to adapt — and the phase brief calls this file out by name as
building "from `ChartSpec`". Introducing a parallel `ChartContent` shape here would be a
second description of the same nine fields, with nothing to keep the two in step.

## Editable, not baked — how, mechanically

`python-pptx`'s `add_chart` does not draw a picture of the numbers: it writes a distinct
`c:chartSpace` XML part whose `c:val`/`c:cat` elements carry both a cached value (`c:v`, for
software that never opens the workbook) *and* a formula reference (`c:f`, e.g.
`Sheet1!$B$2:$B$5`) into a **second part it also writes** — a real `.xlsx` workbook embedded
in the package (`ppt/embeddings/*.xlsx`), populated with exactly `spec.categories` and
`spec.series[*].values`. That second part is what "Edit Data" opens. `tests/test_charts.py`
checks both halves directly: it unzips the produced `.pptx`, asserts the chart part is a
`c:chartSpace` referenced by formula rather than a `p:pic`, and opens the embedded workbook
with `openpyxl` to confirm its cells hold the real numbers rather than a rendering of them.

Per B31, no PowerPoint exists in this environment to confirm PowerPoint's own "Edit Data in
Excel" dialog against this file — that is owner debt, stated precisely in this task's
handover rather than assumed. What *is* checked here, locally, is what LibreOffice does with
the same file: `tests/test_charts.py`'s `render`-marked tests run the produced deck through
`autodeck.render.qa.libreoffice.render_pptx` (confirming LibreOffice opens and rasterises it
without falling back to nothing) and through a headless round trip (`--convert-to pptx`)
that re-saves the file and re-checks what survives. The chart does: the round-tripped file
is still a `c:chartSpace` with the same category and value caches, never flattened to a
picture. The embedded workbook does not — LibreOffice edits chart data through its own
in-XML "Data Table" rather than through the external `.xlsx` PowerPoint's "Edit Data in
Excel" opens, and its save path reflects that. That asymmetry is itself evidence worth
having: it is a concrete, observed instance of the exact PowerPoint/LibreOffice OOXML
disagreement B31 already names as owner debt, not a defect in what this module writes (see
`TestRenderedByLibreOffice`'s round-trip test for the detail).

## Budgets — the one rule this kit will not break, applied to a medium it does not draw

`layout_kit`'s governing rule is that nothing here shrinks, truncates or reflows text to
make it fit; overflow is an error raised before render, not a smaller font discovered after
it. A native chart's title, axis titles, category labels and legend entries are drawn by
whichever application opens the file, not by this codebase's own text engine — there is no
`add_text` call to refuse. `_check_labels_fit` is the render-time safety net for that gap: it
measures every piece of chart text this module sets against the real font metrics
`budgets.py` already uses everywhere else, using the exact same `TextStyle` this module then
hands to python-pptx, and raises `LayoutOverflowError` — the same exception every other
renderer raises, for the same reason — the moment one would not read as a single line at the
size and width the chart is actually given. Like `catalog.py`'s slot budgets, these are
conservative upper-bound estimates of the space a native chart's own layout engine will
actually grant a given label (it also reserves room for tick marks, gridlines and a legend,
which this module cannot see in advance) — overflow becomes rare by construction here, not
impossible by construction, and the golden-preview render is still what catches the rest.

## Where a `ChartKind` resists native construction

Six of the seven `ChartKind` values map onto a `CategoryChartData`-backed native chart with
no friction: the IR's (categories, series-of-values) shape *is* a category chart's shape.
`scatter` is the exception. An XY scatter chart has no category axis at all — python-pptx
models it with `XyChartData`, where every point is a real `(x, y)` pair — so
`ChartSpec.categories` can only become the x-axis if every category actually parses as a
number. When it does not, this is not "hard to build", it is a genuine mismatch between what
the IR calls a scatter chart's independent variable (a label) and what a scatter chart's
independent variable actually is (a number), and `place_chart` raises
`ChartConstructionError` naming it rather than guessing — e.g. by silently using the
category's *position* as x, which would draw a chart whose x-axis lies about what it
measures.

Owning phase: 3a (task 3a.5).
"""

from __future__ import annotations

from collections.abc import Sequence

from pptx.chart.chart import Chart
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.shapes.graphfrm import GraphicFrame
from pptx.text.text import Font
from pptx.util import Pt

from autodeck.design.draw import theme_color
from autodeck.design.layout_kit import (
    CAPTION_STRIP_FACTOR,
    Box,
    Canvas,
    Frame,
    LayoutOverflowError,
    TextStyle,
)
from autodeck.ir.models import ChartKind, ChartSpec, Citation

#: `ChartKind` -> the native `python-pptx` chart type it builds as.
#:
#: `line` picks `LINE_MARKERS` rather than bare `LINE`: A1 treats a line series as a
#: sequence of distinct point-claims, and a marker at each one is what makes every point
#: individually findable rather than only the trend between them. `scatter` picks the
#: unconnected `XY_SCATTER` — a scatter chart asserts a set of points, not a path between
#: them, which `XY_SCATTER_LINES` would draw whether or not the IR meant to imply one.
_CHART_TYPE: dict[ChartKind, XL_CHART_TYPE] = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
    "area": XL_CHART_TYPE.AREA,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
}

#: Cycled across series (or, for `pie`, across points) so a chart with more series than
#: accents repeats rather than raising — a client's palette is not this module's business
#: to second-guess, and a seventh series sharing accent1's colour with the first is a much
#: smaller problem than a chart that refuses to render.
_SERIES_COLORS: tuple[str, ...] = (
    "accent1",
    "accent2",
    "accent3",
    "accent4",
    "accent5",
    "accent6",
)

#: Chart types with no category axis at all (D10's pie has one series, sliced by category,
#: not an X/Y pair) — `spec.x_axis_label`/`spec.y_axis_label` are silently not drawn for
#: these rather than rejected, the same way an irrelevant optional field is elsewhere: a pie
#: chart with an axis label set is a harmless, plausible upstream mistake, not a data-loss
#: risk the way a dropped citation would be.
_NO_AXES: frozenset[ChartKind] = frozenset({"pie"})


class ChartConstructionError(ValueError):
    """`spec` cannot be built as a native chart of its declared `chart_type`.

    Raised in place of ever falling back to a picture (D10): a chart this module cannot
    build natively from what it was given is an escalation for a human — fix the content,
    or extend this module — never a `.png` standing in for the real thing.
    """


def place_chart(frame: Frame, box: Box, spec: ChartSpec) -> GraphicFrame:
    """Draw `spec` as a native chart inside `box`, with its citations as the source line.

    `box` is split the same way `Canvas.body_and_caption` splits the slide content area —
    chart above, one caption band below — so a chart dropped into a component's own content
    region reads exactly like every other component's citation strip. `ChartSpec` always
    has at least one citation (`source_citations` has `min_length=1`), so the caption is
    always drawn; there is no empty-source case to special-case, unlike a plain `str`
    `source` field elsewhere in the catalog.

    Raises:
        ChartConstructionError: `spec` cannot be built as its declared `chart_type`.
        LayoutOverflowError: the chart's title, an axis title, a category label or a series
            name will not read as a single line at the size and width this chart is given.
    """
    canvas = frame.canvas
    body, caption_area = box.split_bottom(
        canvas.size("caption") * CAPTION_STRIP_FACTOR, gutter=canvas.gutter
    )
    _check_labels_fit(canvas, body, spec)

    chart_data = _chart_data(spec)
    left, top, width, height = body.as_emu()
    # python-pptx's own stub annotates `add_chart` as returning `Chart` and its
    # `chart_data` parameter as the narrower `CategoryChartData`; its docstring says (and
    # runtime confirms) it actually returns the containing `GraphicFrame`, and it accepts
    # any `_BaseChartData` subclass including `XyChartData` — both are stub gaps, not real
    # type errors.
    graphic_frame: GraphicFrame = frame.slide.shapes.add_chart(
        _CHART_TYPE[spec.chart_type],
        left,
        top,
        width,
        height,
        chart_data,  # type: ignore[arg-type]
    )
    _style_chart(graphic_frame.chart, canvas, spec)

    frame.caption(caption_area, _source_line(spec.source_citations))
    return graphic_frame


# ---------------------------------------------------------------------------
# Chart data — the editable half
# ---------------------------------------------------------------------------


def _chart_data(spec: ChartSpec) -> CategoryChartData | XyChartData:
    if spec.chart_type == "scatter":
        return _xy_chart_data(spec)
    data = CategoryChartData()
    data.categories = spec.categories
    for series in spec.series:
        data.add_series(series.name, series.values)
    return data


def _xy_chart_data(spec: ChartSpec) -> XyChartData:
    try:
        x_values = [float(category) for category in spec.categories]
    except ValueError as exc:
        raise ChartConstructionError(
            "a scatter chart's categories are its x-axis values and must all parse as "
            f"numbers (D10: no chart type falls back to a picture instead); got "
            f"{spec.categories!r}. Give numeric x-values, or use a category chart type "
            "('bar', 'column', 'line', 'area', 'stacked_bar') instead."
        ) from exc

    data = XyChartData()
    for series in spec.series:
        series_data = data.add_series(series.name)
        for x, y in zip(x_values, series.values, strict=True):
            series_data.add_data_point(x, y)
    return data


# ---------------------------------------------------------------------------
# Styling — theme colours and the deck's own type, never a chart-library default
# ---------------------------------------------------------------------------


def _style_chart(chart: Chart, canvas: Canvas, spec: ChartSpec) -> None:
    _apply_font(chart.font, _label_style(canvas))

    if spec.title:
        chart.has_title = True
        # `TextFrame.text`'s setter rebuilds the paragraph from scratch, discarding any
        # `defRPr` already on it — so the font has to be applied *after* the text, not
        # before, or it is silently lost (checked directly against the written XML).
        chart.chart_title.text_frame.text = spec.title
        _apply_font(chart.chart_title.text_frame.paragraphs[0].font, _title_style(canvas))
    else:
        chart.has_title = False

    if spec.chart_type not in _NO_AXES:
        if spec.x_axis_label:
            axis = chart.category_axis
            axis.has_title = True
            axis.axis_title.text_frame.text = spec.x_axis_label
            _apply_font(axis.axis_title.text_frame.paragraphs[0].font, _label_style(canvas))
        if spec.y_axis_label:
            axis = chart.value_axis
            axis.has_title = True
            axis.axis_title.text_frame.text = spec.y_axis_label
            _apply_font(axis.axis_title.text_frame.paragraphs[0].font, _label_style(canvas))

    chart.has_legend = len(spec.series) > 1
    legend = chart.legend
    if legend is not None:
        legend.position = XL_LEGEND_POSITION.BOTTOM
        legend.include_in_layout = False
        _apply_font(legend.font, _label_style(canvas))

    _color_series(chart, spec)
    _apply_data_labels(chart, canvas)


def _color_series(chart: Chart, spec: ChartSpec) -> None:
    """Theme-colour every series (D1: a recoloured palette must move the chart with it).

    `scatter` gets a marker colour only, deliberately never a line colour: `XY_SCATTER` was
    chosen in `_CHART_TYPE` precisely because it draws unconnected points, and giving its
    series a coloured line would draw exactly the connecting path that choice was made to
    avoid — `chart.chart_type` would then read back as a line-scatter variant instead of
    the plain one this module asked for, which is how this was actually caught.
    """
    if spec.chart_type == "pie":
        # A pie has one series sliced into points; the colour that varies is per-point.
        series = chart.series[0]
        for index, point in enumerate(series.points):
            point.format.fill.solid()
            point.format.fill.fore_color.theme_color = theme_color(_accent(index))
        return

    for index, series in enumerate(chart.series):
        accent = theme_color(_accent(index))
        if spec.chart_type == "line":
            series.format.line.color.theme_color = accent
            marker = getattr(series, "marker", None)
            if marker is not None:
                marker.format.fill.solid()
                marker.format.fill.fore_color.theme_color = accent
        elif spec.chart_type == "scatter":
            marker = getattr(series, "marker", None)
            if marker is not None:
                marker.format.fill.solid()
                marker.format.fill.fore_color.theme_color = accent
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.theme_color = accent


def _accent(index: int) -> str:
    return _SERIES_COLORS[index % len(_SERIES_COLORS)]


def _apply_data_labels(chart: Chart, canvas: Canvas) -> None:
    """Show every value on the chart's face — A1's "dense set of assertions" made visible.

    `number_format="General"` deliberately does not round or rescale: a label that showed
    `63` for a cited `62.9` would be a numeral the corpus never asserted, and the numeric
    linter's second pass (over text extracted from the rendered deck, per
    `prompts/content.md`) exists precisely to catch that kind of drift.

    Skipped for `scatter`: OOXML's own schema allows `<c:dLbls>` under `<c:scatterChart>`
    (`references/ECMA-376`), but `python-pptx`'s `CT_ScatterChart` does not model that child
    — `Plot.has_data_labels` raises `AttributeError` there (checked directly against
    `pptx/oxml/chart/plot.py`). A scatter point's value is still fully present — in the
    embedded workbook and in `c:xVal`/`c:yVal`'s own cached values — just not also printed
    next to the marker, which is the one cosmetic gap this dependency leaves in an otherwise
    complete implementation.
    """
    if chart.chart_type == XL_CHART_TYPE.XY_SCATTER:
        return

    # `_Plots.__getitem__` is stubbed to return `list[Unknown] | Unknown` rather than the
    # `Plot` union its own docs promise (`plot = plots[i]`) — a stub gap, not a real
    # ambiguity: a chart this module builds always has exactly one plot.
    plot = chart.plots[0]
    plot.has_data_labels = True  # type: ignore[union-attr]
    data_labels = plot.data_labels  # type: ignore[union-attr]
    data_labels.number_format = "General"
    data_labels.number_format_is_linked = False
    data_labels.show_value = True
    if chart.chart_type == XL_CHART_TYPE.PIE:
        data_labels.show_category_name = True
    _apply_font(data_labels.font, _label_style(canvas))


def _apply_font(font: Font, style: TextStyle) -> None:
    """Write `style` onto a python-pptx `Font` — the chart-XML equivalent of `draw.add_text`.

    Sharing `TextStyle` with `_check_labels_fit`'s measurements is the point: the kit's
    central rule is that what is measured and what is drawn must be the same object, and a
    chart's text is no exception just because it is drawn by a different application.
    """
    font.name = style.family
    font.size = Pt(style.size)
    font.bold = style.bold
    font.italic = style.italic
    font.color.theme_color = theme_color(style.color)


def _title_style(canvas: Canvas) -> TextStyle:
    return canvas.style("heading", face="major", bold=True, color="dk1")


def _label_style(canvas: Canvas) -> TextStyle:
    return canvas.style("caption", face="minor", color="dk2")


# ---------------------------------------------------------------------------
# Budgets — measured against the exact styles `_style_chart` draws with
# ---------------------------------------------------------------------------


def _check_labels_fit(canvas: Canvas, body: Box, spec: ChartSpec) -> None:
    title_style = _title_style(canvas)
    label_style = _label_style(canvas)

    if spec.title:
        _require_one_line(canvas, spec.title, body.width, title_style, what="title")

    if spec.chart_type not in _NO_AXES:
        if spec.x_axis_label:
            _require_one_line(
                canvas, spec.x_axis_label, body.width, label_style, what="x-axis title"
            )
        if spec.y_axis_label:
            # Drawn rotated alongside the value axis, so its run is `body.height` long,
            # not `body.width` — the one measurement in this module that is not simply the
            # box's own width.
            _require_one_line(
                canvas, spec.y_axis_label, body.height, label_style, what="y-axis title"
            )

    # Conservative per-item shares of the plot area: a real chart also reserves room for
    # tick marks, gridlines and (for the legend) swatches, none of which this module can
    # see in advance, so these estimates are deliberately generous rather than exact — see
    # the module docstring's note on "rare by construction, not impossible by construction".
    category_width = body.width / max(len(spec.categories), 1)
    for category in spec.categories:
        _require_one_line(canvas, category, category_width, label_style, what="category label")

    if len(spec.series) > 1:
        series_width = body.width / len(spec.series)
        for series in spec.series:
            _require_one_line(
                canvas, series.name, series_width, label_style, what="series name (legend)"
            )


def _require_one_line(
    canvas: Canvas, text: str, width_pt: float, style: TextStyle, *, what: str
) -> None:
    block = canvas.measure_block(text, width_pt, style)
    if block.line_count > 1 or block.too_wide_words:
        raise LayoutOverflowError(
            f"chart {what} {text!r} does not read as one line at {width_pt:.0f}pt wide, "
            f"{style.size:g}pt {style.family}. A native chart's own text is drawn by "
            "whichever application opens the file, not by this kit, so nothing here can "
            "predict how it would wrap — shorten the text, or give the chart a wider slot "
            "(§6.9: nothing shrinks text to fit)."
        )


# ---------------------------------------------------------------------------
# Citations — formatted for display, never resolved or composed here
# ---------------------------------------------------------------------------


def _citation_label(citation: Citation) -> str:
    """`doc_id p.N` — the same shape `numeric_linter._citation_label` and `report._ref` use.

    Not a shared import: design has no dependency on audit (and should not gain one for a
    two-token string), so this is the third module to write this exact convention rather
    than the second. What matters for A1 is that it is a *display* of an already-resolved
    `Citation`, never a new one built from a doc_id/page pair handed to it separately.
    """
    return f"{citation.doc_id} p.{citation.page}"


def _source_line(citations: Sequence[Citation]) -> str:
    """Format as "Source: doc1 p.2; doc2 p.5" — every distinct span the chart's data traces to.

    Deduplicated but not otherwise summarised: a ten-point series with three distinct
    sources names all three, because "the source table cell is the citation" only holds if
    every cell's actual source stays visible, not just the first one found.
    """
    labels = dict.fromkeys(_citation_label(citation) for citation in citations)
    return "Source: " + "; ".join(labels)
