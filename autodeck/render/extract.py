"""Read a rendered PPTX back into text, by IR slide id (Phase 3b, task 3b.8's input).

The post-render audit is only as good as this extraction: text it does not read is text the
audit cannot check. So it reads **everything a reader could see or a chart could plot**:
every text frame on the slide (including table cells and grouped shapes, even though the
renderer groups nothing), every chart's title, axis titles, category labels, series names
**and cached values** from the chart XML (`c:numCache`/`c:strCache`), and the notes page.
Chart values matter most: a renderer that altered a plotted number would otherwise leave no
trace in any text frame.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from autodeck.render.renderer import SLIDE_TAG_PREFIX


class ExtractionError(RuntimeError):
    """The file cannot be attributed back to IR slides."""


@dataclass(frozen=True)
class SlideText:
    face: str
    """Every visible text run and chart string/value on the slide, newline-joined."""
    notes: str
    """The notes page's text."""


def extract_slide_text(pptx: Path) -> dict[str, SlideText]:
    """IR slide id → its rendered text.

    Contract: slide ids come from `cSld/@name` with `renderer.SLIDE_TAG_PREFIX` stripped. A
    slide with no tag, or two slides with the same tag, raises `ExtractionError` — text that
    cannot be attributed cannot be audited, and guessing by position is how a moved slide's
    numbers get checked against the wrong citations. Chart values are formatted with `repr`
    of the float, the same way `charts.py` writes them, so the numeric linter sees what it
    would see in the IR. Order within `face` is the shape tree order.
    """
    presentation = Presentation(str(pptx))
    result: dict[str, SlideText] = {}

    for pptx_slide in presentation.slides:
        raw_name = pptx_slide.name
        if not raw_name or not raw_name.startswith(SLIDE_TAG_PREFIX):
            raise ExtractionError(
                f"slide carries no autodeck tag (cSld/@name={raw_name!r}); text that "
                "cannot be attributed to an IR slide cannot be audited"
            )
        slide_id = raw_name[len(SLIDE_TAG_PREFIX) :]
        if slide_id in result:
            raise ExtractionError(
                f"slide id {slide_id!r} is tagged on more than one PPTX slide; a moved or "
                "duplicated slide's text cannot be attributed to the right one"
            )

        lines: list[str] = []
        for shape in _walk_shapes(pptx_slide.shapes):
            lines.extend(_shape_text(shape))
        face = "\n".join(lines)
        notes = _notes_text(pptx_slide)
        result[slide_id] = SlideText(face=face, notes=notes)

    return result


def _walk_shapes(shapes: Any) -> list[Any]:
    """Every shape in `shapes`, depth-first, descending into groups.

    The renderer groups nothing today, but a group added later must not become a blind spot
    for this extraction — see the module docstring.
    """
    flat: list[Any] = []
    for shape in shapes:
        flat.append(shape)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            flat.extend(_walk_shapes(shape.shapes))
    return flat


def _shape_text(shape: Any) -> list[str]:
    if getattr(shape, "has_chart", False):
        return _chart_text(shape.chart)
    if getattr(shape, "has_table", False):
        lines: list[str] = []
        for row in shape.table.rows:
            for cell in row.cells:
                lines.extend(_text_frame_lines(cell.text_frame))
        return lines
    if getattr(shape, "has_text_frame", False):
        return _text_frame_lines(shape.text_frame)
    return []


def _text_frame_lines(text_frame: Any) -> list[str]:
    return [paragraph.text for paragraph in text_frame.paragraphs if paragraph.text]


def _chart_text(chart: Any) -> list[str]:
    lines: list[str] = []

    if chart.has_title:
        lines.append(chart.chart_title.text_frame.text)

    for axis_name in ("category_axis", "value_axis"):
        axis = getattr(chart, axis_name, None)
        if axis is None:
            continue
        try:
            has_title = axis.has_title
        except (ValueError, AttributeError):
            continue
        if has_title:
            lines.append(axis.axis_title.text_frame.text)

    try:
        categories = chart.plots[0].categories
    except (ValueError, AttributeError, IndexError):
        categories = ()
    lines.extend(str(category) for category in categories)

    for series in chart.series:
        if series.name:
            lines.append(series.name)
        lines.extend(repr(float(value)) for value in series.values)

    return lines


def _notes_text(pptx_slide: Any) -> str:
    if not pptx_slide.has_notes_slide:
        return ""
    return "\n".join(_text_frame_lines(pptx_slide.notes_slide.notes_text_frame))
