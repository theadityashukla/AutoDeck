"""Deterministic render QA — pure arithmetic over a rendered PPTX (Phase 3b, task 3b.2).

Implemented per task 3b.2. Tests: `tests/test_deterministic_qa.py`.

**No model is involved, and none may be.** Every check here is arithmetic on shape
geometry, run sizes and theme colours read from the file. A model judging "does this
overlap" would be the one part of QA that could not be reproduced.

**Every failure must be fixable by a token or a slot adjustment** (PHASE-3B). So each finding
carries a `remedy` naming which, and a finding that neither can fix is typed
`"catalog_gap"`: it goes back to the Phase 3a catalog as a component design gap rather than
being special-cased in the renderer (PHASE-3B's third escalation trigger). The type makes
that routing impossible to skip.

**What this deliberately does not check: text overflowing its box.** That is the budget
engine's job, before render (§6.7). The brief is explicit — if deterministic QA starts
catching overflow routinely, the budgets are wrong and must be fixed upstream, not netted
here. A net that catches it quietly would hide the budget defect.

The four checks:

1. **Overlap** — two *text-bearing* shapes whose bounding boxes intersect by more than
   `OVERLAP_TOLERANCE_PT` in both dimensions. Text over a filled non-text shape (a panel, a
   chevron) is intended and is the contrast check's business, not this one's.
2. **Safe area** — any shape extending outside `Canvas(tokens).safe`.
3. **Minimum size** — any text run whose size is below `tokens.typography.minimum`. A run
   with no explicit size inherits from the theme and is resolved through it, not skipped.
4. **Contrast** — WCAG 2.x contrast ratio between each text run's colour and the fill of the
   top-most filled shape behind it (else the slide background). Threshold `CONTRAST_NORMAL`,
   or `CONTRAST_LARGE` for text ≥ `LARGE_TEXT_PT`, or ≥ `LARGE_BOLD_TEXT_PT` if bold.
   `schemeClr` resolves through the tokens palette; `sysClr` through its `lastClr` — note
   PHASE-3A §6.5: LibreOffice renders `sysClr` literally, so the rendered PNG and this check
   may disagree for `dk1`/`lt1` text, and the check follows the file.

**Decoration is exempt from check 4, not from the other three.** WCAG 1.4.3 does not apply
to text that is pure decoration, and a renderer declares that exemption in the file, by
naming the shape `"decor:<...>"` (`quote`'s oversized opening mark is the first user of
this) — deterministic QA does not guess "does this look decorative" from a shape's size or
style, which would be exactly the kind of judgement call §6.9 keeps out of arithmetic. The
escape hatch this would otherwise open — rename any low-contrast text box `"decor:..."` —
is closed by `_decor_exempt`'s other half: a `decor:` shape whose text contains a letter or
a digit is *itself* an "insufficient contrast" finding, routed `catalog_gap` (a naming
mistake or a real caption dressed up as decoration is a component fix, never a token one).
Overlap and safe area still apply to a `decor:` shape exactly as to any other.
"""

from __future__ import annotations

import colorsys
from collections.abc import Mapping
from dataclasses import dataclass
from itertools import combinations
from pathlib import Path
from typing import Final, Literal

from lxml.etree import QName, _Element
from pptx import Presentation
from pptx.dml.color import ColorFormat
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.oxml.simpletypes import ST_Percentage
from pptx.shapes.base import BaseShape
from pptx.slide import Slide
from pptx.text.text import TextFrame, _Paragraph, _Run  # pyright: ignore[reportPrivateUsage]

from autodeck.design.components import catalog
from autodeck.design.layout_kit import Box, Canvas
from autodeck.design.theme.tokens import DesignTokens, emu_to_points
from autodeck.render.renderer import SLIDE_TAG_PREFIX

OVERLAP_TOLERANCE_PT: Final = 1.0
"""Touching edges and sub-point rounding are not overlap."""

CONTRAST_NORMAL: Final = 4.5
CONTRAST_LARGE: Final = 3.0
LARGE_TEXT_PT: Final = 18.0
LARGE_BOLD_TEXT_PT: Final = 14.0
"""WCAG 2.x AA thresholds and its definition of large text. Named, sourced, not tuned."""

QACheck = Literal["overlap", "outside safe area", "below minimum size", "insufficient contrast"]
RemedyKind = Literal["token", "slot", "catalog_gap"]


@dataclass(frozen=True)
class QAFinding:
    check: QACheck
    slide_id: str
    """From the renderer's `autodeck:<id>` tag."""
    shapes: tuple[str, ...]
    """Shape names involved — one, or two for an overlap."""
    measured: float
    threshold: float
    remedy: RemedyKind
    remedy_detail: str
    """What to change: a token name ("typography.minimum", "palette.accent3") or a slot
    ("quote.attribution"), or for a catalog gap, which component and why no token or slot
    fixes it."""


#: A shape named with this prefix declares itself pure decoration (WCAG 1.4.3) — exempt
#: from check 4, not from the other three. See the module docstring.
DECOR_PREFIX: Final = "decor:"


def run_deterministic_qa(
    pptx: Path, tokens: DesignTokens, *, components: Mapping[str, str] | None = None
) -> list[QAFinding]:
    """All four checks over every tagged slide in `pptx`, in slide then shape order.

    Contract: pure — reads the file, writes nothing, no model, no LibreOffice. Deterministic
    ordering. Remedy assignment: contrast → `token` (the colour); minimum size → `token`
    (`typography.minimum` or the role's size); outside safe area and overlap → `slot` when
    the shapes belong to a component slot the catalog declares, else `catalog_gap`.
    Untagged slides raise `ValueError` — findings must be attributable.

    `components` (optional): slide id → the component name that actually drew it. This
    function otherwise has no `Deck`/IR and so no way to know which one — without it, a
    "slot" remedy can say only that *some* registered component declares a slot at that
    geometry, never which, so `remedy_detail` reads `"slot (component unknown)"` rather
    than naming one that may only coincide by chance. Given the map, a slide's slot/
    catalog_gap decision is attributed against that component's own declared slots only.
    """
    presentation = Presentation(str(pptx))
    canvas = Canvas(tokens)
    any_catalog_boxes = _catalog_slot_boxes(tokens)
    component_boxes: dict[str, list[tuple[str, Box]]] = {}

    findings: list[QAFinding] = []
    for slide in presentation.slides:
        slide_id = _slide_id(slide)
        shapes = list(slide.shapes)
        boxes = [_shape_box(shape) for shape in shapes]
        text_indices = [index for index, shape in enumerate(shapes) if _is_text_bearing(shape)]

        component = components.get(slide_id) if components is not None else None
        if component is None:
            slot_candidates = any_catalog_boxes
        else:
            if component not in component_boxes:
                component_boxes[component] = _catalog_slot_boxes_for(component, tokens)
            slot_candidates = component_boxes[component]
        name_slot = component is not None

        # 1. Overlap — text-bearing shapes against each other (module docstring: a filled
        # non-text shape under text is intended and is check 4's business), and against any
        # connector/line shape whose *visual* footprint crosses it — a two_by_two axis or a
        # process_flow transition arrow is exactly as unreadable struck through a label as
        # another label would be, and the module docstring's four-check list never said this
        # check only ever compares two text shapes, only that it compares *text-bearing*
        # shapes against something. `_visual_box` is what makes this detectable at all: a
        # straight connector's python-pptx bounding box is degenerate (zero-width or
        # zero-height) for exactly the horizontal/vertical lines this codebase draws, so the
        # plain `_intersection` this loop already uses would never see it.
        for i, j in combinations(text_indices, 2):
            dx, dy = _intersection(boxes[i], boxes[j])
            if dx > OVERLAP_TOLERANCE_PT and dy > OVERLAP_TOLERANCE_PT:
                remedy, detail = _slot_remedy(
                    (boxes[i], boxes[j]),
                    slot_candidates,
                    name_slot=name_slot,
                    names=(shapes[i].name, shapes[j].name),
                )
                findings.append(
                    QAFinding(
                        check="overlap",
                        slide_id=slide_id,
                        shapes=(shapes[i].name, shapes[j].name),
                        measured=min(dx, dy),
                        threshold=OVERLAP_TOLERANCE_PT,
                        remedy=remedy,
                        remedy_detail=detail,
                    )
                )

        line_indices = [
            index
            for index, shape in enumerate(shapes)
            if shape.shape_type == MSO_SHAPE_TYPE.LINE
        ]
        visual_line_boxes = {
            index: _visual_box(shapes[index], boxes[index]) for index in line_indices
        }
        for text_index in text_indices:
            for line_index in line_indices:
                dx, dy = _intersection(boxes[text_index], visual_line_boxes[line_index])
                if dx > OVERLAP_TOLERANCE_PT and dy > OVERLAP_TOLERANCE_PT:
                    # A text/line collision is always a positioning fix (the line is drawn
                    # by the same geometry that placed the label), never a new catalog slot
                    # — "slot" regardless of whether a declared slot's box happens to cover
                    # this spot; `_slot_remedy` still names one when it can, for the detail.
                    _, detail = _slot_remedy(
                        (boxes[text_index],),
                        slot_candidates,
                        name_slot=name_slot,
                        names=(shapes[text_index].name, shapes[line_index].name),
                    )
                    findings.append(
                        QAFinding(
                            check="overlap",
                            slide_id=slide_id,
                            shapes=(shapes[text_index].name, shapes[line_index].name),
                            measured=min(dx, dy),
                            threshold=OVERLAP_TOLERANCE_PT,
                            remedy="slot",
                            remedy_detail=detail,
                        )
                    )

        # 2. Safe area — every shape, text-bearing or not.
        for shape, box in zip(shapes, boxes, strict=True):
            overshoot = _safe_area_overshoot(box, canvas.safe)
            if overshoot is not None:
                remedy, detail = _slot_remedy(
                    (box,), slot_candidates, name_slot=name_slot, names=(shape.name,)
                )
                findings.append(
                    QAFinding(
                        check="outside safe area",
                        slide_id=slide_id,
                        shapes=(shape.name,),
                        measured=overshoot,
                        threshold=0.0,
                        remedy=remedy,
                        remedy_detail=detail,
                    )
                )

        # 3 & 4: minimum size and contrast, run by run, within each text-bearing shape.
        for index in text_indices:
            shape = shapes[index]
            is_decor = shape.name.startswith(DECOR_PREFIX)
            background_hex = _background_for(index, shapes, boxes, tokens, slide)
            for paragraph in _text_frame(shape).paragraphs:
                for run in paragraph.runs:
                    if not run.text:
                        continue

                    size_pt = _resolved_run_size(run, paragraph, tokens)
                    if size_pt < tokens.typography.minimum:
                        findings.append(
                            QAFinding(
                                check="below minimum size",
                                slide_id=slide_id,
                                shapes=(shape.name,),
                                measured=size_pt,
                                threshold=tokens.typography.minimum,
                                remedy="token",
                                remedy_detail="typography.minimum",
                            )
                        )

                    threshold = _contrast_threshold(size_pt, bool(run.font.bold))

                    if is_decor:
                        if not any(character.isalnum() for character in run.text):
                            continue  # pure decoration — check 4 does not apply to it
                        # The escape hatch closes here: a `decor:` shape whose text is not
                        # pure decoration is itself a finding, regardless of its actual
                        # ratio — routed `catalog_gap` because the fix is a rename or a
                        # real caption slot, never a token.
                        foreground = _resolve_color(run.font.color, tokens)
                        ratio = contrast_ratio(foreground.hex, background_hex)
                        findings.append(
                            QAFinding(
                                check="insufficient contrast",
                                slide_id=slide_id,
                                shapes=(shape.name,),
                                measured=round(ratio, 2),
                                threshold=threshold,
                                remedy="catalog_gap",
                                remedy_detail=(
                                    f"{shape.name!r} ({run.text!r}) is named as decoration "
                                    "but contains a letter or a digit — decoration may not "
                                    "carry words or numbers, so it cannot claim WCAG 1.4.3's "
                                    "exemption."
                                ),
                            )
                        )
                        continue

                    foreground = _resolve_color(run.font.color, tokens)
                    ratio = contrast_ratio(foreground.hex, background_hex)
                    if ratio < threshold:
                        detail = (
                            f"palette.{foreground.slot}"
                            if foreground.slot is not None
                            else f"text colour on {shape.name!r}"
                        )
                        findings.append(
                            QAFinding(
                                check="insufficient contrast",
                                slide_id=slide_id,
                                shapes=(shape.name,),
                                measured=round(ratio, 2),
                                threshold=threshold,
                                remedy="token",
                                remedy_detail=detail,
                            )
                        )

    return findings


def contrast_ratio(foreground_hex: str, background_hex: str) -> float:
    """WCAG 2.x: (L1 + 0.05) / (L2 + 0.05), L the relative luminance with sRGB linearisation
    (channel ≤ 0.04045 → c/12.92, else ((c+0.055)/1.055)^2.4). Symmetric in its arguments;
    black on white is 21.0."""
    luminances = (_relative_luminance(foreground_hex), _relative_luminance(background_hex))
    lighter, darker = sorted(luminances, reverse=True)
    return (lighter + 0.05) / (darker + 0.05)


def _relative_luminance(hex_color: str) -> float:
    hex_color = hex_color.lstrip("#")
    channels = (int(hex_color[i : i + 2], 16) / 255 for i in (0, 2, 4))

    def linearise(value: float) -> float:
        return value / 12.92 if value <= 0.04045 else ((value + 0.055) / 1.055) ** 2.4

    r, g, b = (linearise(value) for value in channels)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


# ---------------------------------------------------------------------------
# Slide and shape geometry
# ---------------------------------------------------------------------------


def _slide_id(slide: Slide) -> str:
    """The IR slide id from `cSld/@name`, or `ValueError` if the slide carries no tag."""
    name = slide.name
    if not name.startswith(SLIDE_TAG_PREFIX):
        raise ValueError(
            f"slide has no {SLIDE_TAG_PREFIX!r} tag (name={name!r}) — deterministic QA "
            "findings must be attributable to an IR slide id"
        )
    return name[len(SLIDE_TAG_PREFIX) :]


def _shape_box(shape: BaseShape) -> Box:
    """`shape`'s bounding box in points, from its raw EMU geometry."""
    return Box(
        emu_to_points(shape.left or 0),
        emu_to_points(shape.top or 0),
        emu_to_points(shape.width or 0),
        emu_to_points(shape.height or 0),
    )


def _text_frame(shape: BaseShape) -> TextFrame:
    """`shape.text_frame`, typed. `BaseShape` covers every shape kind this module handles
    (autoshapes, textboxes, connectors), only some of which have a text frame at all — the
    caller is expected to have checked `has_text_frame` (`_is_text_bearing` does)."""
    return shape.text_frame  # pyright: ignore[reportAttributeAccessIssue]


def _is_text_bearing(shape: BaseShape) -> bool:
    """Whether `shape` carries at least one run of actual (non-empty) text."""
    if not shape.has_text_frame:
        return False
    frame = _text_frame(shape)
    return any(run.text for paragraph in frame.paragraphs for run in paragraph.runs)


def _intersection(a: Box, b: Box) -> tuple[float, float]:
    """How much `a` and `b` overlap in each dimension. Negative means a gap, not overlap."""
    dx = min(a.right, b.right) - max(a.x, b.x)
    dy = min(a.bottom, b.bottom) - max(a.y, b.y)
    return dx, dy


#: A stroke width to assume for a connector whose `line.width` is unset (`Emu(0)` — never
#: true for anything `draw.add_connector` draws, which always sets one explicitly; only a
#: guard against a hand-built fixture that does not).
_DEFAULT_LINE_STROKE_PT: Final = 1.0


def _line_stroke_pt(shape: BaseShape) -> float:
    """`shape`'s own line width in points — the actual stroke this connector draws with,
    read from the file rather than assumed."""
    width = shape.line.width  # pyright: ignore[reportAttributeAccessIssue]
    if not width:
        return _DEFAULT_LINE_STROKE_PT
    return emu_to_points(int(width))


def _visual_box(shape: BaseShape, box: Box) -> Box:
    """`box`, widened to the thin rectangle a straight connector's stroke actually occupies.

    python-pptx reports a connector's *bounding box* — degenerate (zero-width or
    zero-height) for exactly the horizontal and vertical lines every diagram in this
    codebase draws, since a perfectly horizontal line is zero pixels tall by definition of
    its own two endpoints. A plain `_intersection` against that box could never register an
    overlap with anything, which is not the same claim as "nothing crosses this line" — it
    is the box, not the geometry, being too thin to answer the question. Only a
    `MSO_SHAPE_TYPE.LINE` shape is widened; every other shape's box is its own, unchanged.
    """
    if shape.shape_type != MSO_SHAPE_TYPE.LINE:
        return box
    stroke = _line_stroke_pt(shape)
    if box.width == 0:
        return Box(box.x - stroke / 2, box.y, stroke, box.height)
    if box.height == 0:
        return Box(box.x, box.y - stroke / 2, box.width, stroke)
    return box


def _safe_area_overshoot(box: Box, safe: Box) -> float | None:
    """How far `box` extends past `safe`, or `None` when it is inside (within `Box.contains`'s
    own sub-point tolerance — the exact check that docstring names as built for this)."""
    if safe.contains(box):
        return None
    return max(
        safe.x - box.x,
        safe.y - box.y,
        box.right - safe.right,
        box.bottom - safe.bottom,
        0.0,
    )


# ---------------------------------------------------------------------------
# Remedy routing — token vs. slot vs. catalog_gap
# ---------------------------------------------------------------------------


def _catalog_slot_boxes_for(component: str, tokens: DesignTokens) -> list[tuple[str, Box]]:
    """Every declared slot box for one component, across all its variants, labelled
    `"<component>.<slot>"`. Used when the caller has told `run_deterministic_qa` which
    component drew a slide (its `components` map), so a "slot" remedy can safely name one —
    see `_catalog_slot_boxes` for the case where it has not."""
    boxes: list[tuple[str, Box]] = []
    canvas = Canvas(tokens)
    entry = catalog.registration(component)
    for variant in entry.variants:
        for slot in variant.slots(canvas):
            label = f"{component}.{slot.name}"
            boxes.append((label, slot.box))
            if slot.item_box is not None:
                boxes.append((label, slot.item_box))
    return boxes


def _catalog_slot_boxes(tokens: DesignTokens) -> list[tuple[str, Box]]:
    """Every declared slot box across every registered component — what a "slot" remedy
    matches a finding's geometry against when the caller has not said which component drew
    the slide. This can only answer "some registered component declares a slot here", never
    "the slide's own component does": `_slot_remedy` reports that distinction rather than
    naming a slot that may only coincide by chance."""
    boxes: list[tuple[str, Box]] = []
    for component in catalog.known_components():
        boxes.extend(_catalog_slot_boxes_for(component, tokens))
    return boxes


_SLOT_COMPONENT_UNKNOWN: Final = "slot (component unknown)"


def _slot_remedy(
    boxes: tuple[Box, ...],
    catalog_boxes: list[tuple[str, Box]],
    *,
    name_slot: bool,
    names: tuple[str, ...],
) -> tuple[RemedyKind, str]:
    """`slot` — naming the first catalog-declared slot any of `boxes` overlaps when
    `name_slot` is true (the caller identified the component), else the geometric fact alone
    (`_SLOT_COMPONENT_UNKNOWN`) — or `catalog_gap` when nothing declared matches at all: the
    shapes sit where no registered component ever puts a slot, so a token or a slot edit
    cannot be the fix; the catalog needs a new one."""
    for box in boxes:
        for label, slot_box in catalog_boxes:
            dx, dy = _intersection(box, slot_box)
            if dx > 0 and dy > 0:
                return "slot", (label if name_slot else _SLOT_COMPONENT_UNKNOWN)
    joined = " and ".join(repr(name) for name in names)
    return "catalog_gap", (
        f"{joined} occupies geometry no registered component slot declares — this is a "
        "component design gap, not a token or slot fix."
    )


# ---------------------------------------------------------------------------
# Text size resolution
# ---------------------------------------------------------------------------


def _resolved_run_size(run: _Run, paragraph: _Paragraph, tokens: DesignTokens) -> float:
    """`run`'s effective point size: its own, else its paragraph's default, else the
    theme's own body size — "resolved through it, not skipped" (module docstring)."""
    if run.font.size is not None:
        return run.font.size.pt
    paragraph_size = paragraph.font.size
    if paragraph_size is not None:
        return paragraph_size.pt
    return tokens.typography.body


def _contrast_threshold(size_pt: float, bold: bool) -> float:
    """`CONTRAST_LARGE` for large text (by size, or by size+bold), else `CONTRAST_NORMAL`."""
    if size_pt >= LARGE_TEXT_PT:
        return CONTRAST_LARGE
    if bold and size_pt >= LARGE_BOLD_TEXT_PT:
        return CONTRAST_LARGE
    return CONTRAST_NORMAL


# ---------------------------------------------------------------------------
# Colour resolution — srgbClr, schemeClr (through the tokens palette) and sysClr
# ---------------------------------------------------------------------------

#: OOXML colour-map slot -> the `Palette` attribute it names. `tx1`/`bg1`/`tx2`/`bg2` are
#: the *mapped* names a shape actually references (`draw.theme_color` writes these, via
#: `MSO_THEME_COLOR.TEXT_1`/`BACKGROUND_1`); `dk1`/`lt1`/`dk2`/`lt2` are the underlying
#: `clrScheme` slot names, which a `schemeClr` can also name directly. python-pptx's default
#: master's colour map is the identity one (`tx1->dk1`, `bg1->lt1`, ...), which is what every
#: fixture in this file and every real render both rely on.
_SCHEME_SLOT: Final[dict[str, str]] = {
    "dk1": "dk1",
    "lt1": "lt1",
    "dk2": "dk2",
    "lt2": "lt2",
    "tx1": "dk1",
    "bg1": "lt1",
    "tx2": "dk2",
    "bg2": "lt2",
    "accent1": "accent1",
    "accent2": "accent2",
    "accent3": "accent3",
    "accent4": "accent4",
    "accent5": "accent5",
    "accent6": "accent6",
    "hlink": "hyperlink",
    "folHlink": "followed_hyperlink",
}


@dataclass(frozen=True)
class _ResolvedColor:
    hex: str
    slot: str | None
    """The palette attribute name this colour resolved through (`schemeClr`), or `None` for
    a literal colour (`srgbClr`/`sysClr`) — used to name the token in a contrast finding."""


def _resolve_color(color_format: ColorFormat, tokens: DesignTokens) -> _ResolvedColor:
    """A `pptx.dml.color.ColorFormat`'s effective hex, resolved per the module docstring:
    `schemeClr` through `tokens.palette`, `sysClr` through its own `lastClr`, `srgbClr`
    literally.

    Reads the underlying colour element directly (`ColorFormat._color._xClr`) because the
    public API splits by type — `.rgb` only ever reads `srgbClr`, `.theme_color` only ever
    reads `schemeClr` — and neither reaches `sysClr`'s `lastClr` at all (PHASE-3A §6.5: the
    one place this check must follow the literal file, not a theme lookup). This is the only
    private attribute this module touches, and only to read it.
    """
    if color_format.type is None:
        raise ValueError("no explicit colour is set here; deterministic QA cannot resolve it")
    element = color_format._color._xClr  # pyright: ignore[reportPrivateUsage]
    return _hex_from_color_element(element, tokens)


def _hex_from_color_element(element: _Element, tokens: DesignTokens) -> _ResolvedColor:
    tag = QName(element).localname
    if tag == "srgbClr":
        value = element.get("val")
        if value is None:
            raise ValueError("<a:srgbClr> with no val cannot be resolved")
        base_hex, slot = value.upper(), None
    elif tag == "schemeClr":
        raw = element.get("val")
        slot = _SCHEME_SLOT.get(raw or "")
        if slot is None:
            raise ValueError(f"unresolvable theme colour slot {raw!r}")
        base_hex = getattr(tokens.palette, slot).upper()
    elif tag == "sysClr":
        last_clr = element.get("lastClr")
        if not last_clr:
            raise ValueError("<a:sysClr> with no lastClr cannot be resolved")
        base_hex, slot = last_clr.upper(), None
    else:
        raise ValueError(f"unsupported colour element <a:{tag}>")

    lum_mod, lum_off = _luminance_adjustment(element)
    return _ResolvedColor(hex=_apply_luminance(base_hex, lum_mod, lum_off), slot=slot)


def _luminance_adjustment(element: _Element) -> tuple[float, float]:
    """`(lumMod, lumOff)` as fractions — `(1.0, 0.0)`, the identity, when neither is present.

    `draw.add_rect`'s `fill_brightness` (`data_card_grid`'s tinted card panels, among
    others) writes exactly these two children onto the colour element it adjusts
    (`ColorFormat.brightness`'s setter, in python-pptx) — reading only the bare `schemeClr`/
    `srgbClr` value and ignoring them would resolve a heavily tinted pale panel as its full-
    strength base colour, which is the wrong colour to check contrast against.
    """
    lum_mod_elm = element.find(qn("a:lumMod"))
    lum_off_elm = element.find(qn("a:lumOff"))
    lum_mod = _percentage(lum_mod_elm, default=1.0)
    lum_off = _percentage(lum_off_elm, default=0.0)
    return lum_mod, lum_off


def _percentage(element: _Element | None, *, default: float) -> float:
    if element is None:
        return default
    value = element.get("val")
    return default if value is None else ST_Percentage.convert_from_xml(value)


def _apply_luminance(hex_color: str, lum_mod: float, lum_off: float) -> str:
    """OOXML's luminance-modulate/-offset transform: convert to HSL, move the L channel
    (`l' = l * lum_mod + lum_off`, clamped), convert back. This is the actual PowerPoint
    "Lighter/Darker %" swatch formula, not an approximation of it — DrawingML's `lumMod`/
    `lumOff` are defined to operate on luminance, not per-channel RGB."""
    if lum_mod == 1.0 and lum_off == 0.0:
        return hex_color
    r, g, b = (int(hex_color[i : i + 2], 16) / 255 for i in (0, 2, 4))
    hue, lightness, saturation = colorsys.rgb_to_hls(r, g, b)
    lightness = min(max(lightness * lum_mod + lum_off, 0.0), 1.0)
    r, g, b = colorsys.hls_to_rgb(hue, lightness, saturation)
    return "".join(f"{round(channel * 255):02X}" for channel in (r, g, b))


def _shape_fill_hex(shape: BaseShape, tokens: DesignTokens) -> str | None:
    """`shape`'s solid fill colour, or `None` when it has none (transparent, a line-only
    shape, or a fill type this check does not need to understand, e.g. a gradient)."""
    fill = getattr(shape, "fill", None)
    if fill is None:
        return None
    if fill.type != MSO_FILL.SOLID:
        return None
    return _resolve_color(fill.fore_color, tokens).hex


def _background_for(
    index: int, shapes: list[BaseShape], boxes: list[Box], tokens: DesignTokens, slide: Slide
) -> str:
    """The fill of the top-most filled shape behind `shapes[index]`, else the slide
    background — check 4's own definition, read straight off the shape tree's z-order
    (earlier shapes are drawn first, so painted over by anything later)."""
    box = boxes[index]
    for candidate_index in range(index - 1, -1, -1):
        candidate = shapes[candidate_index]
        if candidate.shape_type == MSO_SHAPE_TYPE.GROUP:
            continue
        hex_value = _shape_fill_hex(candidate, tokens)
        if hex_value is None:
            continue
        dx, dy = _intersection(boxes[candidate_index], box)
        if dx > 0 and dy > 0:
            return hex_value
    return _slide_background_hex(slide, tokens)


def _slide_background_hex(slide: Slide, tokens: DesignTokens) -> str:
    """The slide's own explicit background fill, or `tokens.palette.lt1` when it has none.

    Reads the raw XML rather than `slide.background.fill`: python-pptx's own docs call that
    property destructive on an inherited background — merely accessing it replaces the
    inheritance with an explicit `noFill` — which this module's purity contract forbids.
    """
    cSld = slide.element.find(qn("p:cSld"))
    bg = cSld.find(qn("p:bg")) if cSld is not None else None
    bg_pr = bg.find(qn("p:bgPr")) if bg is not None else None
    solid_fill = bg_pr.find(qn("a:solidFill")) if bg_pr is not None else None
    if solid_fill is None or len(solid_fill) == 0:
        return tokens.palette.lt1.upper()
    return _hex_from_color_element(solid_fill[0], tokens).hex
