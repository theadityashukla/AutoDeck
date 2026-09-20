"""`autodeck.design.diagrams` — the three geometries rendered as native shapes (task 3a.6b).

The brief's done-when is "nodes are individually selectable and recolourable in
PowerPoint", so this file checks it the way 3a.7 checked the same requirement for icons:
real python-pptx objects, one `<p:sp>`/`<p:cxnSp>` per node/connector, never a `<p:grpSp>`,
every fill and line a `schemeClr` theme reference and never `srgbClr`. `TestSelectability`
carries that check for all three geometries; the rest of the file is the other two
done-when clauses (labels obey their physical budget, not just their word count; edges are
drawn only where the geometry actually implies one) plus B31's local half — LibreOffice,
since there is no PowerPoint here.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from pptx.presentation import Presentation
from pptx.slide import Slide

from autodeck.design.diagrams import place_diagram
from autodeck.design.fonts import is_available
from autodeck.design.layout_kit import Box, Canvas, Frame, LayoutOverflowError
from autodeck.design.theme.master_builder import new_presentation, save_themed
from autodeck.design.theme.tokens import DesignTokens
from autodeck.ir.models import (
    Citation,
    Claim,
    DiagramAxis,
    DiagramSpec,
    LabelFraming,
    LayeredStackSpec,
    ProcessFlowSpec,
    ProcessStep,
    QuadrantItem,
    StackLayer,
    TwoByTwoSpec,
)
from autodeck.render.qa.libreoffice import render_pptx

TOKENS_DIR = Path(__file__).resolve().parents[1] / "config" / "tokens"

#: Same choice `test_catalog.py` makes and for the same reason: any real installed font
#: does, since these tests check the diagram engine's plumbing, never a font-specific
#: value; `Liberation Sans` is what is actually installed here (B11).
TEST_FAMILY = "Liberation Sans"

requires_test_font = pytest.mark.skipif(
    not is_available(TEST_FAMILY), reason=f"{TEST_FAMILY} not installed"
)

QUOTE = "Inference cost per request fell from $0.41 to $0.24 after the cutover."


def tokens_for(family: str = TEST_FAMILY) -> DesignTokens:
    base = DesignTokens.load(TOKENS_DIR / "dev.json")
    return base.model_copy(
        update={
            "typography": base.typography.model_copy(update={"major": family, "minor": family})
        }
    )


def _citation() -> Citation:
    return Citation.for_quote(
        doc_id="vendor-report",
        page=4,
        bbox=(10.0, 20.0, 300.0, 44.0),
        quote=QUOTE,
        retrieved_by="writer",
    )


def _claim(text: str = "Inference cost per request fell by 41%.") -> Claim:
    return Claim(text=text, citations=[_citation()])


def _framing() -> LabelFraming:
    return LabelFraming(reason="category_name")


def _step(
    node_id: str, label: str, order: int, *, transition: str | None = None
) -> ProcessStep:
    return ProcessStep(
        id=node_id, label=label, order=order, transition=transition, framing=_framing()
    )


def _item(
    node_id: str, label: str, x: float, y: float, *, factual: bool = False
) -> QuadrantItem:
    if factual:
        return QuadrantItem(id=node_id, label=label, x=x, y=y, claim=_claim())
    return QuadrantItem(id=node_id, label=label, x=x, y=y, framing=_framing())


def _layer(node_id: str, label: str, level: int) -> StackLayer:
    return StackLayer(id=node_id, label=label, level=level, framing=_framing())


def _process_flow(*, with_transitions: bool = True) -> DiagramSpec:
    steps = [
        _step("s1", "Discover", 1, transition="Kickoff done" if with_transitions else None),
        _step("s2", "Design", 2, transition="Sign-off" if with_transitions else None),
        _step("s3", "Launch", 3),
    ]
    return DiagramSpec(
        relationship="sequence", kind="process_flow", process_flow=ProcessFlowSpec(steps=steps)
    )


def _two_by_two() -> DiagramSpec:
    return DiagramSpec(
        relationship="classification",
        kind="two_by_two",
        two_by_two=TwoByTwoSpec(
            x_axis=DiagramAxis(name="Cost to serve", low="Low", high="High"),
            y_axis=DiagramAxis(name="Adoption speed", low="Slow", high="Fast"),
            items=[
                _item("q1", "Self-serve", 0.15, 0.85),
                _item("q2", "Enterprise", 0.85, 0.25, factual=True),
                _item("q3", "Mid-market", 0.55, 0.55),
            ],
        ),
    )


def _layered_stack(support: str = "rests_on") -> DiagramSpec:
    return DiagramSpec(
        relationship="hierarchy_foundation",
        kind="layered_stack",
        layered_stack=LayeredStackSpec(
            support=support,  # type: ignore[arg-type]
            layers=[
                _layer("l1", "Data platform", 1),
                _layer("l2", "ML pipeline", 2),
                _layer("l3", "Product layer", 3),
            ],
        ),
    )


def _frame(tokens: DesignTokens) -> tuple[Presentation, Slide, Frame]:
    presentation = new_presentation(tokens)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = Canvas(tokens).on(slide)
    return presentation, slide, frame


# ---------------------------------------------------------------------------
# Individual selectability + theme-linked colour — the brief's done-when
# ---------------------------------------------------------------------------


class TestSelectability:
    """Every node its own `<p:sp>`, every fill/line a `schemeClr`, never a `<p:grpSp>`."""

    @requires_test_font
    @pytest.mark.parametrize(
        ("build", "expected_shapes", "expected_connectors"),
        [
            # 3 chevrons + 3 labels + 2 transition labels, 2 transition arrows
            (_process_flow, 3 + 3 + 2, 2),
            # 3 dots + 3 item labels + 6 axis/pole/name labels, 2 axis lines
            (_two_by_two, 3 + 3 + 6, 2),
            # 3 bands + 3 labels, no edges
            (_layered_stack, 3 + 3, 0),
        ],
    )
    def test_every_node_is_its_own_shape(
        self, build, expected_shapes, expected_connectors
    ) -> None:
        tokens = tokens_for()
        _, slide, frame = _frame(tokens)
        place_diagram(frame, frame.canvas.content, build())

        from pptx.oxml.ns import qn

        shape_elements = slide.shapes._spTree.findall(qn("p:sp"))
        connector_elements = slide.shapes._spTree.findall(qn("p:cxnSp"))
        group_elements = slide.shapes._spTree.findall(qn("p:grpSp"))

        assert len(shape_elements) == expected_shapes
        assert len(connector_elements) == expected_connectors
        assert group_elements == []
        # Real, distinct shape objects — the python-pptx-level version of the same check.
        assert len(slide.shapes) == expected_shapes + expected_connectors
        assert len({id(shape._element) for shape in slide.shapes}) == len(slide.shapes)

    @requires_test_font
    @pytest.mark.parametrize("build", [_process_flow, _two_by_two, _layered_stack])
    def test_every_fill_and_line_is_a_theme_reference_never_an_rgb(self, build) -> None:
        tokens = tokens_for()
        presentation, _slide, frame = _frame(tokens)
        place_diagram(frame, frame.canvas.content, build())

        pptx_path = Path("/tmp") / "diagram-selectability-check.pptx"
        save_themed(presentation, tokens, pptx_path)
        with zipfile.ZipFile(pptx_path) as package:
            xml = package.read("ppt/slides/slide1.xml").decode("utf-8")

        assert "schemeClr" in xml
        assert "srgbClr" not in xml
        pptx_path.unlink()

    @requires_test_font
    def test_no_image_no_smartart_ever(self) -> None:
        """D10/D11: never a picture, never a graphicFrame diagram, for any of the three."""
        tokens = tokens_for()
        presentation, _slide, frame = _frame(tokens)
        for build in (_process_flow, _two_by_two, _layered_stack):
            place_diagram(frame, frame.canvas.content, build())

        pptx_path = Path("/tmp") / "diagram-no-image-check.pptx"
        save_themed(presentation, tokens, pptx_path)
        with zipfile.ZipFile(pptx_path) as package:
            names = package.namelist()
            xml = package.read("ppt/slides/slide1.xml").decode("utf-8")
        assert not [n for n in names if n.startswith("ppt/media/")]
        assert "<p:pic>" not in xml
        assert "dgm:" not in xml  # SmartArt's own namespace prefix
        pptx_path.unlink()


# ---------------------------------------------------------------------------
# Edges: drawn only where the geometry implies one
# ---------------------------------------------------------------------------


class TestEdges:
    @requires_test_font
    def test_process_flow_draws_one_arrow_per_declared_transition(self) -> None:
        tokens = tokens_for()
        _, slide, frame = _frame(tokens)
        place_diagram(frame, frame.canvas.content, _process_flow(with_transitions=True))
        from pptx.oxml.ns import qn

        assert len(slide.shapes._spTree.findall(qn("p:cxnSp"))) == 2

    @requires_test_font
    def test_process_flow_with_no_transitions_draws_no_arrows(self) -> None:
        tokens = tokens_for()
        _, slide, frame = _frame(tokens)
        place_diagram(frame, frame.canvas.content, _process_flow(with_transitions=False))
        from pptx.oxml.ns import qn

        assert slide.shapes._spTree.findall(qn("p:cxnSp")) == []

    @requires_test_font
    def test_two_by_two_draws_only_its_own_axis_lines_never_a_node_to_node_edge(self) -> None:
        """`TwoByTwoSpec.edges` is empty by construction (an arrow would claim flow between
        positions, which the data does not support) — the two connectors here are the axis
        lines, not edges between items."""
        tokens = tokens_for()
        spec = _two_by_two()
        assert spec.edges == ()
        _, slide, frame = _frame(tokens)
        place_diagram(frame, frame.canvas.content, spec)
        from pptx.oxml.ns import qn

        assert len(slide.shapes._spTree.findall(qn("p:cxnSp"))) == 2

    @requires_test_font
    def test_layered_stack_draws_no_connectors_at_all(self) -> None:
        """`LayeredStackSpec.edges` is empty (adjacency carries "rests on"; an arrow would
        claim flow instead) — nothing here should ever draw a connector."""
        tokens = tokens_for()
        spec = _layered_stack()
        assert spec.edges == ()
        _, slide, frame = _frame(tokens)
        place_diagram(frame, frame.canvas.content, spec)
        from pptx.oxml.ns import qn

        assert slide.shapes._spTree.findall(qn("p:cxnSp")) == []


# ---------------------------------------------------------------------------
# Physical fit — over budget by word count is caught at construction; this file checks
# the OTHER failure, a label short enough by word count that still does not fit its box.
# ---------------------------------------------------------------------------


class TestPhysicalFit:
    @requires_test_font
    def test_process_flow_raises_when_the_box_is_far_too_small(self) -> None:
        tokens = tokens_for()
        _, _, frame = _frame(tokens)
        with pytest.raises(LayoutOverflowError):
            place_diagram(frame, Box(0, 0, 40, 20), _process_flow())

    @requires_test_font
    def test_two_by_two_raises_when_axis_labels_dont_fit_the_box(self) -> None:
        tokens = tokens_for()
        _, _, frame = _frame(tokens)
        with pytest.raises(LayoutOverflowError):
            place_diagram(frame, Box(0, 0, 10, 10), _two_by_two())

    @requires_test_font
    def test_layered_stack_raises_when_a_band_is_too_short_for_its_label(self) -> None:
        """15pt split three ways with a 3pt gutter between each is a genuine (if tiny) 3pt
        band — `split_rows` succeeds — so the raise proven here is `_place_label`'s, not
        `split_rows`' unrelated "rows do not fit the box at all" precondition failure."""
        tokens = tokens_for()
        _, _, frame = _frame(tokens)
        with pytest.raises(LayoutOverflowError, match="layered_stack layer"):
            place_diagram(frame, Box(0, 0, 300, 15), _layered_stack())

    @requires_test_font
    def test_a_single_overlong_word_cannot_be_rescued_by_wrapping(self) -> None:
        """The `too_wide_words` path — `_measure`'s own message, not just `Box.reserve`'s."""
        tokens = tokens_for()
        _, _, frame = _frame(tokens)
        spec = DiagramSpec(
            relationship="hierarchy_foundation",
            kind="layered_stack",
            layered_stack=LayeredStackSpec(
                support="peers",
                layers=[
                    _layer("l1", "Antidisestablishmentarianism", 1),
                    _layer("l2", "Sales", 2),
                ],
            ),
        )
        with pytest.raises(LayoutOverflowError, match="wider than"):
            place_diagram(frame, Box(0, 0, 60, 200), spec)


# ---------------------------------------------------------------------------
# spec.title is never drawn here — the component embedding the diagram owns it
# ---------------------------------------------------------------------------


@requires_test_font
def test_the_diagram_title_is_not_rendered_by_this_module() -> None:
    tokens = tokens_for()
    _, slide, frame = _frame(tokens)
    spec = DiagramSpec(
        relationship="sequence",
        kind="process_flow",
        title="A Wildly Distinctive Headline Nobody Else Would Type",
        process_flow=ProcessFlowSpec(
            steps=[_step("s1", "Discover", 1), _step("s2", "Launch", 2)]
        ),
    )
    place_diagram(frame, frame.canvas.content, spec)

    texts = [
        run.text
        for shape in slide.shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs  # type: ignore[attr-defined]
        for run in paragraph.runs
    ]
    assert not any("Wildly Distinctive Headline" in text for text in texts)


# ---------------------------------------------------------------------------
# claim_nodes()'s identity contract survives a render — nothing here rebuilds a node
# ---------------------------------------------------------------------------


@requires_test_font
def test_rendering_does_not_disturb_claim_node_identity() -> None:
    spec = _two_by_two()
    before = {id(site.node) for site in spec.claim_nodes()}
    assert before  # the factual item in `_two_by_two()` gives at least one claim node

    tokens = tokens_for()
    _, _, frame = _frame(tokens)
    place_diagram(frame, frame.canvas.content, spec)

    after = {id(site.node) for site in spec.claim_nodes()}
    assert before == after


# ---------------------------------------------------------------------------
# Rendered by LibreOffice — B31: no PowerPoint here
# ---------------------------------------------------------------------------


class TestRenderedByLibreOffice:
    @pytest.mark.render
    @pytest.mark.parametrize("build", [_process_flow, _two_by_two, _layered_stack])
    def test_libreoffice_rasterises_each_geometry_without_error(
        self, build, tmp_path: Path
    ) -> None:
        tokens = tokens_for(family="Inter") if is_available("Inter") else tokens_for()
        presentation, _slide, frame = _frame(tokens)
        place_diagram(frame, frame.canvas.content, build())

        pptx_path = tmp_path / "deck.pptx"
        save_themed(presentation, tokens, pptx_path)

        result = render_pptx(pptx_path, tmp_path, tokens)
        assert result.page_count == 1
        assert result.images[0].stat().st_size > 0
