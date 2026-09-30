"""Components 13-15: `framework_diagram`, `timeline`, `chart_focus`.

What is distinctive about these three, and therefore what the tests are for: none of them
draws its main content itself. Two hand the body to `diagrams.place_diagram` and one to
`charts.place_chart`. So the tests pin the *seams* — what the component must not do on the
engine's behalf — rather than re-testing the engines, which have their own suites
(`test_diagram_renderers.py`, `test_charts.py`).
"""

from __future__ import annotations

import zipfile
from dataclasses import replace
from pathlib import Path

import pytest
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from autodeck.design.charts import source_line
from autodeck.design.components import preview
from autodeck.design.components.catalog import (
    check_overflow,
    components_missing_previews,
    known_components,
    spec_for,
)
from autodeck.design.components.renderers import chart_focus, framework_diagram, timeline
from autodeck.design.fonts import is_available
from autodeck.design.grammar import lint_slide_ir
from autodeck.design.layout_kit import Canvas, LayoutOverflowError
from autodeck.design.theme.master_builder import new_presentation, save_themed
from autodeck.design.theme.tokens import DesignTokens
from autodeck.ir.models import (
    Block,
    Citation,
    DiagramAxis,
    DiagramSpec,
    LabelFraming,
    ProcessFlowSpec,
    ProcessStep,
    QuadrantItem,
    Slide,
    TwoByTwoSpec,
)

TOKENS_DIR = Path(__file__).resolve().parents[1] / "config" / "tokens"

#: Same choice `test_catalog.py` and `test_diagram_renderers.py` make: any real installed
#: font does, since these tests check the seam plumbing, never a font-specific value.
TEST_FAMILY = "Liberation Sans"

requires_test_font = pytest.mark.skipif(
    not is_available(TEST_FAMILY), reason=f"{TEST_FAMILY} not installed"
)

NAMES = ("framework_diagram", "timeline", "chart_focus")

QUOTE = "Inference cost per request fell from $0.41 to $0.24 after the cutover."


def tokens_for(family: str = TEST_FAMILY) -> DesignTokens:
    base = DesignTokens.load(TOKENS_DIR / "dev.json")
    return base.model_copy(
        update={
            "typography": base.typography.model_copy(update={"major": family, "minor": family})
        }
    )


def _citation() -> Citation:
    return Citation.for_quote(
        doc_id="vendor-report",
        page=4,
        bbox=(10.0, 20.0, 300.0, 44.0),
        quote=QUOTE,
        retrieved_by="writer",
    )


def _blank_slide(tokens: DesignTokens):
    presentation = new_presentation(tokens)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    return presentation, slide


def _all_shape_text(slide) -> list[str]:
    return [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]


# ---------------------------------------------------------------------------
# Every registered component renders its preview example and answers to the registry
# ---------------------------------------------------------------------------


@requires_test_font
@pytest.mark.parametrize("name", NAMES)
def test_each_renders_its_preview_example(name: str) -> None:
    """`<module>.render` draws `preview.EXAMPLES[name]` onto a blank dev-token slide
    without raising."""
    tokens = tokens_for()
    canvas = Canvas(tokens)
    _presentation, slide = _blank_slide(tokens)
    renderer = {
        "framework_diagram": framework_diagram.render,
        "timeline": timeline.render,
        "chart_focus": chart_focus.render,
    }[name]
    renderer(slide, canvas, preview.EXAMPLES[name])


@requires_test_font
@pytest.mark.parametrize("name", NAMES)
def test_the_registry_answers_for_each(name: str) -> None:
    """`spec_for(name, tokens)` returns a spec, and `name` is not in
    `components_missing_previews()` — its golden PNG is committed."""
    tokens = tokens_for()
    spec = spec_for(name, tokens)
    assert spec.name == name
    assert name not in components_missing_previews()


@requires_test_font
@pytest.mark.parametrize("name", NAMES)
def test_an_overlong_headline_is_rejected_before_render(name: str) -> None:
    """`check_overflow` reports the `headline` slot for a headline far past its budget,
    without any render happening."""
    tokens = tokens_for()
    findings = check_overflow({"headline": "word " * 300}, name, tokens)
    assert any(finding.startswith(f"{name}.headline:") for finding in findings)


# ---------------------------------------------------------------------------
# The diagram-led pair's seams
# ---------------------------------------------------------------------------


@requires_test_font
@pytest.mark.parametrize("name", ["framework_diagram", "timeline"])
def test_the_diagram_title_is_never_drawn(name: str) -> None:
    """Give the example's `DiagramSpec` the title "UNMISTAKABLE-DIAGRAM-TITLE", render, and
    assert that string appears in no shape's text. The headline is the slide's one voice."""
    tokens = tokens_for()
    canvas = Canvas(tokens)
    _presentation, slide = _blank_slide(tokens)

    content = preview.EXAMPLES[name]
    marked_diagram = content.diagram.model_copy(update={"title": "UNMISTAKABLE-DIAGRAM-TITLE"})
    marked_content = replace(content, diagram=marked_diagram)

    renderer = framework_diagram.render if name == "framework_diagram" else timeline.render
    renderer(slide, canvas, marked_content)

    texts = _all_shape_text(slide)
    assert not any("UNMISTAKABLE-DIAGRAM-TITLE" in text for text in texts)


@requires_test_font
def test_timeline_refuses_a_diagram_that_is_not_a_sequence() -> None:
    """A `two_by_two` `DiagramSpec` in `TimelineContent` makes `timeline.render` raise
    `ValueError` whose message names "two_by_two". A timeline must never draw a 2x2."""
    tokens = tokens_for()
    canvas = Canvas(tokens)
    _presentation, slide = _blank_slide(tokens)

    two_by_two = DiagramSpec(
        relationship="classification",
        kind="two_by_two",
        two_by_two=TwoByTwoSpec(
            x_axis=DiagramAxis(name="Cost to serve", low="Low", high="High"),
            y_axis=DiagramAxis(name="Adoption speed", low="Slow", high="Fast"),
            items=[
                QuadrantItem(
                    id="q1",
                    label="Self-serve",
                    x=0.15,
                    y=0.85,
                    framing=LabelFraming(reason="category_name"),
                ),
                QuadrantItem(
                    id="q2",
                    label="Enterprise",
                    x=0.85,
                    y=0.25,
                    framing=LabelFraming(reason="category_name"),
                ),
            ],
        ),
    )
    content = timeline.TimelineContent(headline="Not a sequence", diagram=two_by_two)

    with pytest.raises(ValueError, match="two_by_two"):
        timeline.render(slide, canvas, content)


@requires_test_font
@pytest.mark.parametrize("name", ["framework_diagram", "timeline"])
def test_diagram_components_group_nothing_and_embed_no_picture(name: str) -> None:
    """Saved PPTX: zero `<p:grpSp>` on the slide, zero `ppt/media/` parts, and at least one
    `<a:schemeClr` — every node individually selectable and theme-coloured (D10/D11)."""
    tokens = tokens_for()
    canvas = Canvas(tokens)
    presentation, slide = _blank_slide(tokens)
    renderer = framework_diagram.render if name == "framework_diagram" else timeline.render
    renderer(slide, canvas, preview.EXAMPLES[name])

    group_elements = slide.shapes._spTree.findall(qn("p:grpSp"))
    assert group_elements == []

    pptx_path = Path("/tmp") / f"{name}-selectability-check.pptx"
    save_themed(presentation, tokens, pptx_path)
    try:
        with zipfile.ZipFile(pptx_path) as package:
            names = package.namelist()
            xml = package.read("ppt/slides/slide1.xml").decode("utf-8")
        assert not [n for n in names if n.startswith("ppt/media/")]
        assert "schemeClr" in xml
    finally:
        pptx_path.unlink(missing_ok=True)


# ---------------------------------------------------------------------------
# chart_focus's own seams
# ---------------------------------------------------------------------------


@requires_test_font
def test_chart_focus_draws_one_chart_and_one_source_line() -> None:
    """Exactly one chart graphic frame on the slide, zero `ppt/media/` parts, and the
    source text appears exactly once — the chart's own, not a second `frame.caption`."""
    tokens = tokens_for()
    canvas = Canvas(tokens)
    presentation, slide = _blank_slide(tokens)
    chart_focus.render(slide, canvas, preview.EXAMPLES["chart_focus"])

    chart_shapes = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.CHART]
    assert len(chart_shapes) == 1

    expected_source = source_line(preview.EXAMPLES["chart_focus"].chart.source_citations)
    texts = _all_shape_text(slide)
    occurrences = [t for t in texts if expected_source in t]
    assert len(occurrences) == 1

    pptx_path = Path("/tmp") / "chart-focus-media-check.pptx"
    save_themed(presentation, tokens, pptx_path)
    try:
        with zipfile.ZipFile(pptx_path) as package:
            names = package.namelist()
        assert not [n for n in names if n.startswith("ppt/media/")]
    finally:
        pptx_path.unlink(missing_ok=True)


@requires_test_font
def test_chart_focus_reserves_the_takeaway_before_the_chart() -> None:
    """A takeaway far past one line raises `LayoutOverflowError` from `chart_focus.render`
    rather than being drawn with the chart squeezed to fit around it."""
    tokens = tokens_for()
    canvas = Canvas(tokens)
    _presentation, slide = _blank_slide(tokens)

    base = preview.EXAMPLES["chart_focus"]
    content = chart_focus.ChartFocusContent(
        headline=base.headline,
        chart=base.chart,
        takeaway="word " * 300,
        accent=base.accent,
    )
    with pytest.raises(LayoutOverflowError):
        chart_focus.render(slide, canvas, content)


# ---------------------------------------------------------------------------
# The documented gap: word-budget-safe but not physically safe until render
# ---------------------------------------------------------------------------


@requires_test_font
def test_a_node_label_under_its_word_budget_can_still_overflow_at_render() -> None:
    """Build a valid `DiagramSpec` whose node label is within `max_label_words` but uses
    words long enough not to fit its box, and assert `framework_diagram.render` raises
    `LayoutOverflowError`. Pins the documented limitation: diagram labels are protected
    physically only at render, not by `check_overflow`."""
    tokens = tokens_for()
    canvas = Canvas(tokens)
    _presentation, slide = _blank_slide(tokens)

    # One word (well within `process_flow`'s 4-word budget) long enough that it cannot fit
    # ANY chevron this slide can offer once there are enough steps sharing the row —
    # `DiagramSpec` accepts this at construction (it is one word, under budget); only
    # physical measurement at render, against the real width one of many narrow chevrons
    # gets, can catch it.
    long_word = "Pneumonoultramicroscopicsilicovolcanoconiosis"
    steps = [
        ProcessStep(
            id="s1",
            label=long_word,
            order=1,
            framing=LabelFraming(reason="stage_name"),
        ),
        *[
            ProcessStep(
                id=f"s{n}", label="Step", order=n, framing=LabelFraming(reason="stage_name")
            )
            for n in range(2, 13)
        ],
    ]
    diagram = DiagramSpec(
        relationship="sequence", kind="process_flow", process_flow=ProcessFlowSpec(steps=steps)
    )
    content = framework_diagram.FrameworkDiagramContent(headline="A sequence", diagram=diagram)

    with pytest.raises(LayoutOverflowError):
        framework_diagram.render(slide, canvas, content)


# ---------------------------------------------------------------------------
# The phase's exit criterion
# ---------------------------------------------------------------------------


def test_all_fifteen_components_are_registered_and_previewed() -> None:
    """`known_components()` has exactly the fifteen PHASE-3A names and
    `components_missing_previews()` is empty — the phase's exit criterion, asserted."""
    expected = {
        "agenda",
        "before_after",
        "big_number",
        "bullets_supporting",
        "callout_takeaway",
        "chart_focus",
        "closing_cta",
        "data_card_grid",
        "evidence_with_figure",
        "framework_diagram",
        "quote",
        "section_divider",
        "timeline",
        "title",
        "two_column_compare",
    }
    assert set(known_components()) == expected
    assert len(expected) == 15
    assert components_missing_previews() == []


# ---------------------------------------------------------------------------
# D13's grammar lints, on the diagram-led examples
# ---------------------------------------------------------------------------


@requires_test_font
@pytest.mark.parametrize("name", ["framework_diagram", "timeline"])
def test_diagram_led_examples_pass_the_grammar_lints(name: str) -> None:
    """The preview example, as a slide, produces no blocking D13 finding from
    `autodeck.design.grammar` in `diagram_led` mode (<=1 diagram, <=60 words)."""
    content = preview.EXAMPLES[name]
    slide = Slide(
        id="s1",
        narrative_role="test",
        component=name,
        communication_mode="diagram_led",
        blocks=[
            Block(id="headline", kind="framing", slot="headline", text=content.headline),
            Block(id="diagram", kind="diagram", slot="diagram", diagram=content.diagram),
        ],
    )
    findings = lint_slide_ir(slide)
    assert not [f for f in findings if f.severity == "blocking"]
