"""Design system: fonts, budgets, layout kit, theme XML, and the SVG converter.

Everything here is pure computation or XML generation, so it runs in CI. The parts that
need LibreOffice and an installed font stack carry the `render` marker and stay local
per B10.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx.presentation import Presentation
from pptx.slide import Slide

from autodeck.design.budgets import compute_budget, load_metrics, measure_height, wrap_text
from autodeck.design.fonts import FontFile, FontNotFoundError, is_available, resolve_face
from autodeck.design.icons.consistency import (
    IconConsistencyError,
    check_deck_icon_consistency,
    placed_icons,
)
from autodeck.design.icons.custgeom import custgeom_xml, place_icon
from autodeck.design.icons.library import (
    Icon,
    IconNotFoundError,
    UnsupportedIconElementError,
    available_concepts,
    icon_names,
    load_icon,
    parse_svg,
    resolve_icon,
)
from autodeck.design.icons.style import DEFAULT_ICON_STYLE, IconStyle, UnknownIconSizeError
from autodeck.design.icons.svg_path import (
    Close,
    Cubic,
    Line,
    Move,
    SvgPathError,
    circle_subpath,
    parse_path,
)
from autodeck.design.layout_kit import (
    Box,
    Canvas,
    Frame,
    LayoutOverflowError,
    TextStyle,
    pt_to_emu,
)
from autodeck.design.theme.master_builder import apply_theme, build_theme_xml, new_presentation
from autodeck.design.theme.tokens import DesignTokens, Typography

TOKENS_DIR = Path(__file__).resolve().parents[1] / "config" / "tokens"

#: The container/CI font. Present because tests need *a* real font file to measure;
#: nothing here asserts anything about Aptos, which cannot be installed here (B11).
TEST_FAMILY = "Inter"

requires_test_font = pytest.mark.skipif(
    not is_available(TEST_FAMILY), reason=f"{TEST_FAMILY} not installed"
)


# ---------------------------------------------------------------------------
# Fonts — a missing family must be loud
# ---------------------------------------------------------------------------


def test_a_missing_family_raises_rather_than_substituting() -> None:
    """The invariant of this module. A fallback here corrupts every budget silently."""
    with pytest.raises(FontNotFoundError, match="not found"):
        resolve_face("Definitely Not An Installed Typeface")


def test_the_error_explains_the_aptos_case() -> None:
    """Whoever hits this is usually setting the project up for the first time."""
    with pytest.raises(FontNotFoundError, match="Microsoft 365"):
        resolve_face("Aptos Nonexistent Variant")


def test_tokens_declaring_a_missing_font_fail_up_front() -> None:
    tokens = DesignTokens(
        name="broken", typography=Typography(major="No Such Face", minor="No Such Face")
    )
    with pytest.raises(FontNotFoundError):
        tokens.require_fonts()


@requires_test_font
def test_resolve_returns_a_real_file() -> None:
    assert resolve_face(TEST_FAMILY).path.exists()


@requires_test_font
def test_each_face_resolves_to_its_own_file() -> None:
    """The 3a.6 defect in one assertion: asking for bold used to hand back the regular file.

    Four distinct paths, each verified against the style bits the file itself declares —
    so this fails both if the resolver falls back to regular and if it picks up a file
    whose name says Bold while its `OS/2`/`head` tables say otherwise.
    """
    faces = {
        (bold, italic): resolve_face(TEST_FAMILY, bold=bold, italic=italic)
        for bold in (False, True)
        for italic in (False, True)
    }
    assert len({face.path for face in faces.values()}) == 4
    for (bold, italic), face in faces.items():
        assert face.bold is bold and face.italic is italic
        assert face.path.exists()


def test_a_missing_face_is_as_loud_as_a_missing_family() -> None:
    """B11's rule applies per face: substituting regular for bold is the same silent
    corruption as substituting one family for another, and it is how every bold budget in
    the component library came to be measured against the wrong file. Inter Display is the
    real case — it genuinely ships no bold-italic — so this asserts the decision rather
    than a hypothetical: the resolver raises, naming the face, instead of quietly
    measuring something it can measure."""
    if not is_available("Inter Display"):
        pytest.skip("Inter Display not installed")
    assert is_available("Inter Display", bold=True)
    assert is_available("Inter Display", italic=True)
    with pytest.raises(FontNotFoundError, match="bold italic"):
        resolve_face("Inter Display", bold=True, italic=True)


@requires_test_font
def test_require_fonts_checks_bold_and_italic_not_just_regular(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A family present in regular alone cannot carry a bold headline's budget, and the
    build should say so before it measures anything rather than at the first component
    that needs it. Recorded as "which faces were asked for", because that is the whole
    behaviour — the old version asked for one face per family and was satisfied."""
    import autodeck.design.fonts as fonts_module

    asked: list[tuple[str, bool, bool]] = []
    real = fonts_module.resolve_face

    def _spy(family: str, *, bold: bool = False, italic: bool = False) -> FontFile:
        asked.append((family, bold, italic))
        return real(family, bold=bold, italic=italic)

    monkeypatch.setattr(fonts_module, "resolve_face", _spy)
    DesignTokens(
        name="faces", typography=Typography(major=TEST_FAMILY, minor=TEST_FAMILY)
    ).require_fonts()

    assert (TEST_FAMILY, False, False) in asked
    assert (TEST_FAMILY, True, False) in asked
    assert (TEST_FAMILY, False, True) in asked


# ---------------------------------------------------------------------------
# Budgets — measured from real glyph metrics
# ---------------------------------------------------------------------------


@requires_test_font
def test_metrics_come_from_the_font_file() -> None:
    metrics = load_metrics(TEST_FAMILY)
    assert metrics.units_per_em > 0
    assert metrics.advances
    assert metrics.char_width("W", 16) > metrics.char_width("i", 16)


@requires_test_font
def test_bold_text_measures_wider_than_regular() -> None:
    """The arithmetic half of 3a.6: bold advances really are wider, so a budget computed
    from the regular face over-promises. Asserted as a strict inequality rather than
    against a pinned percentage — the gap is a property of whichever family is installed
    (+2.8% for Inter, +6.1% for Inter Display at 32pt), and pinning one family's number
    here would make this test a fact about this container instead of about the fix."""
    text = "Inference costs fall when the model is small enough to serve"
    regular = load_metrics(TEST_FAMILY).text_width(text, 32)
    bold = load_metrics(TEST_FAMILY, bold=True).text_width(text, 32)
    assert bold > regular


@requires_test_font
def test_bold_text_wraps_sooner_than_regular_at_the_same_width() -> None:
    """The consequence that actually bites, stated in the units the budget is written in.

    A width is chosen so the regular face fits on one line with only a few percent to
    spare — the same few percent `quote`'s headline had — and the bold face does not.
    Before this fix `wrap_text` had no `bold` argument at all and both calls returned the
    identical measurement, which is what let a bold headline be promised a line it did not
    have."""
    text = "Inference costs fall when the model is small enough to serve"
    regular_width = load_metrics(TEST_FAMILY).text_width(text, 32)
    box = regular_width * 1.02
    assert wrap_text(text, TEST_FAMILY, 32, box).line_count == 1
    assert wrap_text(text, TEST_FAMILY, 32, box, bold=True).line_count == 2


@requires_test_font
def test_a_bold_budget_holds_less_than_the_same_box_in_regular() -> None:
    """`SlotBudget` carries the face, so `fits()` and `max_chars` both move with it."""
    common = {"slot": "headline", "family": TEST_FAMILY, "size_pt": 32, "height_pt": 90}
    regular = compute_budget(width_pt=600, **common)  # type: ignore[arg-type]
    bold = compute_budget(width_pt=600, bold=True, **common)  # type: ignore[arg-type]
    assert bold.max_chars < regular.max_chars
    assert "bold" in bold.describe()


@requires_test_font
def test_wrapping_respects_the_measured_width() -> None:
    text = "The quick brown fox jumps over the lazy dog and keeps on running."
    narrow = wrap_text(text, TEST_FAMILY, 16, 120)
    wide = wrap_text(text, TEST_FAMILY, 16, 600)
    assert narrow.line_count > wide.line_count
    assert narrow.width_pt <= 120


@requires_test_font
def test_a_word_wider_than_the_box_is_reported_not_broken() -> None:
    """The fix is a smaller size or a different component, not a hyphen we cannot render."""
    measurement = wrap_text(
        "Pneumonoultramicroscopicsilicovolcanoconiosis", TEST_FAMILY, 40, 60
    )
    assert measurement.overflow_width


@requires_test_font
def test_newlines_start_new_lines() -> None:
    assert wrap_text("one\ntwo\nthree", TEST_FAMILY, 12, 500).line_count == 3


@requires_test_font
def test_budget_fits_short_text_and_rejects_long_text() -> None:
    budget = compute_budget(
        slot="body", family=TEST_FAMILY, size_pt=16, width_pt=300, height_pt=100
    )
    assert budget.max_lines >= 1
    assert budget.fits("A short line.")
    assert not budget.fits("word " * 400)


@requires_test_font
def test_budget_describes_itself_for_a_prompt() -> None:
    """§6.7: budgets reach the content agent as hard constraints in its prompt."""
    budget = compute_budget(
        slot="headline", family=TEST_FAMILY, size_pt=32, width_pt=600, height_pt=90
    )
    assert "headline" in budget.describe()
    assert str(budget.max_lines) in budget.describe()


@requires_test_font
def test_equal_character_counts_can_need_very_different_widths() -> None:
    """v1's `chars_per_line` heuristic assumed a character count predicts a line's width;
    this is the property that assumption gets wrong, and the property `wrap_text` has to
    get right instead. Ten narrow glyphs and ten wide ones are the same length by
    `len()` and nothing alike by advance width — a heuristic keyed on the former cannot
    represent the latter, which is why none exists anywhere in this module (PHASE-2B.md
    2b.1: "v1's `chars_per_line` heuristic is not reproduced")."""
    narrow = wrap_text("l" * 10, TEST_FAMILY, 24, 2000)
    wide = wrap_text("W" * 10, TEST_FAMILY, 24, 2000)
    assert wide.width_pt > narrow.width_pt * 2


def test_no_character_count_heuristic_is_exposed() -> None:
    """Static companion to the property test above, and the one part of it that needs no
    font at all: `chars_per_line` (or any spelling of it) must not exist as an importable
    name, so nothing downstream can reach for the shortcut budgets.py exists to replace."""
    import autodeck.design.budgets as budgets_module

    names = " ".join(dir(budgets_module)).lower()
    assert "charsperline" not in names.replace("_", "")


@requires_test_font
def test_measure_height_never_undershoots_the_calibrated_line_height() -> None:
    """2b.1's bias-direction guard, and the one budgets test that needs no LibreOffice.

    `budgets.py`'s own module docstring says why under-prediction is the dangerous
    direction: a budget that under-predicts height passes text that then overflows at
    render. `_line_height_factor` used to under-predict by ~7.4% (hhea ascent+descent,
    ~1.117x, against LibreOffice's real ~1.2x single-spaced pitch — see
    `budgets.LINE_HEIGHT_FACTOR`'s docstring for the full derivation).

    `1.20` here is written as a literal, not imported as `LINE_HEIGHT_FACTOR`: importing it
    would make this test check `measure_height` against whatever the constant currently
    says, which passes trivially no matter how that constant later changes — including a
    change back in the under-predicting direction. The literal is what was actually
    calibrated against a real render (`test_budget_check.py`'s render-marked pitch tests);
    a future edit that quietly shrinks `LINE_HEIGHT_FACTOR` should fail *this* test, not
    sail through because the test moved with it. Over-predicting by a little is fine — see
    `budget_check.py`'s `RENDER_HEADROOM_FRACTION` for how much — so this only guards the floor.
    """
    calibrated_factor = 1.20
    for size_pt, line_spacing in [(16, 1.0), (22, 1.15), (40, 1.25)]:
        # A single short line in a very wide box: exactly one line, so the floor below is
        # exactly the quantity `measure_height` must not fall under, not an approximation
        # of it diluted by however many lines the text happens to wrap to.
        height = measure_height("Single line.", TEST_FAMILY, size_pt, 5000, line_spacing)
        floor = size_pt * line_spacing * calibrated_factor
        assert height >= floor - 1e-6, (
            f"measure_height({size_pt=}, {line_spacing=}) = {height:.3f}pt under-predicts "
            f"the calibrated floor of {floor:.3f}pt — the dangerous direction."
        )


# ---------------------------------------------------------------------------
# Layout kit
# ---------------------------------------------------------------------------


def test_columns_share_the_width_and_the_gutter() -> None:
    columns = Box(0, 0, 100, 50).split_columns(3, gutter=5)
    assert len(columns) == 3
    assert columns[0].width == pytest.approx(30)
    assert columns[1].x == pytest.approx(35)
    assert columns[-1].right == pytest.approx(100)


def test_columns_that_cannot_fit_raise() -> None:
    with pytest.raises(ValueError, match="do not fit"):
        Box(0, 0, 10, 10).split_columns(4, gutter=20)


def test_split_top_returns_the_remainder_below_the_gutter() -> None:
    taken, rest = Box(0, 0, 100, 100).split_top(30, gutter=10)
    assert (taken.height, rest.y, rest.height) == (30, 40, 60)


def test_split_bottom_mirrors_split_top() -> None:
    rest, taken = Box(0, 0, 100, 100).split_bottom(30, gutter=10)
    assert (taken.y, taken.height, rest.height) == (70, 30, 60)


def test_grid_produces_rows_of_columns() -> None:
    grid = Box(0, 0, 100, 100).grid(2, 3, gutter=4)
    assert len(grid) == 2
    assert all(len(row) == 3 for row in grid)


def test_align_within_centres() -> None:
    outer = Box(0, 0, 100, 100)
    centred = Box(0, 0, 20, 10).align_within(outer, horizontal="center", vertical="middle")
    assert (centred.center_x, centred.center_y) == (50, 50)


def test_snap_to_baseline_rounds_down() -> None:
    assert Box(0, 17, 10, 10).snap_to_baseline(6).y == 12


def test_contains_is_the_safe_area_check() -> None:
    outer = Box(0, 0, 100, 100)
    assert outer.contains(Box(10, 10, 50, 50))
    assert not outer.contains(Box(80, 80, 50, 50))


def test_reserve_aligns_a_content_sized_block_inside_its_region() -> None:
    band = Box(0, 0, 100, 100).reserve(40, valign="middle", what="a block")
    assert (band.y, band.height, band.width) == (30, 40, 100)


def test_reserve_refuses_rather_than_returning_a_shorter_box() -> None:
    """The kit's one rule. A shorter box would leave the caller free to draw into it
    anyway, which is how silent overflow gets back in."""
    with pytest.raises(LayoutOverflowError, match="shrink text to fit"):
        Box(0, 0, 100, 50).reserve(80, what="too much")


# ---------------------------------------------------------------------------
# Stacks — one declaration is both the measurement and the drawing
# ---------------------------------------------------------------------------


def _frame() -> Frame:
    """A real slide to draw on. python-pptx only — no LibreOffice, so this runs in CI."""
    tokens = DesignTokens.load(TOKENS_DIR / "dev.json")
    presentation = new_presentation(tokens)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    return Canvas(tokens).on(slide)


@requires_test_font
def test_a_stack_draws_each_item_where_it_measured_it() -> None:
    """The property the whole second layer exists for: the block that was measured and the
    block that was drawn are the same declaration, so they cannot drift apart."""
    frame = _frame()
    style = frame.canvas.style("body")
    stack = frame.stack("test", 300)
    stack.text("First item.", style)
    stack.text("Second item.", style, gap=10)
    stack.place(Box(0, 0, 300, 400))

    first, second = frame.slide.shapes
    heights = stack.item_heights
    assert first.top == pt_to_emu(0)
    assert second.top == pytest.approx(pt_to_emu(heights[0] + 10), abs=2)
    assert stack.height == pytest.approx(heights[0] + 10 + heights[1])


@requires_test_font
def test_place_returns_the_region_below_the_block() -> None:
    frame = _frame()
    stack = frame.stack("test", 300).text("One line.", frame.canvas.style("body"))
    remainder = stack.place(Box(0, 0, 300, 400), gutter=12)
    assert remainder.y == pytest.approx(stack.height + 12)
    assert remainder.bottom == pytest.approx(400)


@requires_test_font
def test_a_stack_taller_than_its_box_raises_rather_than_shrinking() -> None:
    frame = _frame()
    stack = frame.stack("a headline", 300)
    stack.text(
        "A sentence long enough to wrap several times over. " * 6, frame.canvas.style("body")
    )
    with pytest.raises(LayoutOverflowError, match="a headline"):
        stack.place(Box(0, 0, 300, 30))


@requires_test_font
def test_a_word_too_wide_to_wrap_is_refused_as_it_is_added() -> None:
    """Height alone cannot see this failure: an over-wide word measures as exactly one line
    and renders straight past the edge. The stack refuses it where the word can be named."""
    frame = _frame()
    stack = frame.stack("a caption", 40)
    with pytest.raises(LayoutOverflowError, match="Supercalifragilistic"):
        stack.text("Supercalifragilistic", frame.canvas.style("body"))


@requires_test_font
def test_min_height_holds_two_columns_rows_level() -> None:
    """How a comparison stays a comparison: a short row on one side takes the height its
    opposite number needs, so row n is level with row n across the gutter."""
    frame = _frame()
    style = frame.canvas.style("body")
    tall = frame.stack("left", 200)
    tall.text("A point long enough to wrap onto two lines in this column.", style)
    short = frame.stack("right", 200)
    short.text("Short.", style, min_height=tall.item_heights[0])
    assert short.item_heights == tall.item_heights


@requires_test_font
def test_space_reserves_height_and_draws_nothing() -> None:
    frame = _frame()
    stack = frame.stack("test", 300).space(50)
    stack.place(Box(0, 0, 300, 400))
    assert stack.height == 50
    assert len(frame.slide.shapes) == 0


@requires_test_font
def test_items_stack_a_list_with_its_own_row_gap() -> None:
    frame = _frame()
    stack = frame.stack("points", 300).items(
        ["One.", "Two.", "Three."], frame.canvas.style("body"), row_gap=8
    )
    assert stack.height == pytest.approx(sum(stack.item_heights) + 16)
    stack.place(Box(0, 0, 300, 400))
    # Each item is a marker shape and a text shape.
    assert len(frame.slide.shapes) == 6


def test_the_caption_strip_and_the_body_come_from_one_split() -> None:
    """Two definitions of where the body ends is how the catalog's `source` budget and the
    renderer's source line would come to disagree."""
    canvas = Canvas(DesignTokens.load(TOKENS_DIR / "dev.json"))
    body, caption = canvas.body_and_caption()
    assert caption == canvas.caption_strip
    assert caption.y == pytest.approx(body.bottom + canvas.gutter)
    assert caption.bottom == pytest.approx(canvas.content.bottom)


@requires_test_font
def test_a_component_with_no_source_draws_no_caption() -> None:
    frame = _frame()
    assert frame.caption(Box(0, 0, 300, 20), "") is None
    assert len(frame.slide.shapes) == 0


def test_style_scale_multiplies_the_roles_token_size() -> None:
    """`big_number`'s figure is twice `display` — a multiple of a token, never a loose
    point size, so a client's type scale still moves it."""
    canvas = Canvas(DesignTokens.load(TOKENS_DIR / "dev.json"))
    assert canvas.style("display", scale=2.0).size == pytest.approx(canvas.size("display") * 2)


# ---------------------------------------------------------------------------
# Measurement and rendering must not diverge — the bug TextStyle exists to prevent
# ---------------------------------------------------------------------------


@requires_test_font
def test_measurement_accounts_for_line_spacing() -> None:
    """The first spike render put an accent rule through a headline because it did not."""
    canvas = Canvas(DesignTokens.load(TOKENS_DIR / "dev.json"))
    text = "A headline long enough to wrap onto a second line in this column."
    single = TextStyle(family=TEST_FAMILY, size=16, line_spacing=1.0)
    loose = single.with_(line_spacing=1.5)
    assert canvas.measure(text, 300, loose) > canvas.measure(text, 300, single)


@requires_test_font
def test_measurement_accounts_for_paragraph_spacing() -> None:
    canvas = Canvas(DesignTokens.load(TOKENS_DIR / "dev.json"))
    text = "First paragraph.\nSecond paragraph."
    tight = TextStyle(family=TEST_FAMILY, size=16)
    spaced = tight.with_(space_after=12.0)
    assert canvas.measure(text, 300, spaced) == pytest.approx(
        canvas.measure(text, 300, tight) + 12.0
    )


@requires_test_font
def test_fit_returns_a_box_of_the_measured_height() -> None:
    canvas = Canvas(DesignTokens.load(TOKENS_DIR / "dev.json"))
    style = canvas.style("body")
    region = Box(0, 0, 300, 400)
    fitted = canvas.fit("Some text that wraps a couple of times over here.", region, style)
    assert fitted.height == pytest.approx(
        canvas.measure("Some text that wraps a couple of times over here.", 300, style)
    )
    assert fitted.height < region.height


# ---------------------------------------------------------------------------
# Tokens and theme XML
# ---------------------------------------------------------------------------


def test_both_shipped_token_sets_load() -> None:
    for name in ("aptos", "dev"):
        assert DesignTokens.load(TOKENS_DIR / f"{name}.json").name


def test_aptos_is_the_deliverable_default() -> None:
    """B11. The dev set exists for containers; it must not become the default."""
    aptos = DesignTokens.load(TOKENS_DIR / "aptos.json")
    assert aptos.typography.major == "Aptos Display"
    assert aptos.typography.minor == "Aptos"


def test_the_dev_token_set_mirrors_the_aptos_type_structure() -> None:
    """Same major/minor split and type scale, so layout work transfers between them."""
    aptos = DesignTokens.load(TOKENS_DIR / "aptos.json")
    dev = DesignTokens.load(TOKENS_DIR / "dev.json")
    assert aptos.typography.model_dump(exclude={"major", "minor"}) == dev.typography.model_dump(
        exclude={"major", "minor"}
    )
    assert aptos.palette == dev.palette


def test_theme_xml_carries_the_palette_and_fonts() -> None:
    tokens = DesignTokens.load(TOKENS_DIR / "dev.json")
    xml = build_theme_xml(tokens)
    assert f'<a:latin typeface="{tokens.typography.major}"/>' in xml
    assert tokens.palette.accent1.upper() in xml
    assert "<a:clrScheme" in xml and "<a:fontScheme" in xml


def test_theme_xml_declares_all_twelve_colour_slots() -> None:
    xml = build_theme_xml(DesignTokens.load(TOKENS_DIR / "dev.json"))
    for slot in ("dk1", "lt1", "dk2", "lt2", "hlink", "folHlink"):
        assert f"<a:{slot}>" in xml
    for index in range(1, 7):
        assert f"<a:accent{index}>" in xml


def test_theme_xml_has_three_entries_in_every_format_list() -> None:
    """Fewer than three makes the package invalid and PowerPoint offers to repair it."""
    xml = build_theme_xml(DesignTokens.load(TOKENS_DIR / "dev.json"))
    assert xml.count("<a:effectStyle>") == 3
    assert xml.count("<a:ln ") == 3


def test_text_and_background_slots_use_sysclr() -> None:
    """PowerPoint's text/background UI pairs them wrongly with a plain srgbClr."""
    xml = build_theme_xml(DesignTokens.load(TOKENS_DIR / "dev.json"))
    assert '<a:sysClr val="windowText"' in xml
    assert '<a:sysClr val="window"' in xml


@requires_test_font
def test_applying_a_theme_replaces_only_the_theme_part() -> None:
    import io
    import zipfile

    tokens = DesignTokens.load(TOKENS_DIR / "dev.json")
    buffer = io.BytesIO()
    new_presentation(tokens).save(buffer)
    original = buffer.getvalue()

    themed = apply_theme(original, build_theme_xml(tokens))

    with (
        zipfile.ZipFile(io.BytesIO(original)) as before,
        zipfile.ZipFile(io.BytesIO(themed)) as after,
    ):
        assert before.namelist() == after.namelist()
        changed = [n for n in before.namelist() if before.read(n) != after.read(n)]
        assert changed == ["ppt/theme/theme1.xml"]


def test_applying_a_theme_to_a_package_without_one_fails() -> None:
    import io
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as package:
        package.writestr("hello.txt", "not a pptx")
    with pytest.raises(ValueError, match="no ppt/theme"):
        apply_theme(buffer.getvalue(), "<a:theme/>")


# ---------------------------------------------------------------------------
# SVG path parsing
# ---------------------------------------------------------------------------


def test_absolute_and_relative_lines_agree() -> None:
    absolute = parse_path("M 10 10 L 20 20")[0]
    relative = parse_path("m 10 10 l 10 10")[0]
    assert absolute == relative == [Move(10, 10), Line(20, 20)]


def test_horizontal_and_vertical_shorthands() -> None:
    assert parse_path("M0 0 H 10 V 5")[0] == [Move(0, 0), Line(10, 0), Line(10, 5)]


def test_implicit_repeat_after_moveto_is_a_lineto() -> None:
    assert parse_path("M0 0 10 10 20 20")[0] == [Move(0, 0), Line(10, 10), Line(20, 20)]


def test_compound_paths_produce_several_subpaths() -> None:
    assert len(parse_path("M0 0 L1 1 Z M5 5 L6 6 Z")) == 2


def test_quadratics_are_elevated_to_cubics() -> None:
    segments = parse_path("M0 0 Q 10 0 10 10")[0]
    assert isinstance(segments[1], Cubic)
    assert (segments[1].x, segments[1].y) == (10, 10)


def test_smooth_cubic_reflects_the_previous_control_point() -> None:
    segments = parse_path("M0 0 C 0 5 5 10 10 10 S 20 5 20 0")[0]
    reflected = segments[2]
    assert isinstance(reflected, Cubic)
    assert (reflected.x1, reflected.y1) == (15, 10)  # mirror of (5, 10) about (10, 10)


def test_arcs_become_cubics() -> None:
    segments = parse_path("M0 0 A 10 10 0 0 1 10 10")[0]
    assert all(isinstance(s, Cubic) for s in segments[1:])
    assert (segments[-1].x, segments[-1].y) == pytest.approx((10, 10))  # type: ignore[union-attr]


def test_packed_arc_flags_parse_correctly() -> None:
    """The bug Lucide's `zap` icon exposed: `0 00-2.474` is three values, not two.

    A regex tokeniser reads `00` as the single number 0 and every argument after it shifts.
    """
    segments = parse_path("M15.914 4a1.5 1.5 0 00-2.474-1.561")[0]
    assert (segments[-1].x, segments[-1].y) == pytest.approx((13.44, 2.439))  # type: ignore[union-attr]


def test_a_degenerate_arc_becomes_a_line() -> None:
    segments = parse_path("M0 0 A 0 0 0 0 1 10 10")[0]
    assert (segments[-1].x, segments[-1].y) == (10, 10)  # type: ignore[union-attr]


def test_closepath_returns_to_the_subpath_start() -> None:
    segments = parse_path("M5 5 L10 10 Z L20 20")[0]
    assert isinstance(segments[2], Close)
    assert segments[3] == Line(20, 20)


def test_empty_path_data_is_empty() -> None:
    assert parse_path("") == []
    assert parse_path("   ") == []


def test_path_data_starting_with_a_number_is_rejected() -> None:
    with pytest.raises(SvgPathError, match="starts with a number"):
        parse_path("10 10 L 20 20")


def test_an_unsupported_command_is_rejected() -> None:
    with pytest.raises(SvgPathError):
        parse_path("M0 0 X 5")


def test_a_circle_becomes_four_cubics_returning_to_the_start() -> None:
    segments = circle_subpath(10, 10, 5)
    assert isinstance(segments[0], Move)
    assert sum(isinstance(s, Cubic) for s in segments) == 4
    assert (segments[-2].x, segments[-2].y) == pytest.approx((15, 10))  # type: ignore[union-attr]


# ---------------------------------------------------------------------------
# Icon library and custGeom
# ---------------------------------------------------------------------------


def test_the_vendored_library_is_present_and_licensed() -> None:
    """Lucide is ISC — redistributable, but only with the notice alongside it."""
    assert len(icon_names()) >= 5
    license_file = Path(__file__).resolve().parents[1] / "autodeck/design/icons/lucide/LICENSE"
    assert license_file.exists()
    assert "ISC License" in license_file.read_text(encoding="utf-8")


@pytest.mark.parametrize("name", icon_names())
def test_every_vendored_icon_parses_and_converts(name: str) -> None:
    """Ten icons exercising lines, arcs, circles and compound paths."""
    icon = load_icon(name)
    assert icon.subpaths
    assert icon.view_width == icon.view_height == 24
    for subpath in icon.subpaths:
        xml = custgeom_xml(subpath, icon.view_width, icon.view_height)
        assert xml.startswith("<a:custGeom>")
        assert "<a:moveTo>" in xml


def test_an_unknown_icon_lists_what_is_available() -> None:
    with pytest.raises(IconNotFoundError, match="Available"):
        load_icon("no-such-icon")


def test_svg_primitives_all_convert() -> None:
    source = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" stroke-width="2">
      <path d="M1 1 L2 2"/><circle cx="12" cy="12" r="6"/>
      <rect x="1" y="1" width="4" height="4"/><line x1="0" y1="0" x2="9" y2="9"/>
      <polyline points="1,1 2,2 3,3"/><polygon points="4,4 5,5 6,4"/>
      <ellipse cx="8" cy="8" rx="3" ry="2"/>
    </svg>"""
    icon = parse_svg(source, name="mixed")
    assert len(icon.subpaths) == 7
    assert icon.stroke_width == 2


def test_custgeom_declares_an_unfilled_path() -> None:
    """`fill="none"` is what keeps a stroke a stroke instead of a filled wedge."""
    icon = load_icon("trending-up")
    xml = custgeom_xml(icon.subpaths[0], icon.view_width, icon.view_height)
    assert 'fill="none"' in xml
    assert "<a:pathLst>" in xml


def test_custgeom_coordinates_scale_into_the_path_space() -> None:
    from autodeck.design.icons.custgeom import PATH_SPACE

    xml = custgeom_xml([Move(0, 0), Line(24, 24)], 24, 24)
    assert 'x="0"' in xml
    assert f'x="{PATH_SPACE}"' in xml


# ---------------------------------------------------------------------------
# Semantic library — concept -> glyph (§9, task 3a.7)
# ---------------------------------------------------------------------------


def test_a_concept_resolves_to_its_mapped_icon() -> None:
    assert resolve_icon("risk").name == "circle-alert"
    assert resolve_icon("growth").name == "trending-up"


def test_resolve_icon_falls_back_to_a_literal_filename() -> None:
    """A caller that already knows the exact glyph is not forced through the concept table."""
    assert resolve_icon("clock").name == "clock"


def test_an_unknown_concept_lists_both_concepts_and_filenames() -> None:
    with pytest.raises(IconNotFoundError, match="Known concepts") as excinfo:
        resolve_icon("no-such-thing")
    assert "Known icon files" in str(excinfo.value)


def test_every_mapped_concept_resolves_to_a_vendored_icon() -> None:
    """The semantic table is only as good as its entries actually resolving."""
    for concept in available_concepts():
        assert resolve_icon(concept).name in icon_names()


# ---------------------------------------------------------------------------
# No image fallback: an inconvertible element escalates, it does not vanish
# ---------------------------------------------------------------------------


def test_an_unsupported_svg_element_raises_rather_than_being_dropped() -> None:
    """D10/D11: a glyph that cannot convert is an escalation, not a `.png` — and not a
    silently incomplete shape either, which is what dropping the element would produce."""
    source = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24">
      <mask id="m"><rect x="0" y="0" width="24" height="24"/></mask>
    </svg>"""
    with pytest.raises(UnsupportedIconElementError, match="<mask>"):
        parse_svg(source, name="masked")


def test_benign_metadata_elements_are_skipped_not_escalated() -> None:
    source = """<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24">
      <title>An icon</title><desc>Description</desc>
      <path d="M1 1 L2 2"/>
    </svg>"""
    icon = parse_svg(source, name="titled")
    assert len(icon.subpaths) == 1


# ---------------------------------------------------------------------------
# Family on the shape name — what makes deck-wide consistency checkable at all
# ---------------------------------------------------------------------------


def test_a_vendored_icon_carries_the_default_family() -> None:
    assert load_icon("clock").family == "lucide"


def test_placed_shapes_are_named_with_family_and_icon() -> None:
    frame = _frame()
    place_icon(frame.slide, Box(0, 0, 40, 40), load_icon("clock"))
    assert frame.slide.shapes
    assert all(shape.name.startswith("icon:lucide:clock") for shape in frame.slide.shapes)


# ---------------------------------------------------------------------------
# Frame.icon — closing 3a.2's carried finding
# ---------------------------------------------------------------------------


def test_frame_icon_places_a_semantic_icon_without_reaching_past_the_frame() -> None:
    """The finding 3a.2 carried: icon placement lived outside the `Frame` vocabulary,
    taking a slide directly. `Frame.icon` is the fix — a renderer never touches
    `icons.custgeom` or a bare `slide` to place one."""
    frame = _frame()
    shapes = frame.icon(Box(0, 0, 40, 40), "risk")
    assert shapes
    assert all(shape.name.startswith("icon:lucide:circle-alert") for shape in shapes)


def test_frame_icon_accepts_a_literal_icon_name_too() -> None:
    frame = _frame()
    shapes = frame.icon(Box(0, 0, 40, 40), "clock", color="accent2")
    assert shapes
    assert shapes[0].line.color.theme_color is not None


# ---------------------------------------------------------------------------
# IconStyle — the deck-wide vocabulary a size scale is picked from
# ---------------------------------------------------------------------------


def test_default_icon_style_has_a_three_step_size_scale() -> None:
    assert DEFAULT_ICON_STYLE.size("sm") < DEFAULT_ICON_STYLE.size("md")
    assert DEFAULT_ICON_STYLE.size("md") < DEFAULT_ICON_STYLE.size("lg")


def test_an_unknown_icon_size_role_lists_whats_available() -> None:
    with pytest.raises(UnknownIconSizeError, match="Available"):
        DEFAULT_ICON_STYLE.size("xl")


def test_stroke_pt_scales_with_the_size_it_is_placed_at() -> None:
    """One weight *ratio* per deck, not one fixed point value — a fixed value would look
    right at one size and wrong at another."""
    style = IconStyle()
    assert style.stroke_pt(24.0) == pytest.approx(2.0)
    assert style.stroke_pt(48.0) == pytest.approx(4.0)


# ---------------------------------------------------------------------------
# check_deck_icon_consistency — enforced against real placed shapes, not hoped for
# ---------------------------------------------------------------------------

#: Lucide's own design ratio — stroke-width 2 on a 24-unit viewBox.
_LUCIDE_RATIO = 2.0 / 24.0


def _icon_with_ratio(name: str, ratio: float, family: str = "lucide") -> Icon:
    """A synthetic icon whose stroke-to-viewBox ratio is exactly `ratio`.

    Reuses a real vendored icon's geometry (so it still converts to real subpaths) but
    substitutes `stroke_width`/`view_width` so the *placed* stroke-to-size ratio is under
    the test's control rather than Lucide's actual design ratio.
    """
    from dataclasses import replace

    return replace(load_icon("clock"), name=name, stroke_width=ratio * 24.0, family=family)


def _new_slide() -> tuple[Presentation, Slide]:
    """A fresh presentation and a slide on it — real python-pptx objects throughout."""
    tokens = DesignTokens.load(TOKENS_DIR / "dev.json")
    presentation = new_presentation(tokens)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    return presentation, slide


def test_placed_icons_reads_family_name_and_geometry_off_the_shape_name() -> None:
    """`clock` is two subpaths (a circle and the hands), so two shapes — same icon."""
    presentation, slide = _new_slide()
    place_icon(slide, Box(0, 0, 24, 24), load_icon("clock"))
    icons = placed_icons(presentation)
    assert len(icons) == len(load_icon("clock").subpaths)
    for icon in icons:
        assert icon.family == "lucide"
        assert icon.name == "clock"
        assert icon.slide_index == 0
        assert icon.side_pt == pytest.approx(24.0)
        assert icon.stroke_pt == pytest.approx(icon.side_pt * _LUCIDE_RATIO)


def test_a_single_icon_cannot_be_inconsistent_with_itself() -> None:
    presentation, slide = _new_slide()
    place_icon(slide, Box(0, 0, 24, 24), load_icon("clock"))
    check_deck_icon_consistency(presentation)  # does not raise


def test_one_family_one_weight_across_sizes_passes() -> None:
    """`place_icon`'s own default already scales stroke with size — this is the case that
    proves the checker accepts it rather than mistaking a scaled-up icon for a heavier one."""
    presentation, slide = _new_slide()
    place_icon(slide, Box(0, 0, 24, 24), load_icon("clock"))
    place_icon(slide, Box(0, 0, 48, 48), load_icon("target"))
    check_deck_icon_consistency(presentation)  # does not raise


def test_mixing_two_icon_families_raises() -> None:
    presentation, slide = _new_slide()
    place_icon(slide, Box(0, 0, 24, 24), load_icon("clock"))
    place_icon(slide, Box(0, 0, 24, 24), _icon_with_ratio("clock", _LUCIDE_RATIO, "phosphor"))
    with pytest.raises(IconConsistencyError, match="mixes icon families"):
        check_deck_icon_consistency(presentation)


def test_a_drifting_stroke_weight_raises() -> None:
    presentation, slide = _new_slide()
    place_icon(slide, Box(0, 0, 24, 24), load_icon("clock"))  # ratio ~0.083
    place_icon(slide, Box(0, 0, 24, 24), _icon_with_ratio("target", 0.25))
    with pytest.raises(IconConsistencyError, match="stroke weight"):
        check_deck_icon_consistency(presentation)


def test_too_many_distinct_sizes_raises() -> None:
    """One weight ratio held constant (the family default), sizes fanned out past the
    deck-wide scale's step count."""
    presentation, slide = _new_slide()
    for side in (10.0, 11.0, 12.0, 13.0, 14.0, 15.0):
        place_icon(slide, Box(0, 0, side, side), load_icon("clock"))
    with pytest.raises(IconConsistencyError, match="distinct sizes"):
        check_deck_icon_consistency(presentation)
