"""`icon_pillars` (B35): the first component that can place an icon on a slide face.

Covers the renderer, its catalog registration, the IR adapter, and the end-to-end paths B35
exists to unblock: an icon on a rendered deck that the aesthetic loop's icon actions can
reach, and GATE 3 check 2 (an icon selectable and recolourable from the theme palette).
"""

from __future__ import annotations

import dataclasses
import itertools
from pathlib import Path
from typing import Any, cast

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.presentation import Presentation as PresentationType

from autodeck.design import grammar
from autodeck.design.components import catalog, preview
from autodeck.design.components.renderers import icon_pillars
from autodeck.design.components.renderers.icon_pillars import (
    IconPillarsContent,
    Pillar,
    PillarIcon,
)
from autodeck.design.icons.consistency import ICON_NAME_PREFIX
from autodeck.design.icons.library import CONCEPT_TO_ICON, IconNotFoundError
from autodeck.design.layout_kit import Canvas
from autodeck.design.theme.master_builder import new_presentation, save_themed
from autodeck.design.theme.tokens import DesignTokens
from autodeck.ir.actions import SetIconColour, SwapGlyph, apply_action, fact_fingerprint
from autodeck.ir.models import (
    Block,
    Citation,
    Claim,
    Deck,
    IconRef,
    Slide,
    SlideStyle,
)
from autodeck.render.qa.aesthetic import catalog_slot_lookup, library_concept_lookup
from autodeck.render.renderer import (
    RenderStageError,
    adapt_slide,
    placeable_slots,
    render_deck,
)

TOKENS_DIR = Path(__file__).resolve().parents[1] / "config" / "tokens"

#: Any real installed font does: these tests check structure, not a font-specific value.
TEST_FAMILY = "Liberation Sans"


def tokens_for() -> DesignTokens:
    base = DesignTokens.load(TOKENS_DIR / "dev.json")
    return base.model_copy(
        update={
            "typography": base.typography.model_copy(
                update={"major": TEST_FAMILY, "minor": TEST_FAMILY}
            )
        }
    )


#: Framing text may not carry digits (A5 would demote it to an uncited claim), so test labels
#: are words.
_NAMES = ("Alpha", "Bravo", "Charlie", "Delta", "Echo", "Foxtrot")

_ids = itertools.count(1)


def _id() -> str:
    return f"p{next(_ids)}"


def _framing(slot: str, text: str) -> Block:
    return Block(id=_id(), kind="framing", slot=slot, text=text)


def _icon_block(glyph: str, *, concept: str | None = None, color: str = "accent1") -> Block:
    icon = IconRef(concept=concept or glyph, glyph_id=glyph, color_token=color)
    return Block(id=_id(), kind="icon", slot="pillar_icon", icon=icon)


def _claim(slot: str, text: str) -> Block:
    citation = Citation.for_quote(
        quote=text,
        doc_id="programme-review",
        page=3,
        bbox=(10.0, 20.0, 300.0, 44.0),
        retrieved_by="validator",
    )
    return Block(
        id=_id(),
        kind="claim",
        slot=slot,
        claim=Claim(text=text, citations=[citation], verdict="supported"),
    )


def _slide(blocks: list[Block], slide_id: str = "s1", **style: Any) -> Slide:
    return Slide(
        id=slide_id,
        narrative_role="test",
        component="icon_pillars",
        message_ids=["m1"],
        blocks=blocks,
        style=SlideStyle(**style),
    )


def _deck(*slides: Slide) -> Deck:
    return Deck(
        run_id="r1",
        project="p",
        client="c",
        audience="a",
        version=1,
        theme_ref="t",
        component_lib_version=catalog.COMPONENT_LIB_VERSION,
        slides=list(slides),
    )


def _blocks(
    count: int, *, points: int | None = None, glyphs: tuple[str, ...] | None = None
) -> list[Block]:
    """A headline, `count` icons, `count` labels and `points` points (default: none)."""
    glyphs = glyphs or ("users", "shield-check", "zap", "trending-up", "clock")
    blocks = [_framing("headline", "Capabilities that compound")]
    blocks.extend(_icon_block(glyphs[i % len(glyphs)]) for i in range(count))
    blocks.extend(_framing("pillar_label", f"Pillar {_NAMES[i]}") for i in range(count))
    blocks.extend(_framing("pillar_point", f"Point {_NAMES[i]}") for i in range(points or 0))
    return blocks


def _example(count: int) -> IconPillarsContent:
    example = preview.EXAMPLES["icon_pillars"]
    assert isinstance(example, IconPillarsContent)
    return dataclasses.replace(example, pillars=example.pillars[:count])


def _render_content(content: IconPillarsContent, tmp_path: Path) -> PresentationType:
    tokens = tokens_for()
    presentation = new_presentation(tokens)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    icon_pillars.render(slide, Canvas(tokens), content)
    path = save_themed(presentation, tokens, tmp_path / "pillars.pptx")
    return Presentation(str(path))


def _icon_boxes(slide: Any) -> list[tuple[float, float, float, float]]:
    """One (left, top, width, height) per distinct icon, in left-to-right order.

    A multi-stroke glyph is several shapes sharing one square, so shapes are de-duplicated
    by their box.
    """
    boxes = {
        (shape.left.pt, shape.top.pt, shape.width.pt, shape.height.pt)
        for shape in slide.shapes
        if shape.name.startswith(ICON_NAME_PREFIX)
    }
    return sorted(boxes)


def _text_shape_containing(slide: Any, text: str) -> Any:
    matches = [
        shape
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip() == text
    ]
    assert len(matches) == 1, (text, len(matches))
    return matches[0]


@pytest.mark.parametrize("count", [3, 4])
def test_three_and_four_pillars_render_with_icon_label_and_optional_point(
    count: int, tmp_path: Path
) -> None:
    """Both counts render without overflow at standard scale with the preview example's
    content; each pillar's icon shapes carry the `icon:` name prefix
    (`icons.consistency.ICON_NAME_PREFIX`) and the label's text frame is in the same column."""
    content = _example(count)
    presentation = _render_content(content, tmp_path)
    slide = presentation.slides[0]

    boxes = _icon_boxes(slide)
    assert len(boxes) == count

    for (left, top, width, height), pillar in zip(boxes, content.pillars, strict=True):
        icon_centre = left + width / 2
        label = _text_shape_containing(slide, pillar.label)
        assert label.left.pt <= icon_centre <= label.left.pt + label.width.pt
        assert label.top.pt >= top + height  # directly beneath, never beside or above

        assert pillar.point is not None
        point = _text_shape_containing(slide, pillar.point)
        assert point.left.pt <= icon_centre <= point.left.pt + point.width.pt
        assert point.top.pt >= label.top.pt + label.height.pt - 0.5

    # All icons share one top edge: the row reads as a set, not a staircase.
    assert len({round(top, 1) for _, top, _, _ in boxes}) == 1


def test_a_pillar_without_a_point_renders_icon_and_label_only(tmp_path: Path) -> None:
    content = _example(3)
    content = dataclasses.replace(
        content, pillars=[dataclasses.replace(p, point=None) for p in content.pillars]
    )
    slide = _render_content(content, tmp_path).slides[0]

    assert len(_icon_boxes(slide)) == 3
    # Nothing is composed in the point's place.
    texts = [cast(Any, s).text_frame.text for s in slide.shapes if s.has_text_frame]
    assert not any("None" in text for text in texts)


@pytest.mark.parametrize("count", [2, 5])
def test_two_or_five_pillars_are_refused(count: int, tmp_path: Path) -> None:
    """`render` raises `ValueError`; the adapter raises `RenderStageError`."""
    pillar = Pillar(icon=PillarIcon(glyph="zap"), label="Speed")
    content = IconPillarsContent(headline="Too few or too many", pillars=[pillar] * count)

    tokens = tokens_for()
    presentation = new_presentation(tokens)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    with pytest.raises(ValueError, match=str(count)):
        icon_pillars.render(slide, Canvas(tokens), content)

    content_type = catalog.registration("icon_pillars").content_type
    with pytest.raises(RenderStageError, match=str(count)):
        adapt_slide(_slide(_blocks(count)), content_type)


def test_a_colour_that_is_not_a_theme_slot_is_refused() -> None:
    """D1: icons take a theme slot. A literal colour would survive a theme change."""
    pillars = [Pillar(icon=PillarIcon(glyph="zap", color="#FF0000"), label="Speed")] * 3
    tokens = tokens_for()
    presentation = new_presentation(tokens)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    with pytest.raises(ValueError, match="#FF0000"):
        icon_pillars.render(slide, Canvas(tokens), IconPillarsContent("H", pillars))


def test_the_drawn_glyph_is_the_irs_glyph_id_not_its_concept(tmp_path: Path) -> None:
    """An icon block whose `concept` resolves to glyph A but whose `glyph_id` is B renders
    glyph B (compare the placed shape name to B)."""
    assert CONCEPT_TO_ICON["risk"] == "circle-alert"  # glyph A
    blocks = _blocks(3)
    blocks[1] = _icon_block("zap", concept="risk")  # glyph B recorded for concept "risk"

    out = tmp_path / "deck.pptx"
    render_deck(_deck(_slide(blocks)), tokens=tokens_for(), out_path=out)

    names = {
        shape.name
        for shape in Presentation(str(out)).slides[0].shapes
        if shape.name.startswith(ICON_NAME_PREFIX)
    }
    assert "icon:lucide:zap" in names
    assert "icon:lucide:circle-alert" not in names


def test_a_glyph_whose_file_name_is_also_a_concept_still_draws_the_named_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The renderer loads the glyph file directly. If a concept were created with the same
    name as a vendored file but mapped to a DIFFERENT glyph, the named file still wins;
    going through the concept table would draw `circle-alert` here."""
    monkeypatch.setitem(CONCEPT_TO_ICON, "zap", "circle-alert")
    assert CONCEPT_TO_ICON["zap"] != "zap"  # the collision exists

    blocks = _blocks(3, glyphs=("zap", "users", "clock"))
    out = tmp_path / "deck.pptx"
    render_deck(_deck(_slide(blocks)), tokens=tokens_for(), out_path=out)

    names = {
        shape.name
        for shape in Presentation(str(out)).slides[0].shapes
        if shape.name.startswith(ICON_NAME_PREFIX)
    }
    assert "icon:lucide:zap" in names
    assert "icon:lucide:circle-alert" not in names


def test_an_unknown_glyph_file_is_refused_not_guessed(tmp_path: Path) -> None:
    pillar = Pillar(icon=PillarIcon(glyph="no-such-glyph"), label="Speed")
    content = IconPillarsContent(headline="Unknown glyph", pillars=[pillar] * 3)
    with pytest.raises(IconNotFoundError):
        _render_content(content, tmp_path)


def test_panels_hug_the_tallest_stack_and_the_row_is_centred_under_the_headline(
    tmp_path: Path,
) -> None:
    """One panel per pillar, all the same height - the tallest pillar's stack plus padding
    top and bottom - and the row is vertically centred in the region under the headline.
    Each pillar's icon starts a fixed padding below its panel's top edge."""
    tokens = tokens_for()
    canvas = Canvas(tokens)
    body, _ = canvas.body_and_caption()
    content = _example(4)
    slide = _render_content(content, tmp_path).slides[0]

    panels = sorted(
        (s for s in slide.shapes if s.name.startswith("Rect")), key=lambda s: s.left
    )
    assert len(panels) == 4
    assert len({round(p.height.pt, 1) for p in panels}) == 1
    assert len({round(p.top.pt, 1) for p in panels}) == 1

    # Expected height: tallest stack + 2 x padding, measured independently of the renderer.
    frame = canvas.on(
        new_presentation(tokens).slides.add_slide(new_presentation(tokens).slide_layouts[6])
    )
    inner_width = panels[0].width.pt - 2 * icon_pillars.PANEL_PADDING
    heights = [
        icon_pillars._pillar_stack(frame, canvas, i, pillar, inner_width).height
        for i, pillar in enumerate(content.pillars)
    ]
    expected = max(heights) + 2 * icon_pillars.PANEL_PADDING
    assert panels[0].height.pt == pytest.approx(expected, abs=0.5)
    assert panels[0].height.pt < body.height / 2  # hugs content; not a full-height slab

    # Centred in the region between the headline's bottom and the caption band's top.
    headline = frame.stack("h", body.width)
    headline.text(content.headline, canvas.style("title", face="major", bold=True))
    region = headline.place(body, gutter=canvas.baseline * 4)
    row_centre = panels[0].top.pt + panels[0].height.pt / 2
    assert row_centre == pytest.approx(region.y + region.height / 2, abs=0.5)

    for panel, (left, top, width, _height) in zip(panels, _icon_boxes(slide), strict=True):
        assert top == pytest.approx(panel.top.pt + icon_pillars.PANEL_PADDING, abs=0.5)
        assert left + width / 2 == pytest.approx(panel.left.pt + panel.width.pt / 2, abs=1.0)

    # The panel tint is a theme colour, never a literal.
    for panel in panels:
        assert panel._element.xpath(".//a:srgbClr") == []
        assert panel._element.xpath(".//a:schemeClr")


def _line_colours(slide: Any) -> list[tuple[str, str]]:
    """(icon shape name, scheme colour or literal) for each icon shape's outline."""
    found: list[tuple[str, str]] = []
    for shape in slide.shapes:
        if not shape.name.startswith(ICON_NAME_PREFIX):
            continue
        fill = shape._element.spPr.find(qn("a:ln")).find(qn("a:solidFill"))
        assert fill is not None, shape.name
        scheme = fill.find(qn("a:schemeClr"))
        assert scheme is not None, f"{shape.name} is not a theme-slot colour"
        assert fill.find(qn("a:srgbClr")) is None, shape.name
        found.append((shape.name, scheme.get("val")))
    return found


def test_icons_are_theme_recolourable(tmp_path: Path) -> None:
    """GATE 3 check 2, as far as code can check it: every icon shape's fill/line uses a
    scheme colour (`a:schemeClr`) equal to the pillar's `color`, never `a:srgbClr`."""
    colours = ("accent1", "accent2", "accent3", "accent4")
    pillars = [
        Pillar(icon=PillarIcon(glyph=glyph, color=colour), label=f"Pillar {_NAMES[i]}")
        for i, (glyph, colour) in enumerate(
            zip(("users", "shield-check", "zap", "trending-up"), colours, strict=True)
        )
    ]
    slide = _render_content(IconPillarsContent("Four colours", pillars), tmp_path).slides[0]

    found = _line_colours(slide)
    assert found
    by_glyph = {name.rsplit(":", 1)[-1]: colour for name, colour in found}
    assert by_glyph == dict(
        zip(("users", "shield-check", "zap", "trending-up"), colours, strict=True)
    )

    # No literal colour anywhere in an icon's XML.
    for shape in slide.shapes:
        if shape.name.startswith(ICON_NAME_PREFIX):
            assert shape._element.xpath(".//a:srgbClr") == []


def test_adapter_pairs_by_position_and_refuses_ambiguous_points() -> None:
    """3 icons + 3 labels + 0 points → ok; + 3 points → ok; + 2 points → RenderStageError;
    3 icons + 2 labels → RenderStageError; a non-icon block in `pillar_icon` →
    RenderStageError naming it."""
    content_type = catalog.registration("icon_pillars").content_type

    content, warnings = adapt_slide(_slide(_blocks(3)), content_type)
    assert isinstance(content, IconPillarsContent)
    assert [p.point for p in content.pillars] == [None, None, None]
    assert [p.label for p in content.pillars] == [
        "Pillar Alpha",
        "Pillar Bravo",
        "Pillar Charlie",
    ]
    assert [p.icon.glyph for p in content.pillars] == ["users", "shield-check", "zap"]
    assert warnings == []

    content, _ = adapt_slide(_slide(_blocks(3, points=3)), content_type)
    assert isinstance(content, IconPillarsContent)
    assert [p.point for p in content.pillars] == ["Point Alpha", "Point Bravo", "Point Charlie"]

    with pytest.raises(RenderStageError, match="pillar_point"):
        adapt_slide(_slide(_blocks(3, points=2)), content_type)

    unbalanced = [b for b in _blocks(3) if b.text != "Pillar Charlie"]
    with pytest.raises(RenderStageError, match=r"3 .*2 "):
        adapt_slide(_slide(unbalanced), content_type)

    wrong_kind = _blocks(3)
    wrong_kind[1] = Block(id="not-an-icon", kind="framing", slot="pillar_icon", text="oops")
    with pytest.raises(RenderStageError, match="not-an-icon"):
        adapt_slide(_slide(wrong_kind), content_type)


def test_adapter_carries_the_irs_color_token_and_the_slides_accent_and_source() -> None:
    content_type = catalog.registration("icon_pillars").content_type
    blocks = _blocks(3)
    blocks[1] = _icon_block("users", color="accent3")
    blocks.extend(_claim("pillar_point", f"Measured gain {_NAMES[i]}.") for i in range(3))

    content, _ = adapt_slide(_slide(blocks, accent="accent2"), content_type)

    assert isinstance(content, IconPillarsContent)
    assert content.pillars[0].icon.color == "accent3"
    assert content.accent == "accent2"
    assert content.source.startswith("Source: programme-review")


def test_style_the_component_cannot_express_is_reported() -> None:
    content_type = catalog.registration("icon_pillars").content_type
    blocks = _blocks(3)
    _, warnings = adapt_slide(
        _slide(blocks, column_balance="lead_left", emphasis_block_id=blocks[1].id),
        content_type,
    )
    assert {w.field for w in warnings} == {"column_balance", "emphasis_block_id"}


def test_placeable_slots_and_catalog_registration() -> None:
    """`placeable_slots("icon_pillars") == {"headline", "pillar_icon", "pillar_label",
    "pillar_point"}`; `catalog.registration("icon_pillars")` exists with a preview path;
    `COMPONENT_LIB_VERSION` was bumped with a changelog line."""
    assert placeable_slots("icon_pillars") == {
        "headline",
        "pillar_icon",
        "pillar_label",
        "pillar_point",
    }

    entry = catalog.registration("icon_pillars")
    assert entry.preview == catalog.PREVIEW_DIR / "icon_pillars.png"
    assert entry.narrative_roles == ("capability pillars", "parallel ideas")

    major, minor, _ = (int(part) for part in catalog.COMPONENT_LIB_VERSION.split("."))
    assert (major, minor) >= (0, 6)
    source = Path(catalog.__file__).read_text(encoding="utf-8")
    assert "#: 0.6.0:" in source

    # Text slots are budgeted; the icon is not text and gets no ComponentSlot.
    names = {slot.name for slot in catalog.spec_for("icon_pillars", tokens_for()).slots}
    assert names == {"headline", "pillar_label", "pillar_point", "source"}
    spec = catalog.spec_for("icon_pillars", tokens_for())
    assert spec.slot("pillar_point").required is False
    assert spec.slot("pillar_label").repeatable and spec.slot("pillar_point").repeatable


def test_the_slot_budget_rejects_a_fifth_pillar_label() -> None:
    labels = [f"Pillar {name}" for name in _NAMES[: icon_pillars.MAX_PILLARS + 1]]
    findings = catalog.check_overflow(
        {"headline": "Too many", "pillar_label": labels}, "icon_pillars", tokens_for()
    )
    assert findings, "a fifth pillar must be reported before render"


def test_grammar_icon_adjacency_passes_on_a_rendered_pillar_slide(tmp_path: Path) -> None:
    """`grammar.check_icon_adjacency` returns no findings for a rendered icon_pillars slide."""
    for count in (3, 4):
        presentation = _render_content(_example(count), tmp_path)
        assert grammar.check_icon_adjacency(presentation, tokens_for()) == []


def test_aesthetic_icon_actions_reach_a_rendered_face_icon(tmp_path: Path) -> None:
    """A deck with an icon_pillars slide: `SwapGlyph` and `SetIconColour` via `apply_action`
    then `render_deck` → the new glyph/colour is what the PPTX holds; fact fingerprint
    unchanged. This is the end-to-end path B35 says was missing."""
    blocks = _blocks(3, glyphs=("users", "shield-check", "zap"))
    target = blocks[1]  # the first pillar's icon: "users"
    deck = _deck(_slide(blocks))

    swapped = apply_action(
        deck,
        SwapGlyph(slide_id="s1", block_id=target.id, concept="growth"),
        slots_of=catalog_slot_lookup(),
        glyph_for=library_concept_lookup(),
    )
    recoloured = apply_action(
        swapped,
        SetIconColour(slide_id="s1", block_id=target.id, color_token="accent4"),
        slots_of=catalog_slot_lookup(),
        glyph_for=library_concept_lookup(),
    )
    assert fact_fingerprint(recoloured) == fact_fingerprint(deck)

    before, after = tmp_path / "before.pptx", tmp_path / "after.pptx"
    render_deck(deck, tokens=tokens_for(), out_path=before)
    render_deck(recoloured, tokens=tokens_for(), out_path=after)

    def icons(path: Path) -> dict[str, str]:
        found = _line_colours(Presentation(str(path)).slides[0])
        return {name.rsplit(":", 1)[-1]: colour for name, colour in found}

    assert icons(before)["users"] == "accent1"
    drawn = icons(after)
    assert "users" not in drawn
    assert drawn["trending-up"] == "accent4"  # the new glyph, in the new colour
    assert drawn["shield-check"] == "accent1"  # the untouched pillars did not move
    assert drawn["zap"] == "accent1"
