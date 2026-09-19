"""A6 — the build manifest, and what "the same output" is allowed to mean.

The Phase 0 manifest tests live in `test_orchestrator.py` alongside the run state they were
written with. This file is the Phase 2b surface: the knowledge commit and its dirty flag,
and the PPTX normalisation.

The normalisation tests come in pairs on purpose. For each thing the canonical digest
ignores there is a test that varying it does **not** change the digest, and a neighbouring
test that something real in the same place *does*:

- entry order, timestamps and compression vary freely — but a removed part does not;
- `dcterms:modified` varies freely — but `dc:title`, in the same file, does not.

An exclusion is a hole in A6 unless it can be shown not to swallow a real difference, and a
one-sided test proves only that the digest is insensitive, which a constant would also be.

`test_naive_byte_comparison_fails_on_two_identical_decks` is the one that matters for the
invariant's wording: two renders of the same deck are not the same bytes, and no amount of
care makes them so.
"""

from __future__ import annotations

import shutil
import subprocess
import zipfile
from datetime import datetime
from pathlib import Path

import pytest
from pptx import Presentation

from autodeck.audit import manifest
from autodeck.audit.manifest import (
    PPTX_NORMALISATION,
    PPTX_NORMALISATION_RULES,
    GitState,
    Manifest,
    build_manifest,
    canonical_pptx_digest,
    git_state,
    normalise_part,
    pptx_differences,
)
from autodeck.design.components.catalog import COMPONENT_LIB_VERSION

needs_git = pytest.mark.skipif(shutil.which("git") is None, reason="git is not installed")


# ---------------------------------------------------------------------------
# Fixtures: a deck, and the ways one archive can differ from another
# ---------------------------------------------------------------------------


def make_pptx(
    path: Path,
    *,
    headline: str = "Throughput rose 60 percent after the migration",
    title: str = "",
    modified: datetime | None = None,
) -> Path:
    """A one-slide PPTX built directly with python-pptx.

    The renderer lands in Phase 3b, so the normaliser is tested against real PowerPoint
    packages built here rather than against an imagined one. `headline` stands in for deck
    content; `title` and `modified` are the two `docProps` cases the rules distinguish.
    """
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    assert slide.shapes.title is not None
    slide.shapes.title.text = headline
    if title:
        presentation.core_properties.title = title
    if modified is not None:
        presentation.core_properties.modified = modified
    presentation.save(str(path))
    return path


def rewrite_zip(
    source: Path,
    target: Path,
    *,
    reverse: bool = False,
    date_time: tuple[int, int, int, int, int, int] = (1980, 1, 1, 0, 0, 0),
    compression: int = zipfile.ZIP_DEFLATED,
    drop: str | None = None,
) -> Path:
    """Re-pack an archive with the same parts stored differently.

    Rewriting rather than sleeping between two saves: the difference under test is then
    exactly the one named, and the test does not depend on how long anything took.
    """
    with zipfile.ZipFile(source) as archive:
        names = [n for n in archive.namelist() if n != drop]
        parts = {name: archive.read(name) for name in names}
    with zipfile.ZipFile(target, "w", compression=compression) as out:
        for name in reversed(names) if reverse else names:
            info = zipfile.ZipInfo(name, date_time=date_time)
            info.compress_type = compression
            out.writestr(info, parts[name])
    return target


# ---------------------------------------------------------------------------
# The claim A6 makes, and the claim that is actually available
# ---------------------------------------------------------------------------


def test_two_renders_of_the_same_ir_are_normalised_identical(tmp_path: Path) -> None:
    """The phase brief's done-when, stated in the terms that are achievable."""
    first = make_pptx(tmp_path / "first.pptx")
    second = make_pptx(tmp_path / "second.pptx")

    assert canonical_pptx_digest(first) == canonical_pptx_digest(second)
    assert pptx_differences(first, second) == []


def test_a_changed_ir_does_not_compare_equal(tmp_path: Path) -> None:
    """The other half: a digest insensitive to everything would pass the test above too."""
    first = make_pptx(tmp_path / "first.pptx")
    changed = make_pptx(tmp_path / "changed.pptx", headline="Throughput rose 80 percent")

    assert canonical_pptx_digest(first) != canonical_pptx_digest(changed)
    assert any("slide1.xml" in difference for difference in pptx_differences(first, changed))


def test_naive_byte_comparison_fails_on_two_identical_decks(tmp_path: Path) -> None:
    """Why A6's "byte-comparable" is not a claim this system can make.

    The two files here hold precisely the same parts with precisely the same bytes. They
    differ only in the wall-clock time stamped into each zip entry header — which `zipfile`
    writes on every save and nothing in the deck can influence.
    """
    original = make_pptx(tmp_path / "original.pptx")
    later = rewrite_zip(original, tmp_path / "later.pptx", date_time=(2031, 5, 4, 12, 0, 0))

    assert original.read_bytes() != later.read_bytes(), "naive byte comparison"
    assert canonical_pptx_digest(original) == canonical_pptx_digest(later)


# ---------------------------------------------------------------------------
# Each exclusion, and the real difference in the same place that survives it
# ---------------------------------------------------------------------------


def test_entry_order_does_not_change_the_digest(tmp_path: Path) -> None:
    """Nothing resolves a part by position; order cannot carry content."""
    original = make_pptx(tmp_path / "original.pptx")
    shuffled = rewrite_zip(original, tmp_path / "shuffled.pptx", reverse=True)

    assert canonical_pptx_digest(original) == canonical_pptx_digest(shuffled)


def test_compression_does_not_change_the_digest(tmp_path: Path) -> None:
    """The DEFLATE stream is a property of the zlib build, not of the deck."""
    original = make_pptx(tmp_path / "original.pptx")
    stored = rewrite_zip(original, tmp_path / "stored.pptx", compression=zipfile.ZIP_STORED)

    assert canonical_pptx_digest(original) == canonical_pptx_digest(stored)


def test_a_removed_part_changes_the_digest(tmp_path: Path) -> None:
    """Order, timestamps and compression are dropped; the set of parts is not."""
    original = make_pptx(tmp_path / "original.pptx")
    without = rewrite_zip(original, tmp_path / "without.pptx", drop="docProps/app.xml")

    assert canonical_pptx_digest(original) != canonical_pptx_digest(without)
    assert any("docProps/app.xml" in d for d in pptx_differences(original, without))


def test_a_volatile_docprops_field_does_not_change_the_digest(tmp_path: Path) -> None:
    """`dcterms:modified` moves on every save and appears nowhere in the deck."""
    first = make_pptx(tmp_path / "first.pptx", modified=datetime(2024, 1, 1, 9, 0, 0))
    second = make_pptx(tmp_path / "second.pptx", modified=datetime(2031, 7, 9, 17, 30, 0))

    assert first.read_bytes() != second.read_bytes()
    assert canonical_pptx_digest(first) == canonical_pptx_digest(second)


def test_an_authored_docprops_field_does_change_the_digest(tmp_path: Path) -> None:
    """The scope test for the rule above: `dc:title` sits in the same file and is content."""
    first = make_pptx(tmp_path / "first.pptx", title="Migration review")
    second = make_pptx(tmp_path / "second.pptx", title="Migration review — final")

    assert canonical_pptx_digest(first) != canonical_pptx_digest(second)
    assert pptx_differences(first, second) == ["differs: docProps/core.xml"]


def test_normalisation_touches_only_the_core_properties_part() -> None:
    """A normalisation nobody can look at is one nobody can check the scope of."""
    core = (
        b"<cp:coreProperties><dc:title>Kept</dc:title>"
        b"<dcterms:modified>2026-01-01T00:00:00Z</dcterms:modified></cp:coreProperties>"
    )
    assert b"2026-01-01" not in normalise_part("docProps/core.xml", core)
    assert b"Kept" in normalise_part("docProps/core.xml", core)
    assert normalise_part("ppt/slides/slide1.xml", core) == core


def test_the_name_content_boundary_cannot_be_moved(tmp_path: Path) -> None:
    """Framing in the digest, so two different archives cannot hash the same.

    Without the length prefix, an archive of `{"a": b"bc"}` and one of `{"ab": b"c"}` would
    feed the hash identical bytes.
    """

    def archive(path: Path, entries: dict[str, bytes]) -> Path:
        with zipfile.ZipFile(path, "w") as out:
            for name, data in entries.items():
                out.writestr(name, data)
        return path

    left = archive(tmp_path / "left.zip", {"a": b"bc"})
    right = archive(tmp_path / "right.zip", {"ab": b"c"})

    assert canonical_pptx_digest(left) != canonical_pptx_digest(right)


# ---------------------------------------------------------------------------
# The normalisation is recorded, not assumed
# ---------------------------------------------------------------------------


def test_the_manifest_records_which_normalisation_was_applied() -> None:
    """A comparison next year has to know what was excluded from "the same output"."""
    manifest = build_manifest(run_id="r1", env="dev")

    assert manifest.render_normalisation == PPTX_NORMALISATION
    assert manifest.render_normalisation_excludes == [r.id for r in PPTX_NORMALISATION_RULES]
    assert "docprops-volatile-fields" in manifest.render_normalisation_excludes


def test_a_build_with_no_render_records_no_normalisation() -> None:
    """Better to record nothing than to describe a comparison that was never available."""
    manifest = build_manifest(run_id="r1", env="dev", render_normalisation=None)

    assert manifest.render_normalisation is None
    assert manifest.render_normalisation_excludes == []


def test_a_different_normalisation_breaks_the_match() -> None:
    """Two digests computed under different rules are not comparable, so nor are the builds."""
    first = build_manifest(run_id="r1", env="dev")
    second = first.model_copy(update={"render_normalisation": "pptx-canonical-v2"})

    assert not first.matches(second)


def test_every_normalisation_rule_argues_that_its_exclusion_is_safe() -> None:
    """The table's own self-check, in the spirit of the numeric linter's collision check.

    A rule with an empty `safe_because` is an exclusion nobody has justified, and an
    exclusion nobody can argue with is one that survives the comparison it should have
    failed.
    """
    ids = [rule.id for rule in PPTX_NORMALISATION_RULES]
    assert len(set(ids)) == len(ids)
    for rule in PPTX_NORMALISATION_RULES:
        assert rule.what.strip()
        assert rule.why.strip()
        assert rule.safe_because.strip()


def test_the_digest_is_bound_to_the_normalisation_it_used(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The rule version is hashed in, so a v1 digest can never be read as a v2 one.

    Two digests computed under different rules answer different questions. If they could
    collide, a future comparison would report "identical" across a change to what identical
    means — the one failure a recorded normalisation exists to prevent.
    """
    deck = make_pptx(tmp_path / "deck.pptx")
    under_v1 = canonical_pptx_digest(deck)

    monkeypatch.setattr(manifest, "PPTX_NORMALISATION", "pptx-canonical-v2")
    assert canonical_pptx_digest(deck) != under_v1


# ---------------------------------------------------------------------------
# Knowledge commit, and whether the tree was dirty
# ---------------------------------------------------------------------------


def git(repo: Path, *args: str) -> None:
    subprocess.run(
        ["git", "-C", str(repo), "-c", "user.email=t@e", "-c", "user.name=t", *args],
        check=True,
        capture_output=True,
    )


@pytest.fixture
def knowledge_repo(tmp_path: Path) -> Path:
    """A git repo with a `knowledge/` folder and an unrelated folder beside it."""
    repo = tmp_path / "repo"
    (repo / "knowledge").mkdir(parents=True)
    (repo / "code").mkdir()
    (repo / "knowledge" / "claims.md").write_text("A cited claim.\n", encoding="utf-8")
    (repo / "code" / "module.py").write_text("x = 1\n", encoding="utf-8")
    git(repo, "init", "-q")
    git(repo, "add", "-A")
    git(repo, "commit", "-q", "-m", "seed")
    return repo


@needs_git
def test_a_clean_knowledge_tree_records_its_commit(knowledge_repo: Path) -> None:
    state = git_state(knowledge_repo / "knowledge")

    assert state.commit is not None and len(state.commit) == 40
    assert state.dirty is False
    assert state.reproducible


@needs_git
def test_an_edited_knowledge_file_makes_the_build_dirty(knowledge_repo: Path) -> None:
    """A deck built against "whatever claims.md said that day" is not reproducible."""
    (knowledge_repo / "knowledge" / "claims.md").write_text("Edited.\n", encoding="utf-8")
    state = git_state(knowledge_repo / "knowledge")

    assert state.dirty is True
    assert not state.reproducible


@needs_git
def test_an_untracked_knowledge_file_counts_as_dirty(knowledge_repo: Path) -> None:
    """A new, uncommitted source is exactly an input the commit does not describe."""
    (knowledge_repo / "knowledge" / "new.md").write_text("New source.\n", encoding="utf-8")

    assert git_state(knowledge_repo / "knowledge").dirty is True


@needs_git
def test_changes_outside_the_knowledge_folder_do_not_make_it_dirty(
    knowledge_repo: Path,
) -> None:
    """The scoping decision, asserted: an edit under `code/` does not change what a source says.

    A flag that went True on every development build would be read as noise within a week.
    """
    (knowledge_repo / "code" / "module.py").write_text("x = 2\n", encoding="utf-8")
    state = git_state(knowledge_repo / "knowledge")

    assert state.dirty is False
    assert state.commit is not None


@needs_git
def test_a_folder_outside_git_records_neither_commit_nor_flag(tmp_path: Path) -> None:
    """None is not False: "there was no tree to ask" is a different finding from "clean"."""
    loose = tmp_path / "loose"
    loose.mkdir()
    state = git_state(loose)

    assert state == GitState(commit=None, dirty=None)
    assert not state.reproducible


@needs_git
def test_the_manifest_carries_the_commit_and_the_dirty_flag(knowledge_repo: Path) -> None:
    manifest = build_manifest(
        run_id="r1", env="dev", knowledge_dir=knowledge_repo / "knowledge"
    )

    assert manifest.knowledge_commit is not None
    assert manifest.knowledge_dirty is False
    assert manifest.reproducible()


@needs_git
def test_a_dirty_build_says_so_rather_than_implying_reproducibility(
    knowledge_repo: Path,
) -> None:
    (knowledge_repo / "knowledge" / "claims.md").write_text("Edited.\n", encoding="utf-8")
    manifest = build_manifest(
        run_id="r1", env="dev", knowledge_dir=knowledge_repo / "knowledge"
    )

    assert manifest.knowledge_dirty is True
    assert not manifest.reproducible()


def test_a_build_with_no_knowledge_folder_claims_nothing() -> None:
    manifest = build_manifest(run_id="r1", env="dev")

    assert manifest.knowledge_commit is None
    assert manifest.knowledge_dirty is None
    assert not manifest.reproducible()


def test_a_dirty_build_and_a_clean_one_are_not_the_same_build() -> None:
    """The flag is in the reproducibility key; nothing else would show the difference."""
    clean = build_manifest(run_id="r1", env="dev").model_copy(
        update={"knowledge_commit": "a" * 40, "knowledge_dirty": False}
    )
    dirty = clean.model_copy(update={"knowledge_dirty": True})

    assert not clean.matches(dirty)


def test_the_manifest_round_trips_the_new_fields(tmp_path: Path) -> None:
    manifest = build_manifest(run_id="r1", env="dev").model_copy(
        update={"knowledge_commit": "b" * 40, "knowledge_dirty": True}
    )
    assert Manifest.load(manifest.save(tmp_path / "build_manifest.json")) == manifest


# ---------------------------------------------------------------------------
# Component library version
# ---------------------------------------------------------------------------


def test_the_component_library_version_has_one_definition() -> None:
    """The manifest and the `Deck` read the same constant, so they cannot drift apart.

    The old default was a literal `"0.0.0"` here against a literal `"0.1.0"` in the CLI —
    two answers to one question, and the manifest's was the one nobody would have noticed
    had gone stale.
    """
    assert build_manifest(run_id="r1", env="dev").component_lib_version == COMPONENT_LIB_VERSION
    assert (
        Manifest(run_id="r1", env="dev", created_at="now").component_lib_version
        == COMPONENT_LIB_VERSION
    )


def test_a_different_component_library_is_a_different_build() -> None:
    """Different slot geometry is a different deck, even from identical IR."""
    first = build_manifest(run_id="r1", env="dev")
    second = build_manifest(run_id="r1", env="dev", component_lib_version="0.2.0")

    assert not first.matches(second)
