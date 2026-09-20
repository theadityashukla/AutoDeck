"""`autodeck.design.draw`'s two new primitives (task 3a.6b): an arbitrary autoshape and a
connector with an arrowhead — the two gaps 3a.6a's author flagged that `diagrams.py` hits
immediately.

Both are checked the way 3a.7 checked icon shapes: real python-pptx objects, inspected for
a `schemeClr` theme reference rather than a baked RGB (D1), and for the raw `<a:ln>` XML an
arrowhead actually needs, since python-pptx has no higher-level accessor for either.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.slide import Slide

from autodeck.design.draw import add_autoshape, add_connector, add_rect
from autodeck.design.layout_kit import Box

TOKENS_DIR = Path(__file__).resolve().parents[1] / "config" / "tokens"


def _slide() -> Slide:
    presentation = Presentation()
    return presentation.slides.add_slide(presentation.slide_layouts[6])


# ---------------------------------------------------------------------------
# add_autoshape
# ---------------------------------------------------------------------------


def test_add_autoshape_builds_the_named_preset() -> None:
    shape = add_autoshape(_slide(), Box(10, 10, 100, 40), "CHEVRON", fill="accent2")
    assert shape.auto_shape_type == MSO_SHAPE.CHEVRON


def test_add_autoshape_fill_is_a_theme_reference_not_an_rgb() -> None:
    """D1: recolouring the theme must move this shape. A baked RGB would not."""
    shape = add_autoshape(_slide(), Box(0, 0, 60, 60), "OVAL", fill="accent3")
    assert shape.fill.fore_color.theme_color is not None
    assert "schemeClr" in shape.fill.fore_color._xFill.xml
    assert "accent3" in shape.fill.fore_color._xFill.xml


def test_add_autoshape_line_is_also_a_theme_reference() -> None:
    shape = add_autoshape(_slide(), Box(0, 0, 60, 60), "TRAPEZOID", line="dk2", line_width=2.0)
    assert shape.line.color.theme_color is not None


def test_add_autoshape_rejects_an_unknown_shape_name() -> None:
    with pytest.raises(ValueError, match="unknown MSO_SHAPE member"):
        add_autoshape(_slide(), Box(0, 0, 10, 10), "NOT_A_SHAPE")


def test_add_autoshape_no_fill_leaves_the_shape_unfilled() -> None:
    from pptx.enum.dml import MSO_FILL_TYPE

    shape = add_autoshape(_slide(), Box(0, 0, 10, 10), "RECTANGLE")
    assert shape.fill.type == MSO_FILL_TYPE.BACKGROUND


def test_add_rect_is_a_rectangle_and_still_works() -> None:
    """`add_rect` now delegates to `add_autoshape`; its own behaviour is unchanged."""
    shape = add_rect(_slide(), Box(0, 0, 40, 40), fill="accent1")
    assert shape.auto_shape_type == MSO_SHAPE.RECTANGLE
    assert "accent1" in shape.fill.fore_color._xFill.xml


def test_each_autoshape_is_its_own_shape_on_the_slide() -> None:
    """The individual-selectability requirement (brief's done-when): N calls, N `<p:sp>`s,
    never one grouped drawing."""
    slide = _slide()
    for _ in range(3):
        add_autoshape(slide, Box(0, 0, 20, 20), "OVAL", fill="accent1")
    assert len(slide.shapes) == 3
    assert len({id(shape._element) for shape in slide.shapes}) == 3


# ---------------------------------------------------------------------------
# add_connector
# ---------------------------------------------------------------------------


def test_add_connector_is_a_straight_connector_between_the_given_points() -> None:
    slide = _slide()
    connector = add_connector(slide, (0.0, 0.0), (100.0, 50.0))
    assert connector.shape_type is not None  # a real shape landed on the slide
    from autodeck.design.theme.tokens import points_to_emu

    assert connector.begin_x == points_to_emu(0.0)
    assert connector.begin_y == points_to_emu(0.0)
    assert connector.end_x == points_to_emu(100.0)
    assert connector.end_y == points_to_emu(50.0)


def test_add_connector_line_colour_is_a_theme_reference() -> None:
    connector = add_connector(_slide(), (0, 0), (10, 10), color="accent4")
    assert connector.line.color.theme_color is not None


def test_add_connector_default_arrow_is_at_the_end_only() -> None:
    connector = add_connector(_slide(), (0, 0), (10, 0))
    xml = connector.line._get_or_add_ln().xml
    assert "<a:tailEnd" in xml
    assert "<a:headEnd" not in xml


def test_add_connector_both_ends_can_carry_an_arrowhead() -> None:
    connector = add_connector(_slide(), (0, 0), (10, 0), arrow="both")
    xml = connector.line._get_or_add_ln().xml
    assert "<a:headEnd" in xml
    assert "<a:tailEnd" in xml
    # Schema order: headEnd precedes tailEnd inside <a:ln>.
    assert xml.index("<a:headEnd") < xml.index("<a:tailEnd")


def test_add_connector_none_draws_no_arrowhead_element() -> None:
    connector = add_connector(_slide(), (0, 0), (10, 0), arrow="none")
    xml = connector.line._get_or_add_ln().xml
    assert "End" not in xml


def test_add_connector_rejects_an_unknown_arrow_style() -> None:
    with pytest.raises(ValueError, match="unknown arrow style"):
        add_connector(_slide(), (0, 0), (10, 0), arrow="triangle")  # type: ignore[arg-type]


def test_add_connector_is_the_straight_connector_type() -> None:
    slide = _slide()
    add_connector(slide, (0, 0), (10, 10))
    element = slide.shapes[0]._element
    # The generic connector-type check python-pptx itself uses.
    assert element.tag.endswith("}cxnSp")


def test_connectors_are_individually_selectable_shapes_too() -> None:
    slide = _slide()
    add_connector(slide, (0, 0), (10, 0))
    add_connector(slide, (0, 10), (10, 10))
    assert len(slide.shapes) == 2
