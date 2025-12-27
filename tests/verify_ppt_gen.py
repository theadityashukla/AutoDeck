import json
import os
import sys

# Add project root to path
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from autodeck_core.ppt_generator import PPTGenerator
from pptx import Presentation

def test_ppt_generation():
    session_file = "sessions/7851b7b8.json"
    if not os.path.exists(session_file):
        print(f"Session file {session_file} not found.")
        return

    with open(session_file, "r") as f:
        session_data = json.load(f)

    print(f"Loaded session: {session_data.get('name')}")

    output_path = "test_presentation.pptx"
    ppt_gen = PPTGenerator(output_path=output_path)
    
    try:
        generated_file = ppt_gen.generate(session_data)
        print(f"Successfully generated PPT at: {generated_file}")
        
        if os.path.exists(generated_file):
            size = os.path.getsize(generated_file)
            print(f"File size: {size} bytes")
            if size > 0:
                print("Verification passed!")
            else:
                print("Verification failed: File is empty.")
                
            # Deep verification
            prs = Presentation(generated_file)
            # Check Slide 0 (assuming no content in session for slide 0)
            if len(prs.slides) > 1:
                slide_0 = prs.slides[1] # Slide 0 is usually index 1 if title slide is 0? No, create_slide is called in loop.
                # Title slide is created before loop. So prs.slides[0] is Title Slide.
                # Loop starts at index 0. So prs.slides[1] corresponds to session outline[0].
                
                text = ""
                for shape in slide_0.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text
                
                print(f"Slide 1 Body Text: {text}")
                if "[Draft]" in text:
                    print("Draft marker found as expected.")
                else:
                    print("WARNING: Draft marker NOT found in slide 1.")

        else:
             print("Verification failed: File not found.")

    except Exception as e:
        print(f"Verification failed with error: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    test_ppt_generation()
