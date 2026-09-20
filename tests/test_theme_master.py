"""Mode (b) — generate from tokens: theme, package, and master inheritance (3a.1, D1).

`test_design.py` already covers `build_theme_xml` and `apply_theme` as pure XML generation.
What is missing, and what this file adds, is proof of the actual done-when: *"new hand-added
slides inherit it."* A theme part that looks right in isolation is not the same claim as a
slide a person adds by hand, through `Presentation.slides.add_slide`, picking up the tokens'
colours and fonts without a single override anywhere in the file.

Two kinds of evidence:

- **Structural** (no `render` marker, runs in CI): the master, its layouts, and a hand-added
  slide contain zero literal `<a:srgbClr>`/`<a:latin typeface=...>` overrides — every colour
  and font reference is scheme-indirected (`schemeClr`, `+mj-lt`/`+mn-lt`). That is what makes
  inheritance possible at all; python-pptx's own default template already has this property,
  and these tests pin it as a regression guard.
- **Rendered** (`render` marker, local-only per B10): the structural claim only matters if a
  renderer actually resolves the indirection. `check_against_render`'s pattern is followed —
  `soffice`, a real font, `pdftoppm`.

**What LibreOffice could and could not confirm here (see B31) — read before trusting a green
render-marked test as proof of anything about PowerPoint:**

- A shape's fill colour set via a theme colour reference (`accent1`-`accent6`) renders as
  that exact RGB. Confirmed here.
- A hand-added slide's title text renders in the token set's declared major font family,
  distinguishably from a different family. Confirmed here (by ink-width, not glyph shape).
- **Not confirmed, and not testable here:** `dk1`/`lt1` — the text/background pair — are
  encoded as `<a:sysClr val="windowText"/>`/`<a:sysClr val="window"/>` with a `lastClr`
  fallback (see `master_builder._colour_value`'s docstring for why: PowerPoint's own Design
  panel pairs `dk1`/`lt1` in the wrong order without it). LibreOffice was observed, while
  writing this task, to render `sysClr`'s literal system colour (plain black/white) rather
  than honouring `lastClr` — a title placeholder's text stayed black in every render tried
  here regardless of what `dk1` was set to, while switching the same slot to a plain
  `srgbClr` made the custom colour appear immediately. That means **this render loop cannot
  confirm that a client's chosen text/background colours actually reach hand-added slide
  text** — only that accent-coloured shapes and fonts do. This is one of the two
  PowerPoint-only checks carried forward to the owner at GATE 3 (see the handover note); do
  not read a passing test here as covering it.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from autodeck.design.fonts import is_available
from autodeck.design.theme.master_builder import new_presentation, save_themed
from autodeck.design.theme.tokens import DesignTokens, Palette, Typography

TOKENS_DIR = Path(__file__).resolve().parents[1] / "config" / "tokens"

needs_soffice = pytest.mark.skipif(
    __import__("shutil").which("soffice") is None
    and __import__("shutil").which("libreoffice") is None,
    reason="LibreOffice not installed",
)


def _needs_family(*families: str) -> pytest.MarkDecorator:
    missing = [f for f in families if not is_available(f)]
    return pytest.mark.skipif(bool(missing), reason=f"{missing} not installed")


def _themed_deck_with_hand_added_slide(tokens: DesignTokens, out: Path, title: str) -> Path:
    """Build a themed deck, then add a slide the way a person would in PowerPoint.

    `Presentation.slides.add_slide` from a stock layout — no formatting is touched — is the
    actual test of a master, per this task's spec: it is what separates a theme from a
    styled deck.
    """
    from pptx import Presentation

    presentation = new_presentation(tokens)
    save_themed(presentation, tokens, out)

    reopened = Presentation(str(out))
    layout = reopened.slide_layouts[1]  # "Title and Content"
    slide = reopened.slides.add_slide(layout)
    assert slide.shapes.title is not None
    slide.shapes.title.text = title
    reopened.save(str(out))
    return out


# ---------------------------------------------------------------------------
# Structural: nothing in the file overrides the theme
# ---------------------------------------------------------------------------


def test_the_stock_master_and_layouts_carry_no_literal_colour_or_font() -> None:
    """Pins the property that makes inheritance possible at all.

    If python-pptx's bundled template ever starts hardcoding a colour or a font on a
    placeholder, theme substitution stops being enough and this whole approach needs
    rethinking — better to find that here than from a client's hand-added slide looking
    wrong.
    """
    import re
    import zipfile

    tokens = DesignTokens.load(TOKENS_DIR / "dev.json")
    presentation = new_presentation(tokens)
    import io

    buffer = io.BytesIO()
    presentation.save(buffer)

    with zipfile.ZipFile(buffer) as package:
        master_and_layout_parts = [
            name
            for name in package.namelist()
            if name.startswith(("ppt/slideMasters/", "ppt/slideLayouts/"))
            and name.endswith(".xml")
        ]
        assert master_and_layout_parts
        for name in master_and_layout_parts:
            xml = package.read(name).decode("utf-8")
            assert not re.search(r'<a:srgbClr val="[0-9A-Fa-f]{6}"', xml), name
            # `+mj-lt`/`+mn-lt`/`+mj-ea` etc. are theme references, not literal faces.
            for literal_font in re.findall(r'<a:latin typeface="([^"]*)"', xml):
                assert literal_font == "" or literal_font.startswith("+"), (name, literal_font)


def test_a_hand_added_slide_carries_no_direct_formatting(tmp_path: Path) -> None:
    """The actual done-when, checked structurally: nothing on the new slide's part overrides
    the theme it was added against."""
    import re
    import zipfile

    tokens = DesignTokens.load(TOKENS_DIR / "dev.json")
    out = _themed_deck_with_hand_added_slide(tokens, tmp_path / "deck.pptx", "Hand-added slide")

    slide_pattern = re.compile(r"ppt/slides/slide\d+\.xml")
    with zipfile.ZipFile(out) as package:
        slide_parts = [n for n in package.namelist() if slide_pattern.match(n)]
        assert len(slide_parts) == 1
        slide_xml = package.read(slide_parts[0]).decode("utf-8")

    assert not re.search(r'<a:srgbClr val="[0-9A-Fa-f]{6}"', slide_xml)
    assert not re.search(r'<a:latin typeface="(?!\+)[^"]', slide_xml)


def test_the_theme_part_still_carries_this_decks_own_tokens(tmp_path: Path) -> None:
    """Adding a slide must not disturb the one part the theme UI actually reads."""
    from autodeck.design.theme.master_builder import read_theme_xml

    tokens = DesignTokens.load(TOKENS_DIR / "dev.json")
    out = _themed_deck_with_hand_added_slide(tokens, tmp_path / "deck.pptx", "Hand-added slide")

    theme_xml = read_theme_xml(out)
    assert tokens.palette.accent1.upper() in theme_xml
    assert f'<a:latin typeface="{tokens.typography.major}"/>' in theme_xml


# ---------------------------------------------------------------------------
# Rendered: LibreOffice actually resolves the indirection (as far as it allows — see the
# module docstring for exactly where it does not)
# ---------------------------------------------------------------------------


@pytest.mark.render
@needs_soffice
@_needs_family("Inter", "Inter Display")
def test_a_hand_added_shapes_theme_colour_fill_renders_as_that_colour(tmp_path: Path) -> None:
    """A shape recoloured from the theme — the everyday case for a client's own edit."""
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu

    from autodeck.render.qa.libreoffice import render_pptx

    tokens = DesignTokens(
        name="Render check",
        palette=Palette(accent1="2E9B33"),
        typography=Typography(major="Inter Display", minor="Inter"),
    )
    presentation = new_presentation(tokens)
    out = tmp_path / "deck.pptx"
    save_themed(presentation, tokens, out)

    from pptx import Presentation

    reopened = Presentation(str(out))
    slide = reopened.slides.add_slide(reopened.slide_layouts[6])  # blank layout
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(457200), Emu(457200), Emu(2000000), Emu(1500000)
    )
    shape.fill.solid()
    shape.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
    reopened.save(str(out))

    result = render_pptx(out, tmp_path / "png", tokens)
    from PIL import Image

    image = Image.open(result.images[0]).convert("RGB")
    # Pillow's stub types getpixel as float | tuple | None; an "RGB" image always gives a
    # 3-tuple of ints.
    r, g, b = image.getpixel((image.width // 8, image.height // 8))  # type: ignore[misc]
    assert f"{r:02X}{g:02X}{b:02X}" == tokens.palette.accent1.upper()


@pytest.mark.render
@needs_soffice
@_needs_family("Inter", "Inter Display", "Liberation Serif")
def test_a_hand_added_titles_font_family_renders_distinguishably(tmp_path: Path) -> None:
    """Two token sets that differ only in typeface must not render identically.

    Ink width rather than glyph shape, because that is what a PNG comparison can check
    without an OCR step — but a sans/serif pair at the same size and string reliably differs
    by a wide margin, which is the point: this is not measuring anti-aliasing noise.
    """
    from autodeck.render.qa.libreoffice import render_pptx

    def ink_width(major: str, minor: str, subdir: str) -> int:
        import numpy as np
        from PIL import Image
        from pptx import Presentation

        typography = Typography(major=major, minor=minor)
        tokens = DesignTokens(name="Font check", typography=typography)
        out = tmp_path / subdir / "deck.pptx"
        presentation = new_presentation(tokens)
        save_themed(presentation, tokens, out)

        reopened = Presentation(str(out))
        slide = reopened.slides.add_slide(reopened.slide_layouts[1])
        assert slide.shapes.title is not None
        slide.shapes.title.text = "AutoDeck Hand-Added Slide"
        reopened.save(str(out))

        result = render_pptx(out, tmp_path / subdir / "png", tokens)
        image = Image.open(result.images[0]).convert("L")
        width, height = image.size
        # Crop to the title placeholder's own box (master's default position).
        x0, x1 = int(width * 457200 / 12192000), int(width * 8686800 / 12192000)
        y0, y1 = int(height * 274638 / 6858000), int(height * 1417638 / 6858000)
        crop = np.array(image.crop((x0, y0, x1, y1)))
        dark_columns = np.where((crop < 128).any(axis=0))[0]
        assert dark_columns.size, "expected some rendered ink in the title box"
        return int(dark_columns.max() - dark_columns.min())

    sans_width = ink_width("Inter Display", "Inter", "sans")
    serif_width = ink_width("Liberation Serif", "Liberation Serif", "serif")

    assert abs(sans_width - serif_width) > 20, (sans_width, serif_width)
