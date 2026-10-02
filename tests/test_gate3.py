"""GATE 3 surface (3b.10): `autodeck render`, `autodeck gate3`, and the final assessment.

Drive the CLI the way `tests/test_owner_run.py` does (fake providers, no live calls). The
aesthetic loop needs LibreOffice for its default rasteriser: CLI tests that reach it are
`render`-marked; the assessment and report tests are not.
"""

from __future__ import annotations

import inspect
import json
from pathlib import Path

import pytest
from pptx import Presentation
from typer.testing import CliRunner, Result

import autodeck.providers.registry as registry_module
from autodeck.agents.art_direction import ArtDirectionPlan
from autodeck.audit.gate3 import (
    HUMAN_CHECKS,
    FinalAssessment,
    _count_icons,
    assess_final,
    render_final_report,
)
from autodeck.audit.manifest import Manifest, canonical_pptx_digest
from autodeck.audit.numeric_linter import NumericReport
from autodeck.audit.post_render import PostRenderFinding, PostRenderReport
from autodeck.cli import app
from autodeck.design.components.catalog import COMPONENT_LIB_VERSION
from autodeck.design.grammar import GrammarFinding
from autodeck.design.headers.flow import HeaderFlowReport
from autodeck.design.headers.prompt import resolve_profile
from autodeck.ir.actions import AssignIcons, SetAccent, apply_action, facts_digest
from autodeck.ir.models import (
    Block,
    ChartSeries,
    ChartSpec,
    Citation,
    Deck,
    DeckBrief,
    DiagramSpec,
    KeyMessage,
    LabelFraming,
    ProcessFlowSpec,
    ProcessStep,
    Slide,
)
from autodeck.ir.store import IRStore
from autodeck.pipeline.orchestrator import (
    ApprovalState,
    ArtifactMissing,
    Gate,
    GateBlocked,
    Orchestrator,
)
from autodeck.providers.base import ProviderError, RateLimitError
from autodeck.render.qa.aesthetic import catalog_slot_lookup, library_concept_lookup
from autodeck.render.qa.deterministic import QAFinding
from autodeck.render.renderer import render_deck
from tests.test_actions import make_deck, make_icon_block
from tests.test_aesthetic import FakeCritic, _bullets_slide, _framing_block, reply, tokens_for
from tests.test_art_direction import FakeArtModel
from tests.test_knowledge import build_knowledge

runner = CliRunner()


def _invoke(runs_root: Path, *args: str) -> Result:
    return runner.invoke(app, [*args, "--runs-root", str(runs_root)])


def _gate3(runs_root: Path, knowledge: Path) -> Result:
    return _invoke(runs_root, "gate3", "r1", "--knowledge-root", str(knowledge))


def _fixture_deck() -> Deck:
    """Four slides that render for real: a cited claim with bullets, a native diagram, an
    `icon_pillars` slide whose icons only art direction can supply, and a native chart."""
    diagram = DiagramSpec(
        relationship="sequence",
        kind="process_flow",
        process_flow=ProcessFlowSpec(
            steps=[
                ProcessStep(
                    id="n1", label="Pilot", order=1, framing=LabelFraming(reason="stage_name")
                ),
                ProcessStep(
                    id="n2", label="Rollout", order=2, framing=LabelFraming(reason="stage_name")
                ),
            ]
        ),
    )
    chart = ChartSpec(
        chart_type="bar",
        categories=["Q4"],
        series=[ChartSeries(name="Cost per token ($)", values=[0.42])],
        source_citations=[
            Citation.for_quote(
                quote="Cost per token fell to $0.42 in Q4.",
                doc_id="finance",
                page=2,
                bbox=(10.0, 20.0, 300.0, 44.0),
                retrieved_by="validator",
            )
        ],
    )
    return Deck(
        run_id="r1",
        project="attention-efficiency",
        client="acme",
        audience="CTO",
        version=1,
        theme_ref="config/tokens/dev.json",
        component_lib_version=COMPONENT_LIB_VERSION,
        slides=[
            _bullets_slide("s1", with_notes=True),
            Slide(
                id="s2",
                narrative_role="plan",
                component="framework_diagram",
                message_ids=["m2"],
                blocks=[
                    _framing_block(
                        "headline", "A two-step rollout gets the platform live", block_id="s2-h"
                    ),
                    Block(id="s2-d", kind="diagram", slot="diagram", diagram=diagram),
                ],
            ),
            Slide(
                id="s3",
                narrative_role="capabilities",
                component="icon_pillars",
                message_ids=["m2"],
                blocks=[
                    _framing_block("headline", "Three things we do well", block_id="s3-h"),
                    _framing_block("pillar_label", "Fast", block_id="s3-l1"),
                    _framing_block("pillar_label", "Safe", block_id="s3-l2"),
                    _framing_block("pillar_label", "Cheap", block_id="s3-l3"),
                ],
            ),
            Slide(
                id="s4",
                narrative_role="evidence",
                component="chart_focus",
                message_ids=["m1"],
                blocks=[
                    _framing_block("headline", "Cost per token fell", block_id="s4-h"),
                    Block(id="s4-c", kind="chart", slot="chart", chart=chart),
                ],
            ),
        ],
    )


def make_gate3_run(tmp_path: Path, *, approve: bool = True) -> tuple[Path, Path]:
    """A run whose claims stage is validated (and, by default, approved): `(runs_root,
    knowledge_root)`. The validated deck is the fixture above, as `validate` would save it."""
    runs_root = tmp_path / "runs"
    knowledge_root = build_knowledge(tmp_path)
    orchestrator = Orchestrator("r1", runs_root=runs_root, env="dev")
    orchestrator.ir.save_brief(
        DeckBrief(
            run_id="r1",
            objective="Decide whether to fund the serving rewrite.",
            audience="CTO",
            key_messages=[
                KeyMessage(
                    id="m1", text="The rewrite cut cost per token.", evidence_status="supported"
                ),
                KeyMessage(
                    id="m2", text="Rollout can start in Q1.", evidence_status="supported"
                ),
            ],
            approved_by="tester",
        )
    )
    deck = _fixture_deck()
    orchestrator.run_stage(
        "validate", lambda: f"wrote {orchestrator.save_ir(deck)}", force=True
    )
    if approve:
        orchestrator.approve(Gate.CLAIMS)
    return runs_root, knowledge_root


def patch_roles(monkeypatch: pytest.MonkeyPatch, **by_role: object) -> None:
    monkeypatch.setattr(
        registry_module.ModelRegistry, "provider_for", lambda self, role, **kw: by_role[role]
    )


def art_plan() -> ArtDirectionPlan:
    return ArtDirectionPlan(
        rationale="Give the capability slide its icons.",
        actions=[
            AssignIcons(
                slide_id="s3", slot="pillar_icon", concepts=["speed", "security", "growth"]
            )
        ],
    )


def _presentation_only_change(deck: Deck) -> Deck:
    """What art direction and the aesthetic loop do: style, component, mode, icons."""
    copy = deck.model_copy(deep=True, update={"version": deck.version + 1})
    slide = copy.slides[0]
    slide.style.type_scale = "compact"
    slide.style.accent = "accent5"
    slide.component = "alt_layout"
    slide.communication_mode = "diagram_led"
    slide.blocks[0].slot = "renamed_slot"
    icon_block = next(block for block in slide.blocks if block.kind == "icon")
    assert icon_block.icon is not None
    icon_block.icon.glyph_id = "new-glyph"
    slide.blocks.append(make_icon_block("b9", slot="icon2"))
    return copy


def test_claims_approval_is_bound_to_facts_so_presentation_changes_keep_it(
    tmp_path: Path,
) -> None:
    """Approve claims; save a new IR version differing only in style/component/mode/icons →
    `approval_state(CLAIMS)` is CURRENT. Save one with a changed claim text → ArtifactMissing
    detail says the facts changed after validation; `require_gate` raises."""
    orchestrator = Orchestrator("r1", runs_root=tmp_path)
    validated = make_deck(run_id="r1")
    orchestrator.run_stage(
        "validate", lambda: f"wrote {orchestrator.save_ir(validated)}", force=True
    )
    orchestrator.approve(Gate.CLAIMS)
    assert orchestrator.approval_state(Gate.CLAIMS)[0] is ApprovalState.CURRENT

    presented = _presentation_only_change(validated)
    orchestrator.save_ir(presented)
    reopened = Orchestrator("r1", runs_root=tmp_path)
    assert reopened.ir.latest_version() == 2
    assert reopened.approval_state(Gate.CLAIMS) == (ApprovalState.CURRENT, "approved")
    reopened.require_gate(Gate.CLAIMS)  # must not raise

    changed = presented.model_copy(deep=True, update={"version": 3})
    changed_claim = changed.slides[0].blocks[0].claim
    assert changed_claim is not None
    changed_claim.text = "A different claim entirely."
    orchestrator.save_ir(changed)

    stale = Orchestrator("r1", runs_root=tmp_path)
    state, detail = stale.approval_state(Gate.CLAIMS)
    assert state is ApprovalState.NO_ARTIFACT
    assert "the facts changed after validation (IR v3)" in detail
    with pytest.raises(GateBlocked, match="facts changed after validation"):
        stale.require_gate(Gate.CLAIMS)
    with pytest.raises(ArtifactMissing, match="facts changed after validation"):
        stale.approve(Gate.CLAIMS)


def test_facts_digest_is_stable_and_sensitive() -> None:
    """Same deck twice → same digest; a deep copy → same; one changed citation quote,
    verdict, or framing text → different; changed style or an added icon block → same."""
    deck = make_deck()
    digest = facts_digest(deck)
    assert len(digest) == 64
    assert facts_digest(deck) == digest
    assert facts_digest(deck.model_copy(deep=True)) == digest
    assert facts_digest(make_deck()) == digest

    quote = deck.model_copy(deep=True)
    claim = quote.slides[0].blocks[0].claim
    assert claim is not None
    claim.citations[0].quote = "A different quote entirely."
    assert facts_digest(quote) != digest

    verdict = deck.model_copy(deep=True)
    claim = verdict.slides[0].blocks[0].claim
    assert claim is not None
    claim.verdict = "contradicted"
    assert facts_digest(verdict) != digest

    framing = deck.model_copy(deep=True)
    framing.slides[0].blocks[1].text = "A completely different framing sentence."
    assert facts_digest(framing) != digest

    assert facts_digest(_presentation_only_change(deck)) == digest


def test_render_refuses_without_a_current_claims_approval(tmp_path: Path) -> None:
    runs_root, _knowledge = make_gate3_run(tmp_path, approve=False)

    refused = _invoke(runs_root, "render", "r1")
    assert refused.exit_code == 3, refused.output
    assert "GATE 'claims'" in refused.output
    assert "autodeck approve r1 claims" in refused.output
    assert not (runs_root / "r1" / "deck.pptx").exists()
    assert IRStore(runs_root, "r1").versions() == [1]

    # An approval that no longer covers the facts does not count either.
    Orchestrator("r1", runs_root=runs_root).approve(Gate.CLAIMS)
    store = IRStore(runs_root, "r1")
    changed = store.load(1).model_copy(deep=True, update={"version": 2})
    claim = changed.slides[0].blocks[0].claim
    assert claim is not None
    claim.text = "A claim nobody approved."
    store.save(changed)
    stale = _invoke(runs_root, "render", "r1")
    assert stale.exit_code == 3, stale.output
    assert "facts changed after validation" in stale.output
    assert not (runs_root / "r1" / "deck.pptx").exists()


def test_an_unknown_run_is_refused_by_both_commands(tmp_path: Path) -> None:
    for command in ("render", "gate3"):
        result = _invoke(tmp_path / "runs", command, "nope")
        assert result.exit_code == 1, result.output
        assert "no such run" in result.output


def test_gate3_refuses_before_render(tmp_path: Path) -> None:
    runs_root, knowledge = make_gate3_run(tmp_path)

    refused = _invoke(runs_root, "gate3", "r1", "--knowledge-root", str(knowledge))

    assert refused.exit_code == 3, refused.output
    assert "autodeck render r1" in refused.output
    assert not (runs_root / "r1" / "final_audit_report.md").exists()
    assert not (runs_root / "r1" / "build_manifest.json").exists()

    # A completed render stage with the deck missing is refused the same way.
    orchestrator = Orchestrator("r1", runs_root=runs_root)
    orchestrator.run_stage("render", lambda: "claimed", force=True)
    missing = _invoke(runs_root, "gate3", "r1", "--knowledge-root", str(knowledge))
    assert missing.exit_code == 3, missing.output
    assert "autodeck render r1" in missing.output


def _assessment(
    *,
    post_render_findings: tuple[PostRenderFinding, ...] = (),
    qa: tuple[QAFinding, ...] = (),
    grammar_findings: tuple[GrammarFinding, ...] = (),
    charts: int = 1,
    diagrams: int = 1,
    icons: int = 3,
) -> FinalAssessment:
    return FinalAssessment(
        pptx=Path("runs/r1/deck.pptx"),
        deck_digest="d" * 64,
        post_render=PostRenderReport(numeric=NumericReport(), findings=post_render_findings),
        qa=list(qa),
        grammar=list(grammar_findings),
        flow=HeaderFlowReport(),
        chart_count=charts,
        diagram_count=diagrams,
        icon_count=icons,
    )


_POST_RENDER = PostRenderFinding(
    check="claim altered in render", slide_id="s2", detail="expected 'A', found 'B'"
)
_QA = QAFinding(
    check="below minimum size",
    slide_id="s3",
    shapes=("Text 4",),
    measured=9.0,
    threshold=12.0,
    remedy="token",
    remedy_detail="typography.minimum",
)
_BLOCKING_GRAMMAR = GrammarFinding(
    check="icon adjacent to text label", severity="blocking", detail="far from text"
)
_ADVISORY_GRAMMAR = GrammarFinding(check="word budget", severity="advisory", detail="long")


def test_checkable_criteria_and_their_order() -> None:
    """Build `FinalAssessment`s by hand: all clean → four PASS; one post-render finding,
    one QA finding, one blocking grammar finding, zero icons — each fails its own line only."""
    clean = _assessment()
    results = clean.checkable()
    assert [label for label, _passed, _detail in results] == [
        "Post-render audit passes (every claim on its slide, every number traced)",
        "No deterministic QA findings",
        "No blocking grammar findings",
        "At least one native chart, one native diagram and one theme-recolourable icon",
    ]
    assert [passed for _label, passed, _detail in results] == [True] * 4
    assert clean.passes
    assert results[3][2] == "1 chart(s), 1 diagram(s), 3 icon(s)"

    broken = {
        0: _assessment(post_render_findings=(_POST_RENDER,)),
        1: _assessment(qa=(_QA,)),
        2: _assessment(grammar_findings=(_BLOCKING_GRAMMAR,)),
        3: _assessment(icons=0),
    }
    # Each of the three milestone counts fails the same line, naming all three counts.
    for missing in (_assessment(charts=0), _assessment(diagrams=0), _assessment(icons=0)):
        assert [ok for _l, ok, _d in missing.checkable()] == [True, True, True, False]
    for index, assessment in broken.items():
        passed = [ok for _label, ok, _detail in assessment.checkable()]
        assert passed == [position != index for position in range(4)], index
        assert not assessment.passes
    assert "expected 'A', found 'B'" in broken[0].checkable()[0][2]
    assert "below minimum size" in broken[1].checkable()[1][2]
    assert "far from text" in broken[2].checkable()[2][2]
    assert broken[3].checkable()[3][2] == "1 chart(s), 1 diagram(s), 0 icon(s)"
    assert _assessment(charts=0).checkable()[3][2] == "0 chart(s), 1 diagram(s), 3 icon(s)"

    # Advisory grammar findings and header-flow findings are reported, never a criterion.
    advisory = _assessment(grammar_findings=(_ADVISORY_GRAMMAR,))
    assert advisory.passes


def test_the_report_lists_human_checks_unticked_and_says_nothing_checked_them() -> None:
    """Even for an assessment that passes every checkable criterion, every `HUMAN_CHECKS`
    item appears as `- [ ]`, and no `[x]` appears anywhere in the report."""
    assessment = _assessment(grammar_findings=(_ADVISORY_GRAMMAR,))
    assert assessment.passes
    report = render_final_report(assessment, audit_report="# Audit report - Demo\n")

    assert "[x]" not in report.lower()
    for check in HUMAN_CHECKS:
        assert f"- [ ] {check}" in report
    assert report.count("- [ ]") == len(HUMAN_CHECKS) == 5
    human_heading = report.index("Human checks")
    assert "by a person" in report[human_heading:]
    assert "Nothing in this report has checked them" in report[human_heading:]

    # Sections in the contracted order, human checks last.
    order = [
        report.index("Deck: "),
        report.index("Deck digest: "),
        report.index("## Checkable criteria"),
        report.index("## Findings"),
        report.index("## Header flow"),
        report.index("# Audit report - Demo"),
        human_heading,
    ]
    assert order == sorted(order)
    assert report.count("[PASS]") == 4
    assert "d" * 64 in report
    assert "[advisory] word budget: long" in report


def test_the_report_shows_every_failure_in_full_and_is_deterministic() -> None:
    assessment = _assessment(
        post_render_findings=(_POST_RENDER,),
        qa=(_QA,),
        grammar_findings=(_BLOCKING_GRAMMAR,),
        icons=0,
    )
    report = render_final_report(assessment, audit_report="audit")
    assert report.count("[FAIL]") == 4
    assert "4 of 4 checkable criteria fail." in report
    assert "slide s2: expected 'A', found 'B'" in report
    assert "measured 9 against 12; fix by token (typography.minimum)" in report
    assert "[BLOCKING] icon adjacent to text label: far from text" in report
    assert report == render_final_report(assessment, audit_report="audit")
    # Still unticked when everything fails.
    assert report.count("- [ ]") == 5


def _deck_with_icons() -> Deck:
    return apply_action(
        _fixture_deck(),
        AssignIcons(
            slide_id="s3", slot="pillar_icon", concepts=["speed", "security", "growth"]
        ),
        slots_of=catalog_slot_lookup(),
        glyph_for=library_concept_lookup(),
    )


def test_icons_are_counted_per_instance_from_the_rendered_file(tmp_path: Path) -> None:
    """`place_icon` makes one shape per subpath, all named `icon:<family>:<name>`; three icons
    on a slide are three, however many strokes each has. A deck with none counts zero."""
    tokens = tokens_for()
    with_icons = tmp_path / "with.pptx"
    render_deck(_deck_with_icons(), tokens=tokens, out_path=with_icons)
    without = tmp_path / "without.pptx"
    render_deck(make_two_slide_deck_for_rendering(), tokens=tokens, out_path=without)

    assert _count_icons(Presentation(str(with_icons))) == 3
    assert _count_icons(Presentation(str(without))) == 0


def make_two_slide_deck_for_rendering() -> Deck:
    return _fixture_deck().model_copy(update={"slides": _fixture_deck().slides[:2]})


def test_assess_final_counts_charts_diagrams_and_icons_from_the_rendered_file(
    tmp_path: Path,
) -> None:
    tokens = tokens_for()
    deck = _deck_with_icons()
    pptx = tmp_path / "deck.pptx"
    render_deck(deck, tokens=tokens, out_path=pptx)

    assessment = assess_final(
        deck, pptx, tokens=tokens, brief=None, profile=resolve_profile(None)
    )

    assert (assessment.chart_count, assessment.diagram_count, assessment.icon_count) == (
        1,
        1,
        3,
    )
    assert assessment.deck_digest == canonical_pptx_digest(pptx)
    assert assessment.post_render.passes
    assert assessment.passes
    assert assessment.checkable()[3][2] == "1 chart(s), 1 diagram(s), 3 icon(s)"


def test_a_deck_without_a_native_chart_fails_only_the_milestone_criterion(
    tmp_path: Path,
) -> None:
    tokens = tokens_for()
    full = _deck_with_icons()
    deck = full.model_copy(update={"slides": [s for s in full.slides if s.id != "s4"]})
    pptx = tmp_path / "deck.pptx"
    render_deck(deck, tokens=tokens, out_path=pptx)

    assessment = assess_final(
        deck, pptx, tokens=tokens, brief=None, profile=resolve_profile(None)
    )

    assert (assessment.chart_count, assessment.diagram_count, assessment.icon_count) == (
        0,
        1,
        3,
    )
    assert [ok for _label, ok, _detail in assessment.checkable()] == [True, True, True, False]
    assert not assessment.passes
    assert "0 chart(s), 1 diagram(s), 3 icon(s)" in render_final_report(
        assessment, audit_report="audit"
    )


def _approved_run_with_models(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    *,
    art: object = None,
    critic: object = None,
) -> tuple[Path, Path]:
    runs_root, knowledge = make_gate3_run(tmp_path)
    patch_roles(
        monkeypatch,
        outline=art if art is not None else FakeArtModel(art_plan()),
        aesthetic=critic if critic is not None else FakeCritic(reply(9.0, rationale="good")),
    )
    return runs_root, knowledge


@pytest.mark.render
def test_render_writes_the_deck_and_previews_and_keeps_the_claims_approval(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runs_root, _knowledge = _approved_run_with_models(tmp_path, monkeypatch)
    before = IRStore(runs_root, "r1").load(1)

    result = _invoke(runs_root, "render", "r1")

    assert result.exit_code == 0, result.output
    run_dir = runs_root / "r1"
    assert (run_dir / "deck.pptx").exists()
    assert len(list((run_dir / "previews" / "final").glob("*.png"))) == 4
    assert "Art direction" in result.output
    assert "s3" in result.output and "icon_anchored" in result.output
    assert "applied  assign_icons" in result.output
    assert "Aesthetic loop" in result.output
    assert "Stopped: target_reached" in result.output
    assert "Next: autodeck gate3 r1" in result.output

    store = IRStore(runs_root, "r1")
    assert store.versions() == [1, 2]
    after = store.load(2)
    assert facts_digest(after) == facts_digest(before)
    assert [b.kind for b in after.slides[2].blocks].count("icon") == 3
    assert [slide.communication_mode for slide in after.slides] == [
        "text_led",
        "diagram_led",
        "icon_anchored",
        "text_led",
    ]

    # The point of binding the approval to facts: the render did not void it.
    orchestrator = Orchestrator("r1", runs_root=runs_root)
    assert orchestrator.approval_state(Gate.CLAIMS) == (ApprovalState.CURRENT, "approved")
    assert not orchestrator.state.is_approved(Gate.FINAL_RENDER)
    for stage in ("art_direction", "render"):
        assert orchestrator.state.is_complete(stage)


@pytest.mark.render
def test_a_critique_that_changes_the_deck_is_saved_as_its_own_ir_version(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    accent = SetAccent(slide_id="s1", accent="accent2")
    critic = FakeCritic(reply(5.0, accent), reply(9.0))
    runs_root, _knowledge = _approved_run_with_models(tmp_path, monkeypatch, critic=critic)

    result = _invoke(runs_root, "render", "r1")

    assert result.exit_code == 0, result.output
    assert "iteration 0: score 5" in result.output
    assert "applied  set_accent" in result.output
    store = IRStore(runs_root, "r1")
    assert store.versions() == [1, 2, 3]
    assert store.load(3).slides[0].style.accent == "accent2"
    assert store.load(2).slides[0].style.accent != "accent2"
    Orchestrator("r1", runs_root=runs_root).require_gate(Gate.CLAIMS)


@pytest.mark.render
def test_a_spent_critique_quota_ends_render_with_exit_5_but_leaves_a_reviewable_deck(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    critic = FakeCritic(RateLimitError("gemini: rate limited (429)"))
    runs_root, _knowledge = _approved_run_with_models(tmp_path, monkeypatch, critic=critic)

    result = _invoke(runs_root, "render", "r1")

    assert result.exit_code == 5, result.output
    assert "Traceback" not in result.output
    assert "role 'aesthetic'" in result.output
    assert "not (fully) critiqued" in result.output
    assert "autodeck render r1" in result.output
    assert (runs_root / "r1" / "deck.pptx").exists()
    # Art direction completed and was saved; the approval still stands.
    assert IRStore(runs_root, "r1").versions() == [1, 2]
    Orchestrator("r1", runs_root=runs_root).require_gate(Gate.CLAIMS)


@pytest.mark.render
def test_a_spent_art_direction_quota_ends_render_with_exit_5_and_says_what_is_saved(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    art = FakeArtModel(RateLimitError("gemini: rate limited (429)"))
    runs_root, _knowledge = _approved_run_with_models(tmp_path, monkeypatch, art=art)

    result = _invoke(runs_root, "render", "r1")

    assert result.exit_code == 5, result.output
    assert "Traceback" not in result.output
    assert "role 'outline'" in result.output
    assert "IR versions up to v2 were written" in result.output
    assert "autodeck render r1" in result.output
    # The rule still assigned modes, but the icon slide has no icons without the model.
    assert not (runs_root / "r1" / "deck.pptx").exists()
    assert IRStore(runs_root, "r1").versions() == [1, 2]
    Orchestrator("r1", runs_root=runs_root).require_gate(Gate.CLAIMS)


@pytest.mark.render
def test_a_model_that_never_answers_validly_is_reported_not_fatal(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    critic = FakeCritic(ProviderError("the reply was not valid JSON"))
    runs_root, _knowledge = _approved_run_with_models(tmp_path, monkeypatch, critic=critic)

    result = _invoke(runs_root, "render", "r1")

    assert result.exit_code == 0, result.output
    assert "Stopped: model_error" in result.output
    assert "was not (fully) critiqued" in result.output
    assert (runs_root / "r1" / "deck.pptx").exists()


@pytest.mark.render
def test_render_then_gate3_produces_deck_report_and_manifest_together(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Full CLI path with fakes on an approved fixture run: `render` writes deck.pptx,
    previews and new IR versions without voiding the claims approval; `gate3` writes
    final_audit_report.md and build_manifest.json and prints all three paths first; then
    `approve <run> final_render` records `canonical_pptx_digest(deck.pptx)`."""
    runs_root, knowledge = _approved_run_with_models(tmp_path, monkeypatch)
    assert _invoke(runs_root, "render", "r1").exit_code == 0

    result = _gate3(runs_root, knowledge)

    run_dir = runs_root / "r1"
    assert result.exit_code == 0, result.output
    deck_path, report_path, manifest_path = (
        run_dir / "deck.pptx",
        run_dir / "final_audit_report.md",
        run_dir / "build_manifest.json",
    )
    for path in (deck_path, report_path, manifest_path):
        assert path.exists(), path
    lines = result.output.splitlines()
    first_paths = [i for i, line in enumerate(lines) if str(run_dir) in line][:3]
    assert [lines[i].split(str(run_dir))[1].lstrip("/") for i in first_paths] == [
        "deck.pptx",
        "final_audit_report.md",
        "build_manifest.json",
    ]
    assert max(first_paths) < next(i for i, line in enumerate(lines) if "[PASS]" in line)
    # The claim-level audit is in the file only; the terminal points to it (3b.10 integrity).
    assert "# Audit report" in report_path.read_text(encoding="utf-8")
    assert "# Audit report" not in result.output
    assert "claim-level audit is in" in result.output
    assert result.output.count("- [ ]") == 5 and "[x]" not in result.output
    assert Manifest.load(manifest_path).run_id == "r1"

    approved = _invoke(runs_root, "approve", "r1", "final_render")
    assert approved.exit_code == 0, approved.output
    state = json.loads((run_dir / "state.json").read_text(encoding="utf-8"))
    assert state["fingerprints"]["final_render"] == canonical_pptx_digest(deck_path)
    assert canonical_pptx_digest(deck_path)[:16] in result.output
    # The digest the report printed is the digest the approval recorded.
    assert canonical_pptx_digest(deck_path) in report_path.read_text(encoding="utf-8")


@pytest.mark.render
def test_gate3_exits_4_and_says_so_when_the_rendered_deck_fails_a_criterion(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    runs_root, knowledge = _approved_run_with_models(tmp_path, monkeypatch)
    assert _invoke(runs_root, "render", "r1").exit_code == 0
    deck_path = runs_root / "r1" / "deck.pptx"
    _tamper_claim(deck_path)

    result = _gate3(runs_root, knowledge)

    assert result.exit_code == 4, result.output
    assert "[FAIL] Post-render audit passes" in result.output
    assert (runs_root / "r1" / "final_audit_report.md").exists()
    assert "Approve with `autodeck approve r1 final_render`" in result.output


def _tamper_claim(pptx: Path) -> None:
    """Change one word of a claim inside the rendered file, the way a rendering bug might."""
    import shutil
    import zipfile

    scratch = pptx.with_suffix(".tampered")
    with zipfile.ZipFile(pptx) as source, zipfile.ZipFile(scratch, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "ppt/slides/slide1.xml":
                assert b"cut cost" in data
                data = data.replace(b"cut cost", b"raised cost")
            target.writestr(item, data)
    shutil.move(scratch, pptx)


def test_neither_command_can_record_an_approval() -> None:
    """Extend the existing walk (`test_none_of_the_new_commands_can_record_an_approval` in
    tests/test_cli_gate2.py) to `render` and `gate3`: no bypass-shaped option, and no call to
    `.approve(` in either command's own source or anything they define."""
    from autodeck import cli

    banned = {"yes", "force", "skip_gates", "no_gates", "auto_approve", "approve", "trust"}
    # No way to skip a stage either: render's only options are the run, where runs live, and
    # which environment's models to use.
    assert set(inspect.signature(cli.render).parameters) == {"run_id", "runs_root", "env"}
    for name in ("render", "gate3"):
        func = getattr(cli, name)
        assert not (set(inspect.signature(func).parameters) & banned), name
        assert ".approve(" not in inspect.getsource(func), name
        assert "approve(" not in inspect.getsource(func).replace("autodeck approve", ""), name
