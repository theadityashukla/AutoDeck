
import os
from pptx import Presentation
# Mock PPTGenerator to isolate the issue, or import it if better. 
# Importing is better to test actual code.
import sys
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
from autodeck_core.ppt_generator import PPTGenerator

def debug_ppt_content():
    # Mock Session Data with explicit bullets and image
    # make sure to use a valid image path or mock one?
    # I'll check if I can find a valid image in the workspace or just skip image key to test bullets first, 
    # then add image key to test interaction.
    
    # Check for an existing image
    img_dir = "extracted_images"
    valid_img = None
    if os.path.exists(img_dir):
        for root, dirs, files in os.walk(img_dir):
            for f in files:
                if f.endswith(".png"):
                    valid_img = os.path.join(root, f)
                    break
            if valid_img: break
    
    print(f"Using image: {valid_img}")

    session_data = {
        "name": "Debug Session",
        "outline": [
            {"title": "Slide 1", "description": "Desc 1"}
        ],
        "content": {
            "0": {
                "title": "Conclusion: The Future is Self-Supervised",
                "bullet_points": [
                    "Self-supervised learning (SSL) is poised to revolutionize AI by enabling machines to learn from unlabeled data.",
                    "LeJEPA introduces a theoretically grounded framework for Joint-Embedding Predictive Architectures (JEPAs), simplifying training and boosting performance.",
                    "Our research demonstrates LeJEPA’s scalability, stability, and impressive results across diverse datasets, architectures, and domains – even surpassing state-of-the-art transfer learning.",
                    "The future of AI lies in learning the world's dynamics through self-supervision, unlocking unprecedented capabilities."
                ],
                "image_suggestion": valid_img,
                "speaker_notes": "Notes here."
            }
        }
    }

    output_path = "debug_presentation.pptx"
    if os.path.exists(output_path):
        os.remove(output_path)

    generator = PPTGenerator(output_path)
    generator.generate(session_data)

    print("PPT Generated.")
    
    # Inspect
    prs = Presentation(output_path)
    slide = prs.slides[1] # Slide 0 is title, Slide 1 is content
    
    # Check shapes
    print(f"Shape count: {len(slide.shapes)}")
    found_text = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                found_text.append(p.text)
                print(f"Found paragraph: '{p.text}'")
    
    expected_snippets = ["Self-supervised learning", "LeJEPA introduces", "Our research demonstrates", "The future of AI"]
    missing = [s for s in expected_snippets if not any(s in t for t in found_text)]
    
    if missing:
        print(f"FAILURE: Missing content: {missing}")
    else:
        print("SUCCESS: All content found in text frames.")

if __name__ == "__main__":
    debug_ppt_content()
